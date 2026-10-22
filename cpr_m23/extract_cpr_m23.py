#!/usr/bin/env python3
"""Extract filename dates and M23 values into an Excel file in Downloads.

Install: python -m pip install openpyxl xlrd pyxlsb
Run:     python extract_cpr_m23.py
Sheet:   python extract_cpr_m23.py --sheet "Summary"

Requires Python 3.10+. Formula cells must already have saved numeric results.
"""

from __future__ import annotations

import argparse
import ctypes
from datetime import date, datetime
from io import BytesIO
import math
from numbers import Real
import os
from pathlib import Path
import re
import sys
from uuid import UUID


DEFAULT_FOLDER = Path(r"G:\FIN_COMM\DeltaOne\Kenny\a")
FILENAME = re.compile(r"457276\s+CPR\s+([0-9]{8})", re.IGNORECASE)
EXTENSIONS = {".xlsx", ".xlsm", ".xls", ".xlsb"}


def downloads_folder() -> Path:
    """Resolve Windows Downloads, including a redirected Downloads folder."""
    if os.name != "nt":
        return Path.home() / "Downloads"

    class GUID(ctypes.Structure):
        _fields_ = [
            ("data1", ctypes.c_uint32),
            ("data2", ctypes.c_uint16),
            ("data3", ctypes.c_uint16),
            ("data4", ctypes.c_ubyte * 8),
        ]

    # FOLDERID_Downloads from the Windows Known Folders API.
    folder_id = GUID.from_buffer_copy(
        UUID("374DE290-123F-4565-9164-39C4925E467B").bytes_le
    )
    shell32 = ctypes.WinDLL("shell32")
    ole32 = ctypes.WinDLL("ole32")
    get_path = shell32.SHGetKnownFolderPath
    get_path.argtypes = [
        ctypes.POINTER(GUID), ctypes.c_uint32, ctypes.c_void_p,
        ctypes.POINTER(ctypes.c_wchar_p),
    ]
    get_path.restype = ctypes.c_long
    ole32.CoTaskMemFree.argtypes = [ctypes.c_void_p]
    ole32.CoTaskMemFree.restype = None
    result = ctypes.c_wchar_p()
    try:
        status = get_path(ctypes.byref(folder_id), 0, None, ctypes.byref(result))
        if status != 0 or not result.value:
            raise OSError(f"Cannot locate Downloads (Windows status {status}).")
        return Path(result.value)
    finally:
        ole32.CoTaskMemFree(ctypes.cast(result, ctypes.c_void_p))


def find_sources(folder: Path) -> list[tuple[date, Path]]:
    if not folder.is_dir():
        raise FileNotFoundError(f"Source folder not found: {folder}")

    def scan_error(error: OSError) -> None:
        raise error

    matches = []
    for root, directories, filenames in os.walk(
        folder, onerror=scan_error, followlinks=False
    ):
        # Do not follow folder links/junctions into loops or other trees.
        directories[:] = [
            name for name in directories
            if not (Path(root) / name).is_symlink()
            and not (
                getattr((Path(root) / name).lstat(), "st_file_attributes", 0)
                & 0x400
            )
        ]
        for filename in filenames:
            path = Path(root) / filename
            if path.suffix.lower() not in EXTENSIONS or path.is_symlink():
                continue
            match = FILENAME.fullmatch(path.stem)
            if not match:
                continue
            try:
                file_date = datetime.strptime(match.group(1), "%Y%m%d").date()
                if file_date.year < 1900:
                    raise ValueError("Excel dates must be 1900 or later")
            except ValueError as error:
                raise ValueError(f"Invalid filename date in {path}: {error}") from error
            matches.append((file_date, path))
    if not matches:
        raise ValueError(f"No files matching '457276 CPR YYYYMMDD' found in {folder}.")
    return sorted(matches, key=lambda item: (item[0], str(item[1]).casefold()))


def read_m23(path: Path, sheet_name: str | None = None) -> float:
    """Read the saved value, without opening Excel or changing the source."""
    extension = path.suffix.lower()
    if extension in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
        try:
            sheet = workbook[sheet_name] if sheet_name else workbook.worksheets[0]
            cell = sheet["M23"]
            if cell.data_type == "e":
                raise ValueError(f"M23 contains an Excel error: {cell.value}")
            value = cell.value
        finally:
            workbook.close()
    elif extension == ".xls":
        import xlrd

        workbook = xlrd.open_workbook(str(path), on_demand=True)
        try:
            sheet = (
                workbook.sheet_by_name(sheet_name)
                if sheet_name else workbook.sheet_by_index(0)
            )
            if sheet.nrows <= 22 or sheet.ncols <= 12:
                value = None
            else:
                cell = sheet.cell(22, 12)
                if cell.ctype != xlrd.XL_CELL_NUMBER:
                    raise ValueError("M23 is blank, an error, or not a numeric cell.")
                value = cell.value
        finally:
            workbook.release_resources()
    elif extension == ".xlsb":
        from pyxlsb import open_workbook

        value = None
        with open_workbook(str(path)) as workbook:
            with workbook.get_sheet(sheet_name if sheet_name else 1) as sheet:
                for row in sheet.rows(sparse=True):
                    if not row:
                        continue
                    if row[0].r > 22:
                        break
                    if row[0].r == 22:
                        value = next((cell.v for cell in row if cell.c == 12), None)
                        break
    else:
        raise ValueError(f"Unsupported Excel extension: {extension}")

    if value is None:
        raise ValueError(
            "M23 is blank or has no saved formula result. If it is a formula, "
            "recalculate and save the source in Excel, then run again."
        )
    if isinstance(value, bool) or not isinstance(value, Real):
        raise ValueError(f"M23 must contain a number; found {value!r}.")
    if not math.isfinite(value):
        raise ValueError("M23 is not a finite number.")
    return float(value)


def extract_rows(folder: Path, sheet_name: str | None = None) -> list[tuple[date, float]]:
    rows = []
    for file_date, path in find_sources(folder):
        try:
            rows.append((file_date, read_m23(path, sheet_name)))
        except Exception as error:
            raise ValueError(f"Could not extract {path}: {error}") from error
    return rows


def write_report(rows: list[tuple[date, float]], directory: Path | None = None) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    if not rows:
        raise ValueError("No rows to write.")
    destination = downloads_folder() if directory is None else directory
    destination.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "CPR M23"
    sheet.append(["YYYYMMDD", "M23"])
    for file_date, value in sorted(rows, key=lambda row: row[0]):
        sheet.append([file_date, value])
        sheet.cell(sheet.max_row, 1).number_format = "yyyymmdd"
    for cell in sheet[1]:
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center")
    sheet.column_dimensions["A"].width = 15
    sheet.column_dimensions["B"].width = 24
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False

    # Build the complete file before publishing it; never overwrite an earlier run.
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    content = buffer.getvalue()
    suffix = 1
    while True:
        tail = "" if suffix == 1 else f"_{suffix}"
        path = destination / f"CPR_M23_summary{tail}.xlsx"
        try:
            stream = path.open("xb")
        except FileExistsError:
            suffix += 1
            continue
        try:
            with stream:
                stream.write(content)
        except BaseException:
            path.unlink(missing_ok=True)
            raise
        return path


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--folder", type=Path, default=DEFAULT_FOLDER,
                        help="Source folder (all subfolders are included).")
    parser.add_argument("--sheet", help="Worksheet name; defaults to the first worksheet.")
    args = parser.parse_args(argv)
    try:
        rows = extract_rows(args.folder, args.sheet)
        output = write_report(rows)
    except Exception as error:
        print(f"Extraction failed: {error}", file=sys.stderr)
        print("No complete report was produced by this run.", file=sys.stderr)
        return 1
    print(f"Saved {len(rows)} rows to {output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
