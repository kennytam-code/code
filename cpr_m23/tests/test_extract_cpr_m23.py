"""Run with: python -m unittest discover -s tests -v"""

from contextlib import redirect_stderr, redirect_stdout
from datetime import date
import hashlib
from io import StringIO
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
from zipfile import ZipFile, ZIP_DEFLATED

from openpyxl import Workbook, load_workbook

import extract_cpr_m23 as cpr


class ExtractionTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.source = self.root / "source"
        self.source.mkdir()
        self.downloads = self.root / "Downloads"

    def make_source(self, name, value, second_value=None):
        path = self.source / name
        path.parent.mkdir(parents=True, exist_ok=True)
        workbook = Workbook()
        workbook.active["M23"] = value
        if second_value is not None:
            workbook.create_sheet("Summary")["M23"] = second_value
        workbook.save(path)
        workbook.close()
        return path

    def test_recursive_matching_sorting_and_numeric_zero(self):
        self.make_source("457276 CPR 20260920.xlsx", 7.25)
        self.make_source("sub/457276 CPR 20260918.XLSX", 0)
        self.make_source("457276 CPR 20260919.xlsx", -3.5)
        self.make_source("457276 CPR 20260919 copy.xlsx", 999)
        self.make_source("~$457276 CPR 20260919.xlsx", 999)
        self.make_source("457277 CPR 20260919.xlsx", 999)
        self.assertEqual(cpr.extract_rows(self.source), [
            (date(2026, 9, 18), 0),
            (date(2026, 9, 19), -3.5),
            (date(2026, 9, 20), 7.25),
        ])

    def test_named_sheet_and_duplicate_dates(self):
        self.make_source("457276 CPR 20260918.xlsx", 10, 100)
        self.make_source("sub/457276 CPR 20260918.xlsx", 20, 200)
        self.assertEqual(cpr.extract_rows(self.source, "Summary"), [
            (date(2026, 9, 18), 100), (date(2026, 9, 18), 200),
        ])

    def test_invalid_date_stops_extraction(self):
        self.make_source("457276 CPR 20260230.xlsx", 1)
        with self.assertRaisesRegex(ValueError, "Invalid filename date"):
            cpr.extract_rows(self.source)

    def test_blank_text_boolean_error_and_uncached_formula_are_rejected(self):
        for value in (None, "123.45", "", True, "#DIV/0!", "=1+2"):
            with self.subTest(value=value):
                path = self.make_source("457276 CPR 20260918.xlsx", value)
                with self.assertRaisesRegex(ValueError, "Could not extract"):
                    cpr.extract_rows(self.source)
                path.unlink()

    def test_saved_formula_value_is_extracted_without_changing_source(self):
        path = self.make_source("457276 CPR 20260918.xlsx", "=1+2")
        # Insert a saved result, as Excel would when saving a calculated workbook.
        with ZipFile(path) as source_zip:
            files = {name: source_zip.read(name) for name in source_zip.namelist()}
        sheet_xml = files["xl/worksheets/sheet1.xml"]
        self.assertIn(b"<f>1+2</f><v></v>", sheet_xml)
        files["xl/worksheets/sheet1.xml"] = sheet_xml.replace(
            b"<f>1+2</f><v></v>", b"<f>1+2</f><v>3</v>"
        )
        with ZipFile(path, "w", ZIP_DEFLATED) as target_zip:
            for name, content in files.items():
                target_zip.writestr(name, content)
        before = hashlib.sha256(path.read_bytes()).hexdigest()
        self.assertEqual(cpr.extract_rows(self.source), [(date(2026, 9, 18), 3)])
        self.assertEqual(hashlib.sha256(path.read_bytes()).hexdigest(), before)

    def test_cli_saves_two_columns_to_downloads_and_preserves_prior_report(self):
        self.make_source("457276 CPR 20260918.xlsx", 123.456)
        with patch.object(cpr, "downloads_folder", return_value=self.downloads):
            with redirect_stdout(StringIO()):
                self.assertEqual(cpr.main(["--folder", str(self.source)]), 0)
                original = (self.downloads / "CPR_M23_summary.xlsx").read_bytes()
                self.assertEqual(cpr.main(["--folder", str(self.source)]), 0)
        self.assertEqual((self.downloads / "CPR_M23_summary.xlsx").read_bytes(), original)
        report = self.downloads / "CPR_M23_summary_2.xlsx"
        self.assertTrue(report.is_file())
        workbook = load_workbook(report, data_only=True)
        try:
            sheet = workbook.active
            self.assertEqual((sheet.max_row, sheet.max_column), (2, 2))
            self.assertEqual([cell.value for cell in sheet[1]], ["YYYYMMDD", "M23"])
            self.assertEqual(sheet["A2"].value.date(), date(2026, 9, 18))
            self.assertEqual(sheet["A2"].number_format, "yyyymmdd")
            self.assertEqual(sheet["B2"].value, 123.456)
            self.assertEqual(sheet.freeze_panes, "A2")
            self.assertEqual(sheet.auto_filter.ref, "A1:B2")
        finally:
            workbook.close()

    def test_failure_does_not_create_partial_report(self):
        self.make_source("457276 CPR 20260918.xlsx", 5)
        self.make_source("457276 CPR 20260919.xlsx", None)
        with patch.object(cpr, "downloads_folder", return_value=self.downloads):
            with redirect_stderr(StringIO()) as errors:
                self.assertEqual(cpr.main(["--folder", str(self.source)]), 1)
        self.assertIn("20260919.xlsx", errors.getvalue())
        self.assertFalse(self.downloads.exists())

    def test_missing_sheet_identifies_source(self):
        self.make_source("457276 CPR 20260918.xlsx", 5)
        with self.assertRaisesRegex(ValueError, "20260918.xlsx"):
            cpr.extract_rows(self.source, "Missing")

    def test_no_matches_is_an_error(self):
        with self.assertRaisesRegex(ValueError, "No files matching"):
            cpr.extract_rows(self.source)

    def test_all_supported_extensions_are_discovered(self):
        for extension in cpr.EXTENSIONS:
            (self.source / f"457276 CPR 20260918{extension}").touch()
        self.assertEqual(len(cpr.find_sources(self.source)), 4)


if __name__ == "__main__":
    unittest.main()
