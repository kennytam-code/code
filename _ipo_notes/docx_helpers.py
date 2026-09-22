"""Small python-docx helper layer for the pair-trade note (no pandas)."""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
GREY = RGBColor(0x59, 0x59, 0x59)
GREEN = RGBColor(0x1E, 0x7B, 0x34)
RED = RGBColor(0xB0, 0x2A, 0x2A)


def new_doc():
    d = Document()
    sec = d.sections[0]
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin = sec.bottom_margin = Cm(1.8)
    st = d.styles['Normal']
    st.font.name = 'Calibri'
    st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft JhengHei')
    for name, size in (('Heading 1', 14), ('Heading 2', 12), ('Heading 3', 11)):
        h = d.styles[name]
        h.font.name = 'Calibri'
        h.font.size = Pt(size)
        h.font.color.rgb = NAVY
        h.font.bold = True
        h.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Microsoft JhengHei')
    t = d.styles['Title']
    t.font.name = 'Calibri'
    t.font.size = Pt(20)
    t.font.color.rgb = NAVY
    t.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Microsoft JhengHei')
    return d


def title(d, text, sub=None):
    d.add_paragraph(text, style='Title')
    if sub:
        p = d.add_paragraph()
        r = p.add_run(sub)
        r.font.size = Pt(9.5)
        r.font.color.rgb = GREY
        r.italic = True


def h1(d, text):
    return d.add_heading(text, level=1)


def h2(d, text):
    return d.add_heading(text, level=2)


def h3(d, text):
    return d.add_heading(text, level=3)


def para(d, text, bold_lead=None, size=None, italic=False, color=None):
    p = d.add_paragraph()
    if bold_lead:
        r = p.add_run(bold_lead)
        r.bold = True
        if size: r.font.size = Pt(size)
    r = p.add_run(text)
    r.italic = italic
    if size: r.font.size = Pt(size)
    if color: r.font.color.rgb = color
    p.paragraph_format.space_after = Pt(4)
    return p


def bullets(d, items, style='List Bullet'):
    for it in items:
        p = d.add_paragraph(style=style)
        if isinstance(it, tuple):
            lead, rest = it
            r = p.add_run(lead)
            r.bold = True
            p.add_run(rest)
        else:
            p.add_run(it)
        p.paragraph_format.space_after = Pt(2)


def numbered(d, items):
    bullets(d, items, style='List Number')


def _shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_fill)
    tcPr.append(shd)


def table(d, header, rows, col_widths=None, font_size=9, first_col_bold=True, align_right_from=1, zebra=True):
    t = d.add_table(rows=1, cols=len(header))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for i, h in enumerate(header):
        hdr[i].text = ''
        p = hdr[i].paragraphs[0]
        r = p.add_run(str(h))
        r.bold = True
        r.font.size = Pt(font_size)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        _shade(hdr[i], '1F3A5F')
        if i >= align_right_from:
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for k, row in enumerate(rows):
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ''
            p = cells[i].paragraphs[0]
            txt = '' if v is None else str(v)
            r = p.add_run(txt)
            r.font.size = Pt(font_size)
            if i == 0 and first_col_bold:
                r.bold = True
            if i >= align_right_from:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            if zebra and k % 2 == 1:
                _shade(cells[i], 'F2F5F9')
    if col_widths:
        # python-docx needs BOTH: autofit off + the width on every cell of the column
        # (Word reads the cell widths, LibreOffice the grid) - otherwise the first column collapses.
        t.autofit = False
        for i, w in enumerate(col_widths):
            t.columns[i].width = Cm(w)
        for row in t.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Cm(w)
    d.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def kv_table(d, pairs, col_widths=(4.5, 12.5), font_size=9.5):
    """Two-column key/value block (trade ticket)."""
    t = d.add_table(rows=0, cols=2)
    t.style = 'Table Grid'
    for k, v in pairs:
        cells = t.add_row().cells
        cells[0].text = ''
        r = cells[0].paragraphs[0].add_run(k)
        r.bold = True
        r.font.size = Pt(font_size)
        _shade(cells[0], 'E8EEF5')
        cells[1].text = ''
        r = cells[1].paragraphs[0].add_run(v)
        r.font.size = Pt(font_size)
    for row in t.rows:
        row.cells[0].width = Cm(col_widths[0])
        row.cells[1].width = Cm(col_widths[1])
    d.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def fmt(x, nd=1, pct=False, sign=False):
    if x is None:
        return 'n/a'
    try:
        v = float(x)
    except Exception:
        return str(x)
    if v != v:
        return 'n/a'
    s = f"{v:+.{nd}f}" if sign else f"{v:,.{nd}f}"
    return s + ('%' if pct else '')
