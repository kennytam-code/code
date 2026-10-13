"""Offline regression checks: no Bloomberg connection or market-data download."""
import ast
import contextlib
import io
import json
import os
from pathlib import Path
import runpy
import tempfile
import unittest
from unittest.mock import patch

os.environ.setdefault('MPLBACKEND', 'Agg')
import numpy as np
import pandas as pd
import matplotlib.pyplot as plt

ROOT = Path(__file__).resolve().parents[1]
SCRIPT = ROOT / 'PCA_fixed.py'
if not SCRIPT.exists():
    SCRIPT = ROOT / 'pca_hedge_v103_full.py'
engine = runpy.run_path(str(SCRIPT), init_globals={'HEDGE_UI_AUTOSTART': False})

def data(n=420, seed=11):
    rng = np.random.default_rng(seed)
    dates = pd.bdate_range('2023-01-02', periods=n)
    market, sector = rng.normal(0, .01, (2, n))
    target = .75 * market + .25 * sector + rng.normal(0, .001, n)
    returns = np.column_stack([
        target, target + rng.normal(0, .002, n),
        .8 * market + rng.normal(0, .003, n),
        .6 * market + .3 * sector + rng.normal(0, .004, n),
        sector + rng.normal(0, .004, n), rng.normal(0, .01, n)])
    columns = ['2889 TT Equity', '2910 TT Equity', '0050 TT Equity',
               '006208 TT Equity', '2911 TT Equity', '2912 TT Equity']
    return {'target': columns[0],
            'prices': pd.DataFrame(100 * np.exp(np.cumsum(returns, axis=0)),
                                   index=dates, columns=columns),
            'mktcap_usd_mm': {}, 'price_observations': 'ACTIVE_DAYS_ONLY'}

OPTIONS = dict(lookback_options=(126,), pc_options=(2,),
               basket_size_options=(1, 2, 3), cv_test_window=21, cv_n_folds=3,
               outer_window=21, outer_folds=2, min_mktcap_usd_mm=0,
               restrict_to_target_tz=True, async_adjust=False,
               rolling_oos=False, boot_n=0, show_plots=False)

def run(dataset=None, **options):
    with contextlib.redirect_stdout(io.StringIO()):
        return engine['find_best_hedge']('2889 TT Equity', data=dataset or data(),
                                        **dict(OPTIONS, **options))

class HedgeSearchTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.dataset = data()
        cls.result = run(cls.dataset)
        cls.report = engine['transcript_text']()

    def tearDown(self):
        plt.close('all')

    def test_every_eligible_single_is_evaluated_alongside_basket_methods(self):
        specs = set(self.result['cv_detail']['pc_spec'])
        expected = {'single:' + c for c in self.dataset['prices'].columns[1:]}
        self.assertTrue(expected <= specs)
        self.assertTrue({'correlation', 'forward', '2'} <= specs)
        self.assertEqual(self.result['search']['single_instruments_scored'], 5)

    def test_stock_etf_and_basket_have_same_unseen_periods(self):
        for family in ('basket', 'single', 'single_stock', 'single_etf'):
            item = self.result['hedge_alternatives'][family]
            self.assertEqual(item['n_unseen_windows'], 2)
            self.assertEqual(item['trade_risk']['n_observations'], 42)
            self.assertTrue(np.isfinite(item['unseen_r2']))
        stock = self.result['best_single_stock_hedge']['weights']
        etf = self.result['best_single_etf_hedge']['weights']
        self.assertEqual(len(stock), 1)
        self.assertFalse(engine['_is_etf'](stock.index[0]))
        self.assertEqual(len(etf), 1)
        self.assertTrue(engine['_is_etf'](etf.index[0]))
        self.assertGreaterEqual(len(self.result['best_basket_hedge']['weights']), 2)

    def test_comparison_is_in_existing_report(self):
        text = self.report
        self.assertIn('BEST BASKET AND SINGLE-INSTRUMENT HEDGES', text)
        self.assertIn('BEST SINGLE', text)
        self.assertIn('STOCK', text)
        self.assertIn('ETF', text)

    def test_forward_selection_can_find_complementary_exposures(self):
        rng = np.random.default_rng(31)
        common, factor = rng.normal(0, 1, (2, 3000))
        frame = pd.DataFrame({'y': factor, 'a': 5*common + factor,
                              'b': 5*common - factor, 'noise': rng.normal(0, 1, 3000)})
        chosen = engine['_forward_order'](frame, 'y', ['a', 'b', 'noise'], max_size=2)
        self.assertEqual(set(chosen), {'a', 'b'})
        _, raw, _ = engine['_fit_basket'](frame, 'y', chosen)
        self.assertGreater(raw, .99)

    def test_failed_tracking_test_cannot_return_trade_weights(self):
        sample = data()
        prices = sample['prices']
        returns = np.log(prices / prices.shift()).fillna(0)
        returns.iloc[-42:, 0] *= -1
        sample['prices'] = 100 * np.exp(returns.cumsum())
        result = run(sample)
        self.assertLess(result['nested_pooled_r2'], 0)
        self.assertFalse(result['trade_eligible'])
        self.assertTrue(result['trade_weights'].empty)
        self.assertFalse(result['best_single_hedge']['qualified'])

    def test_gross_cap_is_never_silently_ignored(self):
        with self.assertRaisesRegex(RuntimeError, 'All CV folds failed'):
            run(max_gross=1e-12)

    def test_nested_score_and_risk_equal_the_actual_unseen_errors(self):
        unseen = self.result['nested_unseen_returns']
        self.assertTrue(unseen.index.is_monotonic_increasing)
        self.assertTrue(unseen.index.is_unique)
        y, residual, spread = (unseen[c].to_numpy() for c in unseen.columns)
        expected = 1 - np.sum(residual**2) / np.sum((y-y.mean())**2)
        self.assertAlmostEqual(self.result['nested_pooled_r2'], expected)
        risk = self.result['trade_risk']
        self.assertAlmostEqual(risk['vol_reduction'], 1 - spread.std()/y.std())
        self.assertAlmostEqual(risk['variance_reduction'], 1 - spread.var()/y.var())

    def test_bootstrap_interval_resamples_its_own_point_statistic(self):
        result = run(boot_n=150)
        interval = result['bootstrap_ci']
        self.assertEqual(interval['source'], 'nested unseen selection procedure')
        self.assertTrue(interval['conditional_on_observed_fits'])
        self.assertAlmostEqual(interval['r2_point'], result['nested_pooled_r2'])
        selected = result['selected_config_bootstrap_ci']
        self.assertAlmostEqual(selected['r2_point'], result['selected_unseen_pooled_r2'])
        self.assertTrue(np.isfinite([interval['r2_lo'], interval['r2_hi']]).all())

    def test_future_test_returns_do_not_change_prior_selection_or_predictions(self):
        changed = data()
        returns = np.log(changed['prices'] / changed['prices'].shift()).fillna(0)
        returns.iloc[-21:, 0] = -.5 * returns.iloc[-21:, 0] + .015
        changed['prices'] = 100 * np.exp(returns.cumsum())
        altered = run(changed)
        columns = ['window', 'lookback', 'pc_spec', 'basket_size']
        pd.testing.assert_frame_equal(self.result['nested_validation'][columns],
                                      altered['nested_validation'][columns])
        first = self.result['nested_unseen_returns']
        second = altered['nested_unseen_returns']
        np.testing.assert_allclose(first['target'] - first['regression_residual'],
                                   second['target'] - second['regression_residual'], atol=1e-12)
        self.assertNotAlmostEqual(self.result['nested_pooled_r2'], altered['nested_pooled_r2'])

    def test_later_missing_quotes_do_not_remove_an_earlier_eligible_stock(self):
        baseline = data()
        baseline['prices'].iloc[:60, 1] = np.nan
        changed = {**baseline, 'prices': baseline['prices'].copy()}
        changed['prices'].iloc[-21:, 1] = np.nan
        first, second = run(baseline), run(changed)
        oldest = first['nested_validation'].iloc[-1]
        second_nested = (second['prior_nested_validation'] if second.get('fallback_reason')
                         else second['nested_validation'])
        changed_oldest = second_nested.iloc[-1]
        for key in ('lookback', 'pc_spec', 'basket_size'):
            self.assertEqual(oldest[key], changed_oldest[key])
        self.assertAlmostEqual(oldest['r2_oos'], changed_oldest['r2_oos'])
        self.assertIn('2910 TT Equity', second['prices'].columns)

    def test_predictive_r2_is_not_given_a_fisher_correlation_interval(self):
        detail = self.result['outer_window_detail'].copy()
        detail['r2'] = [-.1, .7]
        stats = engine['_window_dispersion_stats'](detail, .5, 42)
        self.assertIsNone(stats['band'])
        self.assertEqual(stats['flags'], ['negative', ''])

    def test_singles_are_searched_even_with_basket_only_size_settings(self):
        result = run(basket_size_options=(2, 3))
        self.assertEqual(result['search']['single_instruments_scored'], 5)
        self.assertIsNotNone(result['best_single_hedge'])

    def test_quick_rehedge_preserves_the_primary_basket_configuration(self):
        result = self.result
        self.assertEqual(result['primary_hedge_family'], 'basket')
        with contextlib.redirect_stdout(io.StringIO()):
            quick = engine['quick_rehedge']('2889 TT Equity', self.dataset, result,
                                           show_plots=False, rolling_oos=False)
        self.assertEqual(result['params']['pc_spec'], quick['params']['pc_spec'])
        self.assertEqual(result['weights'].index.tolist(), quick['weights'].index.tolist())
        self.assertFalse(quick['params']['broad_search'])

    def test_explicitly_pinned_single_reference_can_still_be_refitted(self):
        result = run(basket_size_options=(1,), broad_search=False,
                     pc_options=('single:2910 TT Equity',))
        with contextlib.redirect_stdout(io.StringIO()):
            quick = engine['quick_rehedge']('2889 TT Equity', self.dataset, result,
                                           show_plots=False, rolling_oos=False)
        self.assertEqual(result['params']['pc_spec'], quick['params']['pc_spec'])
        self.assertEqual(result['weights'].index.tolist(), quick['weights'].index.tolist())
        self.assertFalse(quick['params']['broad_search'])

    def test_monthly_candidates_are_indicative_and_include_both_forms(self):
        result = run(sample_window='1m')
        self.assertEqual(result['validation_mode'], 'IN-SAMPLE')
        self.assertEqual(result['primary_hedge_family'], 'basket')
        pd.testing.assert_series_equal(result['weights'], result['best_basket_hedge']['weights'])
        for family in ('basket', 'single_stock', 'single_etf'):
            item = result['hedge_alternatives'][family]
            self.assertIn('INDICATIVE', item['status'])
            self.assertTrue(np.isnan(item['unseen_r2']))
            self.assertFalse(item['qualified'])
            self.assertEqual(item['reference_only'], family.startswith('single'))

    def test_short_history_with_many_names_caps_requested_pca_dimensions(self):
        sample = data()
        rng = np.random.default_rng(123)
        for i in range(30):
            sample['prices'][f'{3100+i} TT Equity'] = 100 * np.exp(np.cumsum(rng.normal(0, .01, 420)))
        result = run(sample, sample_window='1m', pc_options=(30,))
        self.assertEqual(result['validation_mode'], 'IN-SAMPLE')
        self.assertTrue(np.isfinite(result['in_sample_r2']))
        self.assertIsNotNone(result['best_basket_hedge'])

    def test_reported_flat_prices_are_retained(self):
        frame = pd.DataFrame({'a': [100.] * 20})
        pd.testing.assert_frame_equal(engine['_mask_stale'](frame, 7, frame), frame)

    def test_unsupported_matplotlib_setting_is_removed_and_charts_draw(self):
        rc = engine['_supported_plot_style']({'font.size': 9, 'axes.titlecolor': 'navy'}, {'font.size': 10})
        self.assertNotIn('axes.titlecolor', rc)
        with plt.rc_context(rc):
            figure, axis = plt.subplots()
            axis.set_title('Hedge comparison', color='navy')
            figure.canvas.draw()

    def test_cache_roundtrip_and_local_export_paths(self):
        with tempfile.TemporaryDirectory() as folder:
            with contextlib.redirect_stdout(io.StringIO()):
                path = engine['save_hedge_data'](data(), cache_dir=folder)
                restored = engine['load_hedge_data'](path)
            pd.testing.assert_frame_equal(restored['prices'], data()['prices'])
            export_folder, warning = engine['_export_folder'](str(Path(folder) / 'reports'))
            self.assertTrue(Path(export_folder).is_dir())
            self.assertIsNone(warning)
        self.assertNotIn('G:', engine['HEDGE_EXPORT_DIR'])

    def test_missing_bloomberg_api_gives_actionable_download_error(self):
        namespace = engine['bbg_connect'].__globals__
        previous = namespace['blpapi']
        try:
            namespace['blpapi'] = None
            with self.assertRaisesRegex(RuntimeError, 'Install the Bloomberg Python API'):
                engine['bbg_connect']()
        finally:
            namespace['blpapi'] = previous

    def test_existing_form_and_word_export_work_without_bloomberg(self):
        from docx import Document
        from IPython.utils.capture import capture_output
        with capture_output():
            state = engine['hedge_ui'](data=self.dataset)
        self.assertIs(state['data'], self.dataset)
        with tempfile.TemporaryDirectory() as folder:
            with contextlib.redirect_stdout(io.StringIO()):
                path = engine['export_result'](self.result, report_text='Offline synthetic validation.', out_dir=folder)
            self.assertEqual(Path(path).suffix, '.docx')
            document = Document(path)
            text = ('\n'.join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
                    + '\n' + '\n'.join(p.text for p in document.paragraphs))
            self.assertIn('Basket', text)
            self.assertIn('Best single stock', text)
            self.assertIn('Best single ETF', text)
            alternatives = pd.read_csv(Path(path).with_suffix('').as_posix() + '_alternatives.csv')
            self.assertEqual(set(alternatives['candidate_type']), {'basket', 'single_stock', 'single_etf', 'single_market'})
            self.assertTrue(alternatives.loc[alternatives['candidate_type'].str.startswith('single'),
                                             'reference_only'].all())
            self.assertFalse(alternatives.loc[alternatives['candidate_type'].eq('basket'),
                                              'reference_only'].any())
            self.assertIn('single instruments are references', '\n'.join(p.text for p in document.paragraphs))

    def test_history_comparison_contains_both_forms_for_each_window(self):
        with contextlib.redirect_stdout(io.StringIO()):
            comparison = engine['compare_history_windows']('2889 TT Equity', self.dataset,
                windows=('1y', '6m', '3m', '1m'), **OPTIONS)
        table = comparison['alternatives_summary']
        self.assertEqual(set(table['period']), {'1y', '6m', '3m', '1m'})
        for period, rows in table.groupby('period'):
            self.assertEqual(set(rows['candidate_type']), {'basket', 'single_stock', 'single_etf', 'single_market'})
            self.assertTrue(rows.loc[rows['candidate_type'].str.startswith('single'), 'reference_only'].all())
            self.assertGreaterEqual(len(comparison['diagnostic_results'][period]['weights']), 2)
        rejected = table[~table['qualified'] & ~table['status'].str.contains('INDICATIVE')]
        self.assertTrue(rejected['weights_per_100'].eq('').all())

    def test_batched_single_fits_match_ridgecv_alphas_coefficients_and_errors(self):
        rng = np.random.default_rng(33)
        for n, p in ((41, 20), (94, 9)):
            X = rng.normal(0, .01, (n, p)) + rng.normal(0, .005, (1, p))
            y = .003 + .7*X[:, 0] - .2*X[:, 1] + rng.normal(0, .002, n)
            frame = pd.DataFrame(X, columns=[str(i) for i in range(p)])
            frame['y'] = y
            fits = engine['_fit_single_candidates'](frame, 'y', list(frame.columns[:-1]))
            for name, (fast, score, residual) in fits.items():
                original, expected, errors = engine['_fit_basket'](frame, 'y', [name])
                self.assertEqual(fast.alpha_, original.alpha_)
                np.testing.assert_allclose(fast.coef_, original.coef_, atol=1e-10)
                self.assertAlmostEqual(fast.intercept_, original.intercept_, places=12)
                self.assertAlmostEqual(score, expected, places=12)
                np.testing.assert_allclose(residual, errors, atol=1e-12)

    def test_thin_pca_matches_sample_and_ledoit_wolf_covariance(self):
        from sklearn.covariance import LedoitWolf
        from sklearn.decomposition import PCA
        rng = np.random.default_rng(22)
        for n, p in ((41, 90), (94, 25)):
            X = rng.normal(size=(n, p)) + rng.normal(size=(n, 1))*np.arange(p)[None, :]/p
            for method in ('none', 'ledoit'):
                fast = engine['_pca_fit'](X, 3, method)
                if method == 'ledoit':
                    covariance = LedoitWolf().fit(X).covariance_
                    ev, vectors = np.linalg.eigh(covariance)
                    ids = np.argsort(ev)[::-1][:3]
                    expected = ev[ids]
                    projection = vectors[:, ids] @ vectors[:, ids].T
                    total = np.trace(covariance)
                else:
                    original = PCA(n_components=3, svd_solver='full').fit(X)
                    expected = original.explained_variance_
                    projection = original.components_.T @ original.components_
                    total = np.sum(np.var(X, axis=0, ddof=1))
                np.testing.assert_allclose(fast.explained_variance_, expected, atol=1e-12)
                np.testing.assert_allclose(fast.explained_variance_ratio_, expected/total, atol=1e-12)
                np.testing.assert_allclose(fast.components_.T @ fast.components_, projection, atol=1e-10)

    def test_two_missing_training_returns_do_not_discard_46_row_window(self):
        frame = pd.DataFrame({'target': np.arange(46, dtype=float),
                              'hedge': np.arange(46, dtype=float)*2},
                             index=pd.bdate_range('2026-04-01', periods=46))
        frame.iloc[[4, 14], 0] = np.nan
        clean, problem, omitted = engine['_training_rows'](frame, 'target', 42)
        self.assertIsNone(problem)
        self.assertEqual(len(clean), 44)
        self.assertEqual(omitted, list(frame.index[[4, 14]]))
        np.testing.assert_allclose(clean['hedge'], frame.loc[clean.index, 'hedge'])
        frame.iloc[:6, 0] = np.nan
        self.assertIn('coverage', engine['_training_rows'](frame, 'target', 42)[1])
        frame['target'] = 0.
        self.assertIn('constant', engine['_training_rows'](frame, 'target', 42)[1])

    def test_isolated_training_quote_gaps_keep_full_nested_validation(self):
        sample = data()
        sample['prices'].iloc[[210, 230], 0] = np.nan
        result = run(sample)
        self.assertEqual(len(result['nested_validation']), 2)
        audit = result['training_data_diagnostics']
        self.assertGreater(len(audit), 0)
        self.assertTrue(audit['accepted'].all())
        self.assertTrue(np.isfinite(result['in_sample_r2']))
        self.assertTrue(np.isfinite(result['nested_pooled_r2']))

    def test_missing_test_target_dates_are_common_and_not_filled_with_zero(self):
        sample = data()
        dates = sample['prices'].index
        sample['prices'].iloc[-10, 0] = np.nan
        result = run(sample)
        unseen = result['nested_unseen_returns']
        self.assertNotIn(dates[-10], unseen.index)
        self.assertNotIn(dates[-9], unseen.index)
        self.assertEqual(len(unseen), 40)
        for family in ('basket', 'single_stock', 'single_etf'):
            self.assertEqual(result['hedge_alternatives'][family]['trade_risk']['n_observations'], 40)
        audit = result['evaluation_data_diagnostics']
        self.assertTrue(audit['accepted'].all())
        self.assertTrue((audit['retained_rows'] >= .9*audit['calendar_rows']).all())
        prior = self.result['nested_unseen_returns'].loc[unseen.index]
        np.testing.assert_allclose(unseen['target']-unseen['regression_residual'],
                                   prior['target']-prior['regression_residual'], atol=1e-12)

    def test_missing_hedge_test_returns_cannot_be_removed_to_save_candidate(self):
        sample = data()
        sample['prices'].iloc[-10, 1] = np.nan
        result = run(sample)
        self.assertNotIn('single:2910 TT Equity', set(result['cv_summary']['pc_spec']))
        partial = result['cv_detail'].query("pc_spec == 'single:2910 TT Equity'")
        self.assertEqual(set(partial['fold']), {1, 2})
        self.assertFalse(result['trade_eligible'])
        self.assertTrue(result['trade_weights'].empty)
        self.assertGreaterEqual(len(result['weights']), 2)
        self.assertIn('unvalidated reference', result['fallback_reason'])
        self.assertTrue(np.isnan(result['ticket']['quality_r2']))
        self.assertFalse(result['best_basket_hedge']['qualified'])
        with tempfile.TemporaryDirectory() as folder:
            with contextlib.redirect_stdout(io.StringIO()):
                path = engine['export_result'](result, report_text='Unvalidated basket regression case.', out_dir=folder)
            from docx import Document
            text = '\n'.join(p.text for p in Document(path).paragraphs)
            self.assertIn('IN-SAMPLE ONLY', text)
            self.assertIn('unvalidated reference', text)
        with self.assertRaisesRegex(RuntimeError, 'All CV folds failed'):
            run(sample, broad_search=False, pc_options=('single:2910 TT Equity',),
                basket_size_options=(1,))

    def test_empty_nested_evidence_has_schema_and_does_not_crash_report(self):
        sample = data(n=264)
        sample['prices'].index = pd.bdate_range(end='2026-10-12', periods=264)
        sample['prices'].loc[pd.to_datetime(['2026-05-11', '2026-05-25']), sample['target']] = np.nan
        result = run(sample, sample_window='1y', lookback_options=(126, 252),
                     cv_test_window=42, outer_window=42, boot_n=150)
        self.assertIn('window', result['nested_validation'])
        self.assertTrue(result['nested_validation'].empty)
        self.assertIsNone(result['bootstrap_ci'])
        self.assertIsNotNone(result['selected_config_bootstrap_ci'])
        audit = result['evaluation_data_diagnostics']
        rejected = audit[~audit['accepted']]
        self.assertGreater(len(rejected), 0)
        self.assertTrue((rejected['retained_rows'] < .9*rejected['calendar_rows']).all())
        self.assertFalse(result['trade_eligible'])
        self.assertIn('INSUFFICIENT', result['ticket']['grade'])
        self.assertTrue(np.isfinite(result['in_sample_r2']))

    def test_rejected_report_never_recommends_full_size_or_claims_generalization(self):
        sample = data()
        returns = np.log(sample['prices']/sample['prices'].shift()).fillna(0)
        returns.iloc[-42:, 0] *= -1
        sample['prices'] = 100*np.exp(returns.cumsum())
        result = run(sample, rolling_oos=True, hedge_horizon_days=21)
        report = engine['transcript_text']()
        self.assertFalse(result['trade_eligible'])
        self.assertNotIn('trade it at full size', report)
        self.assertNotIn('✓ generalizes', report)
        self.assertNotIn('weighting three years ago', report)
        self.assertIn('CANDIDATE RECIPE — NOT APPROVED', report)
        self.assertEqual(result['ticket']['worst_window_r2'],
                         result['nested_validation']['r2_oos'].min())
        self.assertEqual(result['ticket']['worst_window_source'], 'nested selection procedure')

    def test_reported_universe_size_does_not_run_pca_for_each_single(self):
        rng = np.random.default_rng(824)
        n, p = 42, 1132
        y = rng.normal(0, .01, n)
        others = .5*y[:, None] + rng.normal(0, .01, (n, p))
        columns = ['2889 TT Equity', '0050 TT Equity'] + [f'{10000+i} TT Equity' for i in range(p-1)]
        prices = pd.DataFrame(100*np.exp(np.cumsum(np.column_stack([y, others]), axis=0)),
                              index=pd.bdate_range('2026-08-10', periods=n), columns=columns)
        sample = {'target': columns[0], 'prices': prices,
                  'mktcap_usd_mm': {}, 'price_observations': 'ACTIVE_DAYS_ONLY'}
        namespace = engine['find_best_hedge'].__globals__
        with patch.dict(namespace, {'_pca_fit': namespace['_pca_fit']}):
            original = namespace['_pca_fit']
            calls = []
            def counted(*a, **k):
                calls.append(1)
                return original(*a, **k)
            namespace['_pca_fit'] = counted
            result = run(sample, pc_options='auto', basket_size_options=(1, 2, 3, 5),
                         robust_cov='ledoit', winsorize_pct=.01)
        self.assertEqual(result['search']['single_instruments_scored'], 1132)
        self.assertEqual(result['validation_mode'], 'IN-SAMPLE')
        self.assertEqual(len(calls), 1)
        self.assertIsNotNone(result['best_basket_hedge'])
        self.assertIsNotNone(result['best_single_etf_hedge'])

    def test_engine_charts_draw_for_explicitly_pinned_single_reference(self):
        with patch.object(plt, 'show'):
            result = run(self.dataset, basket_size_options=(1,), show_plots=True,
                         broad_search=False, pc_options=('single:2910 TT Equity',))
        self.assertEqual(len(result['weights']), 1)
        self.assertEqual(result['params']['pc_spec'][:7], 'single:')
        for figure_number in plt.get_fignums():
            plt.figure(figure_number).canvas.draw()

    def test_stronger_single_reference_cannot_override_primary_basket(self):
        result = run(gross_penalty=1., basket_size_options=(1, 5))
        summary = result['cv_summary']
        basket = summary.loc[summary['family'].eq('basket')]
        single = summary.loc[summary['family'].str.startswith('single')]
        self.assertGreater(single['composite'].max(), basket['composite'].max())
        self.assertEqual(result['primary_hedge_family'], 'basket')
        self.assertGreaterEqual(len(result['weights']), 2)
        pd.testing.assert_series_equal(result['weights'], result['best_basket_hedge']['weights'])
        self.assertFalse(result['nested_validation']['pc_spec'].str.startswith('single:').any())
        for weights in result['nested_validation']['weights']:
            self.assertGreaterEqual(len(weights), 2)
        for family in ('single', 'single_stock', 'single_etf'):
            self.assertTrue(result['hedge_alternatives'][family]['reference_only'])
        self.assertEqual(result['trade_eligible'], result['best_basket_hedge']['qualified'])

    def test_holdout_selects_basket_without_using_holdout_target_returns(self):
        sample = data(n=76)
        first = run(sample)
        self.assertEqual(first['validation_mode'], 'HOLDOUT')
        self.assertGreaterEqual(len(first['weights']), 2)
        returns = np.log(sample['prices'] / sample['prices'].shift()).fillna(0)
        returns.iloc[-15:, 0] = -returns.iloc[-15:, 0] + .005
        changed = {**sample, 'prices': 100*np.exp(returns.cumsum())}
        second = run(changed)
        self.assertEqual(first['params']['pc_spec'], second['params']['pc_spec'])
        pd.testing.assert_series_equal(first['weights'], second['weights'])
        self.assertEqual(second['trade_eligible'], second['best_basket_hedge']['qualified'])
        self.assertTrue(second['best_single_hedge']['reference_only'])

    def test_no_basket_means_no_implicit_single_fallback(self):
        with self.assertRaisesRegex(ValueError, 'primary hedge is a basket'):
            run(basket_size_options=(1,))
        sample = data()
        sample['prices'] = sample['prices'].iloc[:, :2]
        with self.assertRaisesRegex(ValueError, 'At least two eligible instruments'):
            run(sample)

    def test_required_instrument_cannot_be_ignored_or_dropped(self):
        with self.assertRaisesRegex(ValueError, 'required instruments cannot be ignored'):
            run(force_include=['9999 TT Equity'])
        sample = data()
        sample['prices'].iloc[-1, 3] = np.nan
        with self.assertRaisesRegex(RuntimeError, 'All CV folds failed'):
            run(sample, force_include=['006208 TT Equity'])
        sample['prices'].iloc[:, 3] = np.nan
        with self.assertRaisesRegex(ValueError, 'required instruments cannot be dropped'):
            run(sample, force_include=['006208 TT Equity'])

    def test_absent_etf_is_unavailable_instead_of_relabelled_stock(self):
        sample = data()
        sample['prices'] = sample['prices'].drop(columns=['0050 TT Equity', '006208 TT Equity'])
        result = run(sample)
        self.assertIsNone(result['best_etf'])
        self.assertIsNone(result['best_single_etf_hedge'])
        self.assertIsNotNone(result['best_single_stock_hedge'])
        self.assertIn('Best single ETF: unavailable', engine['transcript_text']())

    def test_single_reference_does_not_change_primary_basket_grade(self):
        for objective in ('variance', 'tracking'):
            grade, reasons = engine['_ticket_verdict'](
                .7, .5, .02, .9, False, .7, .8, False,
                vol_red=.4, worst_vol_red=.3, objective=objective,
                ref_vol_pca=.4, ref_vol_simple=.5)
            self.assertTrue(grade.startswith('✓'))
            self.assertIn('reference', ' '.join(reasons))
        self.assertNotIn('prefer the simple hedge', self.report)
        self.assertNotIn('take the ETF', self.report)
        self.assertIn('single instruments are references', self.report)

    def test_notebook_engine_matches_script_and_every_cell_parses(self):
        notebook_path = ROOT / 'PCA_fixed.ipynb'
        if not notebook_path.exists():
            notebook_path = ROOT / 'pca_hedge_v103_full.ipynb'
        notebook = json.loads(notebook_path.read_text(encoding='utf-8'))
        self.assertEqual(''.join(notebook['cells'][1]['source']), SCRIPT.read_text(encoding='utf-8'))
        for cell in notebook['cells']:
            if cell['cell_type'] == 'code':
                ast.parse(''.join(cell['source']))
                self.assertFalse(cell.get('outputs'))

    def test_peers_only_filters_every_fit_and_reference(self):
        sample = data()
        peers = ['2911 TT Equity', '2912 TT Equity']
        sample['peer_tickers_by_target'] = {'2889 TT Equity': peers}
        result = run(sample, candidate_scope='peers', basket_size_range=(2, 2))
        self.assertEqual(set(result['weights'].index), set(peers))
        self.assertEqual(set(result['prices'].columns), {'2889 TT Equity', *peers})
        for weights in result['nested_validation']['weights']:
            self.assertEqual(set(weights.index), set(peers))
        self.assertIsNone(result['best_single_etf_hedge'])
        self.assertIsNone(result['broad_market_hedge'])
        for item in result['hedge_alternatives'].values():
            self.assertTrue(set(item['weights'].index) <= set(peers))

    def test_peer_filter_fails_closed_for_old_cache_and_conflicting_forced_name(self):
        with self.assertRaisesRegex(ValueError, 'Peers-only needs verified peer metadata'):
            run(candidate_scope='peers')
        peers = ['2910 TT Equity', '2911 TT Equity']
        manual = run(candidate_scope='peers', peer_tickers=peers, basket_size_range=(2, 2))
        self.assertEqual(set(manual['weights'].index), set(peers))
        with self.assertRaisesRegex(ValueError, 'conflict with the candidate filter'):
            run(candidate_scope='peers', peer_tickers=peers, force_include=['2912 TT Equity'])
        sample = dict(self.dataset, peer_tickers_by_target={'2889 TT Equity': peers})
        with self.assertRaisesRegex(ValueError, 'not in the downloaded data'):
            run(sample, candidate_scope='peers', peer_tickers=['UNKNOWN TT Equity'])

    def test_local_market_exception_is_explicit_and_synchronous(self):
        sample = dict(self.dataset, peer_tickers_by_target={'2889 TT Equity': ['2910 TT Equity', '2911 TT Equity']})
        result = run(sample, candidate_scope='peers', include_broad_market=True, basket_size_range=(2, 3))
        allowed = {'2910 TT Equity', '2911 TT Equity', '0050 TT Equity', '006208 TT Equity'}
        self.assertEqual(set(result['prices'].columns) - {'2889 TT Equity'}, allowed)
        market = result['broad_market_hedge']
        self.assertTrue(set(market['weights'].index) <= {'0050 TT Equity', '006208 TT Equity'})
        self.assertTrue(market['reference_only'])
        self.assertEqual(market['trade_risk']['n_observations'], 42)
        self.assertFalse(result['two_day_returns'])

    def test_hk_and_adr_market_mapping_use_listing_venue(self):
        spec = engine['_broad_market_spec']
        self.assertEqual(spec('700 HK Equity')['instruments'], ['2800 HK Equity'])
        self.assertEqual(spec('BABA US Equity')['instruments'], ['SPY US Equity'])
        self.assertEqual(spec('7203 JT Equity')['instruments'], ['1306 JP Equity'])
        self.assertEqual(engine['_close_utc_of']('7203 JT Equity'), 6.5)
        self.assertIn('Taiwan 50', spec('2889 TT Equity')['benchmark'])
        self.assertNotIn('TAIEX', spec('2889 TT Equity')['benchmark'])
        self.assertEqual(spec('UNKNOWN ZZ Equity')['instruments'], [])
        for _, tickers in engine['BROAD_MARKET_HEDGES'].values():
            self.assertTrue(all(engine['_is_etf'](ticker) for ticker in tickers))

    def test_discovery_tags_direct_and_industry_peers_not_arbitrary_members(self):
        ns = engine['auto_discover_universe'].__globals__
        fields = {'COUNTRY_ISO': 'TW', 'GICS_SECTOR_NAME': 'Financials', 'GICS_INDUSTRY_NAME': 'Banks'}
        def members(session, spec):
            return 'TW50 Index', ['2910 TT Equity', '2911 TT Equity', '2912 TT Equity']
        def refs(session, tickers, ref_fields, **kw):
            return {ticker: {'GICS_INDUSTRY_NAME': 'Banks' if ticker == '2911 TT Equity' else 'Hardware'} for ticker in tickers}
        def bulk(session, ticker, field):
            return [{'Peer Ticker': '2910 TT'}] if ticker == '2889 TT Equity' else [{'Peer Ticker': '2912 TT'}]
        meta = {}
        with patch.dict(ns, bbg_ref_single=lambda *a: fields, _index_members=members,
                        bbg_ref_multi=refs, bbg_bulk=bulk), contextlib.redirect_stdout(io.StringIO()):
            universe = engine['auto_discover_universe'](object(), '2889 TT Equity', min_mktcap_usd_mm=0, meta=meta)
        self.assertIn('2912 TT Equity', universe)
        self.assertEqual(set(meta['peer_tickers']), {'2910 TT Equity', '2911 TT Equity'})
        self.assertEqual(meta['broad_market']['market'], 'TW')

    def test_exact_range_is_obeyed_in_cv_nested_and_final_refits(self):
        result = run(basket_size_range=(3, 3), force_include=['2911 TT Equity'])
        self.assertEqual(len(result['weights']), 3)
        baskets = result['cv_detail'].loc[result['cv_detail']['family'].eq('basket')]
        self.assertEqual(set(baskets['basket_size']), {3})
        for weights in result['nested_validation']['weights']:
            self.assertEqual(len(weights), 3)
            self.assertIn('2911 TT Equity', weights.index)
        for row in result['basket_comparison']:
            self.assertEqual(len(row['weights']), 3)
        quick = engine['quick_rehedge']
        with contextlib.redirect_stdout(io.StringIO()):
            refreshed = quick('2889 TT Equity', self.dataset, result, show_plots=False, rolling_oos=False)
        self.assertEqual(len(refreshed['weights']), 3)
        with self.assertRaisesRegex(ValueError, 'exceeds every requested basket size'):
            run(basket_size_range=(2, 2), force_include=['2910 TT Equity', '2911 TT Equity', '2912 TT Equity'])

    def test_unfillable_size_is_never_shrunk_and_short_forced_size_is_indicative(self):
        sample = dict(self.dataset, peer_tickers_by_target={'2889 TT Equity': ['2910 TT Equity', '2911 TT Equity']})
        with self.assertRaisesRegex(ValueError, 'eligible instruments'):
            run(sample, candidate_scope='peers', basket_size_range=(3, 3))
        result = run(sample_window='1m', basket_size_range=(3, 3))
        self.assertEqual(len(result['weights']), 3)
        self.assertEqual(result['validation_mode'], 'IN-SAMPLE')
        self.assertFalse(result['trade_eligible'])
        self.assertTrue(np.isnan(result['ticket']['quality_r2']))
        for row in result['basket_comparison']:
            self.assertEqual(len(row['weights']), 3)
        for invalid in ((1, 3), (4, 2), (2.5, 3), (2, 51)):
            with self.assertRaises(ValueError):
                run(basket_size_range=invalid)

    def test_all_five_windows_show_in_sample_r2_and_keep_notes_in_one_area(self):
        for period in ('1m', '2m', '3m', '6m', '1y'):
            result = run(sample_window=period)
            blocks = engine['_report_blocks'](result)
            headline = next(b for b in blocks if b.get('headers', []) ==
                            ['EVIDENCE', 'UNSEEN R²', 'VOLATILITY CUT', 'IN-SAMPLE R²'])
            self.assertEqual(headline['rows'][0][3], f"{result['in_sample_r2']:.3f}")
            html = engine['_report_html'](result)
            visible, notes = html.split('<details>', 1)
            self.assertIn('IN-SAMPLE R²', visible)
            self.assertIn('metric-grid', visible)
            self.assertEqual(html.count('<details>'), 1)
            self.assertNotIn('Alternative basket scores retain selection bias', visible)
            self.assertIn('Alternative basket scores retain selection bias', notes)
            self.assertIn('costs', engine['_BRIEF_METHODOLOGY'])

    def test_more_distinct_alternatives_do_not_change_the_selected_hedge(self):
        options = dict(lookback_options=(63, 126), pc_options=(2, 3, 4), basket_size_range=(2, 4))
        fewer = run(alternative_baskets=1, **options)
        more = run(alternative_baskets=8, **options)
        pd.testing.assert_series_equal(fewer['weights'], more['weights'])
        self.assertEqual(fewer['nested_pooled_r2'], more['nested_pooled_r2'])
        self.assertLessEqual(len(fewer['basket_comparison']), 2)
        self.assertGreater(len(more['basket_comparison']), len(fewer['basket_comparison']))
        fingerprints = [tuple(sorted((tk, round(float(x), 5)) for tk, x in r['weights'].items()))
                        for r in more['basket_comparison']]
        self.assertEqual(len(fingerprints), len(set(fingerprints)))
        self.assertEqual(sum(row['selected'] for row in more['basket_comparison']), 1)
        self.assertTrue(all(2 <= row['basket_size'] <= 4 for row in more['basket_comparison']))
        for bad in (0, 21, 2.5, True):
            with self.assertRaisesRegex(ValueError, 'alternative_baskets'):
                run(alternative_baskets=bad)

    def test_short_and_long_window_graphs_survive_html_and_word_export(self):
        import base64
        from docx import Document
        for window, count in [('1m', 2), ('1y', 3)]:
            result = run(sample_window=window, show_plots=True, rolling_oos=True)
            charts = result['report_charts']
            self.assertEqual(len(charts), count)
            self.assertTrue(all(base64.b64decode(c['png']).startswith(b'\x89PNG') for c in charts))
            html = engine['_report_html'](result)
            self.assertEqual(html.count('data:image/png;base64,'), count)
            self.assertLess(html.index('<figure'), html.index('<div class="metric-grid">'))
            blocks = engine['_report_blocks'](result)
            chart_positions = [i for i, block in enumerate(blocks) if block['kind'] == 'chart']
            first_table = next(i for i, block in enumerate(blocks) if block['kind'] == 'table')
            self.assertLess(max(chart_positions), first_table)
            with tempfile.TemporaryDirectory() as folder:
                path = Path(folder) / 'report.docx'
                engine['_docx_hedge_report'](result, result['report_text'], str(path))
                self.assertEqual(len(Document(path).inline_shapes), count)
            self.assertEqual(len(plt.get_fignums()), 0)

    def test_recipe_trade_directions_match_the_coefficient_only_spread(self):
        weights = pd.Series([.6, -.25], index=['2910 TT Equity', '0050 TT Equity'])
        result = dict(self.result, weights=weights)
        recipe = next(b for b in engine['_report_blocks'](result)
                      if b.get('headers') == ['ACTION', 'PER $100 TARGET', 'INSTRUMENT'])
        self.assertEqual(recipe['rows'], [['SHORT', '$60.0', '2910 TT Equity'],
                                         ['LONG', '$25.0', '0050 TT Equity']])
        frame = self.result['returns']
        actual = frame['2889 TT Equity'] - frame[weights.index] @ weights
        manual = frame['2889 TT Equity'] - .6*frame['2910 TT Equity'] + .25*frame['0050 TT Equity']
        np.testing.assert_allclose(actual, manual)
        self.assertFalse(engine['_decision_fields'](weights, '? IN-SAMPLE ONLY')['trade_eligible'])
        self.assertTrue(engine['_decision_fields'](weights, '✗ NO VALIDATED HEDGE')['trade_weights'].empty)

    def test_notebook_run_prints_results_after_the_styled_report(self):
        namespace = engine['find_best_hedge'].__globals__
        output = io.StringIO()
        displayed = []
        def display(html):
            displayed.append(str(html.data))
            self.assertEqual(output.getvalue(), '')
        with patch.dict(namespace, _IN_JUPYTER=True, _ipy_display=display):
            with contextlib.redirect_stdout(output):
                result = engine['find_best_hedge']('2889 TT Equity', data=self.dataset, **OPTIONS)
        text = output.getvalue()
        self.assertEqual(len(displayed), 1)
        self.assertIn('UNSEEN R²', text)
        self.assertIn('IN-SAMPLE R²', text)
        self.assertIn('PER $100 TARGET', text)
        self.assertIn('ALTERNATIVE BASKETS', text)
        self.assertIn('costs, borrow, financing and FX', ' '.join(text.split()))
        self.assertNotIn('Alternative basket scores retain selection bias', text)
        self.assertIn('Alternative basket scores retain selection bias', result['report_text'])
        pd.testing.assert_series_equal(result['weights'], self.result['weights'])
        self.assertEqual(result['nested_pooled_r2'], self.result['nested_pooled_r2'])

        output = io.StringIO()
        with patch.dict(namespace, _IN_JUPYTER=True, _ipy_display=lambda _: None):
            with contextlib.redirect_stdout(output):
                detailed = engine['find_best_hedge']('2889 TT Equity', data=self.dataset,
                                                    report='full', **OPTIONS)
        self.assertIn(detailed['report_text'], output.getvalue())
        pd.testing.assert_series_equal(detailed['weights'], self.result['weights'])

    def test_failed_notebook_display_still_prints_the_result(self):
        namespace = engine['find_best_hedge'].__globals__
        for notebook in (False, True):
            output = io.StringIO()
            with patch.dict(namespace, _IN_JUPYTER=notebook,
                            _ipy_display=lambda _: (_ for _ in ()).throw(RuntimeError('Display unavailable'))):
                with contextlib.redirect_stdout(output):
                    result = engine['find_best_hedge']('2889 TT Equity', data=self.dataset, **OPTIONS)
            self.assertIn('UNSEEN R²', output.getvalue())
            self.assertIn('PER $100 TARGET', output.getvalue())
            self.assertIn(f"{result['in_sample_r2']:.3f}", output.getvalue())

    def test_two_day_cumulative_chart_does_not_double_count_overlap(self):
        curves = []
        def capture():
            figure = plt.gcf()
            if len(figure.axes) == 6:
                curves.append([np.asarray(line.get_ydata(), dtype=float).copy()
                               for line in figure.axes[3].lines[:2]])
        try:
            with patch.object(plt, 'show', side_effect=capture):
                result = run(async_adjust=True, show_plots=True)
            self.assertTrue(result['two_day_returns'])
            self.assertEqual(len(curves), 1)
            frame = result['returns']
            target = frame['2889 TT Equity']
            spread = target - frame[result['weights'].index] @ result['weights']
            np.testing.assert_allclose(curves[0][0], np.expm1(target.iloc[::2].cumsum())*100)
            np.testing.assert_allclose(curves[0][1], np.expm1(spread.iloc[::2].cumsum())*100)
        finally:
            plt.close('all')

    def test_spread_path_uses_initial_capital_and_nonoverlapping_returns(self):
        dates = pd.bdate_range('2024-01-01', periods=8)
        series = pd.Series([-.02] * 8, index=dates)
        path = engine['_spread_path_stats'](series)
        self.assertAlmostEqual(path['max_drawdown'], np.expm1(-.16))
        self.assertIsNone(path['peak'])
        two_day = engine['_spread_path_stats'](series, two_day=True)
        self.assertEqual(two_day['n_observations'], 4)
        self.assertAlmostEqual(two_day['max_drawdown'], np.expm1(-.08))
        self.assertTrue(np.isnan(two_day['worst_5d']))
        short = run(sample_window='1m')
        self.assertIn('fitted', short['path']['source'])
        self.assertFalse(short['trade_eligible'])

    def test_rich_report_retains_evidence_and_distinguishes_tracking_from_traded_risk(self):
        result = run(rolling_oos=True, boot_n=100)
        frame = result['rolling_oos']
        self.assertIn('traded_spread', frame.columns)
        self.assertIn('spread_vol_roll', frame.columns)
        self.assertGreater(float((frame['traded_spread'] - frame['resid']).abs().max()), 0)
        expected = engine['_spread_path_stats'](frame['traded_spread'], result['two_day_returns'])
        self.assertAlmostEqual(result['path']['max_drawdown'], expected['max_drawdown'])
        text = ' '.join(result['report_text'].split()).lower()
        for useful in ('VALIDATION AND RESIDUAL RISK', 'Nested selection windows',
                       'STABILITY AND RECENT CONDITIONS', 'REMAINING EXPOSURES',
                       'SEARCH SETTINGS AND DATA QUALITY', 'Coefficient-only spread volatility',
                       'block-bootstrap interval', 'Name overlap'):
            self.assertIn(useful.lower(), text)
        self.assertIn('configuration was selected using the full analysis window', text)
        self.assertIn('not financed portfolio p&l', text)

    def test_compact_report_keeps_full_diagnostics_separate_and_exports_real_alternatives(self):
        from docx import Document
        result = self.result
        self.assertGreater(len(result['full_report']), 2 * len(result['report_text']))
        self.assertNotIn('WORDS USED IN THIS REPORT', result['report_text'])
        self.assertIn('Alternative basket scores retain', result['report_text'])
        self.assertGreater(len(result['basket_comparison']), 1)
        with tempfile.TemporaryDirectory() as folder, contextlib.redirect_stdout(io.StringIO()):
            path = engine['export_result'](result, out_dir=folder)
            document = Document(path)
            paragraphs = '\n'.join(p.text for p in document.paragraphs)
            self.assertIn('Alternative baskets', paragraphs)
            self.assertNotIn('Diagnostic report appendix', paragraphs)
            table_text = '\n'.join(cell.text for t in document.tables for row in t.rows for cell in row.cells)
            for row in result['basket_comparison']:
                self.assertTrue(all(t in table_text for t in row['weights'].index))

    def test_form_presets_and_optional_sizing_do_not_change_validation_preview(self):
        from IPython.utils.capture import capture_output
        with capture_output():
            state = engine['hedge_ui'](data=self.dataset)
        widgets = state['widgets']
        periods = {v for _, v in widgets['analysis_window'].options}
        self.assertTrue({'1m', '2m'} <= periods)
        self.assertEqual(widgets['hedge_horizon'].value, 0)
        self.assertEqual(widgets['notional'].value, 0)
        before = widgets['validation_preview'].value
        widgets['hedge_horizon'].value = 63
        self.assertEqual(widgets['validation_preview'].value, before)
        widgets['candidate_scope'].value = 'peers'
        self.assertFalse(widgets['include_broad_market'].value)
        self.assertIn('verified peer metadata', widgets['validation_preview'].value)

    def test_quiet_download_cache_retains_peer_metadata_and_material_data_warnings(self):
        sample = dict(self.dataset, peer_tickers_by_target={
            '2889 TT Equity': ['2910 TT Equity', '2911 TT Equity']})
        def downloaded(*args, **kwargs):
            print('routine discovery progress')
            print('⚠ historical classification data unavailable')
            return dict(sample)
        ns = engine['download_hedge_data'].__globals__
        summary = io.StringIO()
        with patch.dict(ns, _download_hedge_data_impl=downloaded), contextlib.redirect_stdout(summary):
            result = engine['download_hedge_data']('2889 TT Equity')
        self.assertNotIn('routine discovery progress', summary.getvalue())
        self.assertIn('historical classification data unavailable', summary.getvalue())
        self.assertIn('routine discovery progress', result['download_log'])
        with tempfile.TemporaryDirectory() as folder, contextlib.redirect_stdout(io.StringIO()):
            path = engine['save_hedge_data'](result, cache_dir=folder)
            cached = engine['load_hedge_data'](path)
        peer_result = run(cached, candidate_scope='peers', basket_size_range=(2, 2))
        self.assertEqual(set(peer_result['weights'].index), {'2910 TT Equity', '2911 TT Equity'})

if __name__ == '__main__':
    unittest.main(verbosity=2)
