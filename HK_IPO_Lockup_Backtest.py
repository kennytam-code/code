"""HK IPO lock-up backtest. Copy this ENTIRE file into ONE Jupyter cell and run.

Requires an entitled Bloomberg Desktop API workstation with Bloomberg logged in.
Ticker/date inputs and primary-source review are embedded. No companion files.
Gross short price returns only. No assumed fees or borrow costs.
Prints date checks, strategy results, coverage and charts; exports Excel/CSV/HTML.
The current Hong Kong daily bar is excluded. No substitute price provider.
Live Bloomberg transport/entitlements require verification on your workstation.
"""
from __future__ import annotations

# SETTINGS — defaults use your supplied workbook and Hong Kong trading sessions.
AS_OF = "2026-10-28"
DAY_BASIS = "sessions"           # One month = 21 HK trading sessions
T0_CONVENTION = "last_locked"    # Or "first_unlocked"
EXTENDED_COMBINATIONS = False    # True adds pre-expiry entry to T+7 / T+21 exits
REFRESH_BLOOMBERG = False        # True ignores the exact-request Bloomberg cache
SELECTED_TICKER = "3750 HK Equity"
WORKBOOK_PATH = None             # Optional path to an updated .xlsm; never executes macros
OUTPUT_ROOT = "lockup_backtest_results"

import sys, importlib.util, subprocess
_packages = {"pandas":"pandas>=2.2,<3", "numpy":"numpy>=1.26,<3",
             "matplotlib":"matplotlib>=3.8,<4", "openpyxl":"openpyxl>=3.1,<4"}
_missing = [spec for name,spec in _packages.items() if importlib.util.find_spec(name) is None]
if _missing:
    print("Installing missing packages in this Jupyter kernel:", _missing)
    subprocess.check_call([sys.executable, "-m", "pip", "install", *_missing])

from dataclasses import dataclass, asdict
from pathlib import Path
import hashlib
import html
import json
import re
import time
import warnings
from datetime import datetime, date

import numpy as np
import pandas as pd

VERSION = "1.1.0"
PRICE_FIELDS = ("PX_OPEN", "PX_LAST", "PX_VOLUME")


@dataclass(frozen=True)
class Settings:
    asof: str = "2026-10-28"
    day_basis: str = "sessions"       # sessions or calendar
    anchor: str = "last_locked"       # last_locked or first_unlocked
    borrow_rate: float = 0.0           # No assumed costs in the requested price study
    fee_bps_per_leg: float = 0.0
    borrow_day_count: int = 365
    intraday_borrow_days: int = 1
    require_locate: bool = False       # True requires dated shortability + borrow evidence
    require_complete_30: bool = True
    bootstrap_draws: int = 2000
    random_seed: int = 20261028

    def __post_init__(self):
        cutoff=pd.Timestamp(self.asof)
        if pd.isna(cutoff) or cutoff.tzinfo is not None or cutoff != cutoff.normalize():
            raise ValueError("asof must be a valid date without a timezone or time component")
        if self.day_basis not in {"sessions", "calendar"}:
            raise ValueError("day_basis must be sessions or calendar")
        if self.anchor not in {"last_locked", "first_unlocked"}:
            raise ValueError("anchor must be last_locked or first_unlocked")
        for k in ("borrow_rate", "fee_bps_per_leg"):
            if not np.isfinite(getattr(self, k)) or getattr(self, k) < 0:
                raise ValueError(f"{k} must be finite and nonnegative")
        if self.borrow_day_count <= 0 or self.intraday_borrow_days < 0:
            raise ValueError("Invalid borrow day-count convention")
        if self.bootstrap_draws < 100:
            raise ValueError("Use at least 100 bootstrap draws")


@dataclass(frozen=True)
class Strategy:
    key: str
    entry_offset: int
    exit_offset: int
    entry_field: str = "PX_LAST"
    exit_field: str = "PX_LAST"

    @property
    def label(self):
        def leg(n, field):
            return f"T{n:+d}".replace("T+0", "T0") + (" open" if field == "PX_OPEN" else " close")
        return f"{leg(self.entry_offset, self.entry_field)} → {leg(self.exit_offset, self.exit_field)}"

    def label_for(self, basis):
        return self.label if basis == "sessions" else self.label.replace("T-21", "T−1m").replace("T+21", "T+1m")


def make_strategies(extended=False):
    """19 requested/core combinations, including the explicitly invalid zero hold.

    T-1 close to T-1 close is retained in the audit and is NOT a zero-return trade.
    extended=True additionally holds the seven pre-expiry entries to T+7/T+21.
    """
    ans = [Strategy(f"pre_{n}_{x}", -n, x)
           for n in (21, 14, 10, 7, 3, 2, 1) for x in (-1, 0)]
    ans += [Strategy("expiry_intraday", 0, 0, "PX_OPEN")]
    ans += [Strategy(f"post_close_{n}", 0, n) for n in (7, 21)]
    ans += [Strategy(f"post_open_{n}", 0, n, "PX_OPEN") for n in (7, 21)]
    if extended:
        ans += [Strategy(f"cross_{n}_{x}", -n, x)
                for n in (21, 14, 10, 7, 3, 2, 1) for x in (7, 21)]
    return ans


def normalize_ticker(value):
    m = re.fullmatch(r"\s*0*(\d{1,5})\s+HK\s+EQUITY\s*", str(value), re.I)
    if not m or not 0 < int(m[1]) < 100000:
        raise ValueError(f"Not an explicit Hong Kong equity ticker: {value!r}")
    return f"{int(m[1])} HK Equity"


def parse_dates(value):
    """This workbook uses US m/d/yyyy. No locale inference or day/month guessing."""
    if isinstance(value, (datetime, date, pd.Timestamp)):
        return [pd.Timestamp(value).normalize()]
    if value is None or not str(value).strip():
        raise ValueError("Missing expiry")
    parts = re.split(r"[,;\n]+", str(value))
    result = []
    for token in parts:
        token = token.strip()
        if not token:
            continue
        fmt = "%Y-%m-%d" if re.fullmatch(r"\d{4}-\d{2}-\d{2}", token) else "%m/%d/%Y"
        result.append(pd.Timestamp(datetime.strptime(token, fmt)).normalize())
    if not result:
        raise ValueError("Missing expiry")
    return result


def read_workbook(path, sheet="Sheet1"):
    """Read cached cells only. No macros, formulas or workbook instructions execute."""
    import openpyxl
    rows, issues = [], []
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", UserWarning)
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True, keep_links=False)
        try:
            if sheet not in wb.sheetnames:
                raise ValueError(f"Missing sheet {sheet!r}")
            ws = wb[sheet]
            header = list(next(ws.iter_rows(min_row=1, max_row=1, max_col=3, values_only=True)))
            if header != ["Name", "Lock-Up / Quiet Info", "Ticker"]:
                raise ValueError(f"Unexpected input header: {header}")
            for row, (name, raw, ticker) in enumerate(ws.iter_rows(min_row=4, max_col=3, values_only=True), 4):
                if not ticker and not raw and not name:
                    continue
                try:
                    normalized = normalize_ticker(ticker)
                    days = parse_dates(raw)
                except (ValueError, TypeError) as exc:
                    issues.append(dict(source_row=row, raw_ticker=ticker, raw_dates=raw, reason=str(exc)))
                    continue
                for seq, d in enumerate(days, 1):
                    rows.append(dict(input_id=f"Sheet1:{row}:{seq}", source_row=row,
                                     name=str(name or ""), ticker=normalized,
                                     input_date=d.date().isoformat(), raw_dates=str(raw)))
        finally:
            wb.close()
    events = pd.DataFrame(rows, columns=["input_id", "source_row", "name", "ticker", "input_date", "raw_dates"])
    if events.empty:
        raise ValueError("No usable ticker/date rows")
    events["duplicate_input"] = events.duplicated(["ticker", "input_date"], keep="first")
    return events, pd.DataFrame(issues, columns=["source_row", "raw_ticker", "raw_dates", "reason"])


def true_mask(series):
    return series.astype(str).str.lower().isin(["true", "1", "yes"])


def select_events(inputs, review, settings):
    """Exact ticker + supplied-date join. Corrections require an explicit review row."""
    required = {"ticker", "input_date", "last_day_locked", "category", "confirmed",
                "amendment_blocked", "source_url", "publication_date", "review_note"}
    if not required <= set(review.columns):
        raise ValueError(f"Review missing columns: {sorted(required - set(review.columns))}")
    a, r = inputs.copy().reset_index(drop=True), review.copy().reset_index(drop=True)
    for frame in (a, r):
        frame["ticker"] = frame.ticker.map(normalize_ticker)
        frame["input_date"] = pd.to_datetime(frame.input_date, errors="raise").dt.strftime("%Y-%m-%d")
        if frame.input_date.isna().any():
            raise ValueError("Input and review keys must have nonmissing ticker/date values")
    if r.duplicated(["ticker", "input_date"]).any():
        raise ValueError("Ambiguous review: multiple decisions for the same ticker/date")
    audit = a.merge(r, on=["ticker", "input_date"], how="left", validate="many_to_one")
    audit["status"] = "Unverified date or holder scope"
    confirmed = true_mask(audit.confirmed) & audit.source_url.fillna("").str.startswith("https://")
    audit.loc[confirmed, "status"] = "Outside cornerstone scope"
    cornerstone = confirmed & audit.category.eq("Cornerstone")
    audit.loc[cornerstone, "status"] = "Eligible date"
    audit.loc[true_mask(audit.amendment_blocked), "status"] = "Amended restriction: excluded"
    audit["last_day_locked"] = pd.to_datetime(audit.last_day_locked, errors="coerce")
    audit["publication_date"] = pd.to_datetime(audit.publication_date, errors="coerce")
    audit.loc[cornerstone & audit.last_day_locked.isna(), "status"] = "Missing confirmed last day locked"
    audit.loc[cornerstone & audit.publication_date.isna(), "status"] = "Missing publication date"
    future = audit.last_day_locked.ge(pd.Timestamp(settings.asof)) | (
        audit.last_day_locked.isna() & pd.to_datetime(audit.input_date).ge(pd.Timestamp(settings.asof)))
    audit.loc[future, "status"] = "Future or current-day restriction"
    duplicates = a.duplicated(["ticker", "input_date"], keep="first")
    audit.loc[duplicates, "status"] = "Duplicate input ticker/date"
    accepted = audit[audit.status.eq("Eligible date")].copy()
    accepted["event_id"] = accepted.ticker + "|" + accepted.last_day_locked.dt.strftime("%Y-%m-%d")
    # Two vendor dates cannot give the same underlying cornerstone event two weights.
    duplicates = accepted.duplicated("event_id", keep="first")
    audit.loc[accepted.index[duplicates], "status"] = "Duplicate corrected event"
    accepted = accepted[~duplicates].copy()
    return audit, accepted


def validate_prices(prices):
    required = {"ticker", "date", *PRICE_FIELDS}
    if not required <= set(prices.columns):
        raise ValueError(f"Prices missing columns: {sorted(required - set(prices.columns))}")
    p = prices.copy()
    p["ticker"] = p.ticker.map(normalize_ticker)
    p["date"] = pd.to_datetime(p.date, errors="raise").dt.normalize()
    if p.date.isna().any() or p.duplicated(["ticker", "date"]).any():
        raise ValueError("Missing or duplicate ticker/date price keys")
    for field in PRICE_FIELDS:
        p[field] = pd.to_numeric(p[field], errors="raise")
        if np.isinf(p[field]).any():
            raise ValueError(f"Infinite {field}")
        if (p[field].dropna() < 0).any():
            raise ValueError(f"Negative {field}")
    # Zero and NaN prices stay missing. Repeated real closes are NOT presumed stale.
    for field in ("PX_OPEN", "PX_LAST"):
        p.loc[p[field].eq(0), field] = np.nan
    return p.sort_values(["ticker", "date"]).reset_index(drop=True)


def validate_sessions(sessions):
    days = pd.DatetimeIndex(pd.to_datetime(sessions, errors="raise")).normalize()
    if days.empty or days.hasnans or days.has_duplicates or not days.is_monotonic_increasing:
        raise ValueError("Sessions must be nonempty, unique and chronological")
    if (days.dayofweek > 4).any():
        raise ValueError("A weekend was included as an HK session")
    return days


def anchor_position(days, locked, convention):
    d = pd.Timestamp(locked).normalize()
    if d < days[0] or d > days[-1]:
        raise ValueError("Restriction outside session coverage")
    k = days.searchsorted(d, side="right")
    pos = k - 1 if convention == "last_locked" else k
    if not 0 <= pos < len(days):
        raise ValueError("No session for event anchor")
    return int(pos)


def offset_position(days, anchor, offset, basis):
    if basis == "sessions":
        pos = anchor + offset
    else:
        d = days[anchor]
        # In calendar mode only ±21 denotes a calendar month; other numbers are days.
        target = d + (pd.DateOffset(months=1 if offset > 0 else -1)
                      if abs(offset) == 21 else pd.Timedelta(days=offset))
        if target < days[0] or target > days[-1]:
            raise ValueError("Calendar target outside session coverage")
        pos = days.searchsorted(target, side="right") - 1  # previous actual HK session
    if not 0 <= pos < len(days):
        raise ValueError("Offset outside session coverage")
    return int(pos)


def _positive(value):
    return pd.notna(value) and np.isfinite(value) and value > 0


def backtest(events, prices, sessions, strategies=None, settings=None,
             benchmark=None, locates=None, dividends=None, dividend_coverage=None):
    """Every eligible event × strategy produces either a result or a specific skip.

    Prices must use consistent split/capital-change units, without cash-dividend
    adjustments. Dividend inputs, if supplied, must use those same per-share units.
    A daily locate row is evidence for ENTRY only; recalls and intra-hold changes
    are not modeled. This remains a fixed-share event study, not a portfolio.
    """
    cfg = settings or Settings()
    strategies = make_strategies() if strategies is None else list(strategies)
    if len({s.key for s in strategies}) != len(strategies):
        raise ValueError("Strategy keys must be unique")
    p = validate_prices(prices)
    days = validate_sessions(sessions)
    asof = pd.Timestamp(cfg.asof)
    days = days[days <= asof]
    if days.empty:
        raise ValueError("No sessions before the as-of date")
    groups = {t: g.set_index("date") for t, g in p.groupby("ticker")}
    b = None
    if benchmark is not None and not benchmark.empty:
        b = benchmark.copy()
        b["date"] = pd.to_datetime(b.date).dt.normalize()
        if b.date.duplicated().any():
            raise ValueError("Duplicate benchmark dates")
        b = b.set_index("date")
    l = None
    if locates is not None and not locates.empty:
        l = locates.copy()
        l["date"] = pd.to_datetime(l.date).dt.normalize()
        if l.duplicated(["ticker", "date"]).any():
            raise ValueError("Duplicate locate rows")
        l = l.set_index(["ticker", "date"])
    div = None
    if dividends is not None and not dividends.empty:
        div = dividends.copy()
        div["date"] = pd.to_datetime(div.date).dt.normalize()
        div["cash_per_adjusted_share"] = pd.to_numeric(div.cash_per_adjusted_share, errors="raise")
        if (div.cash_per_adjusted_share < 0).any() or not np.isfinite(div.cash_per_adjusted_share).all():
            raise ValueError("Invalid cash dividends")
    coverage = None
    if dividend_coverage is not None and not dividend_coverage.empty:
        coverage = dividend_coverage.copy()
        for col in ("start", "end"):
            coverage[col] = pd.to_datetime(coverage[col]).dt.normalize()
        if (coverage.start > coverage.end).any():
            raise ValueError("Reversed dividend coverage")
        if dividends is None:
            raise ValueError("Dividend coverage needs a dividend ledger, even if empty")
    results = []
    for e in events.itertuples(index=False):
        locked = pd.Timestamp(e.last_day_locked)
        publication = pd.Timestamp(e.publication_date)
        g = groups.get(e.ticker)
        event_problem, anchor = "", None
        try:
            if locked >= asof:
                raise ValueError("Future or current-day restriction")
            anchor = anchor_position(days, locked, cfg.anchor)
            # Admission always checks ±30 around last-locked anchor, independent of UI T0.
            reference = anchor_position(days, locked, "last_locked")
            if cfg.require_complete_30:
                if reference < 30 or reference + 30 >= len(days):
                    raise ValueError("Incomplete ±30-session study window")
                if g is None:
                    raise ValueError("No price history")
                ends = g.reindex([days[reference - 30], days[reference], days[reference + 30]])
                if not ends.PX_LAST.map(_positive).all() or not ends.PX_VOLUME.map(_positive).all():
                    raise ValueError("Missing or nontraded ±30-session endpoint")
        except ValueError as exc:
            event_problem = str(exc)
        for s in strategies:
            row = dict(event_id=e.event_id, ticker=e.ticker, name=e.name,
                       input_date=e.input_date, last_day_locked=locked,
                       t0=days[anchor] if anchor is not None else pd.NaT,
                       strategy=s.key, strategy_label=s.label_for(cfg.day_basis), entry_offset=s.entry_offset,
                       exit_offset=s.exit_offset, entry_field=s.entry_field, exit_field=s.exit_field,
                       status="Skipped", reason=event_problem, entry_date=pd.NaT, exit_date=pd.NaT,
                       entry_price=np.nan, exit_price=np.nan, gross_return=np.nan,
                       after_cost_price_return=np.nan, net_return=np.nan,
                       borrow_cost=np.nan, transaction_cost=np.nan, dividend_cost=np.nan,
                       benchmark_return=np.nan, short_plus_hsi_return=np.nan,
                       calendar_days=np.nan, sessions_held=np.nan, path_missing_sessions=np.nan,
                       worst_mark_return=np.nan, best_mark_return=np.nan,
                       borrow_rate_used=cfg.borrow_rate, locate_status="Not verified",
                       source_url=e.source_url)
            if event_problem:
                results.append(row)
                continue
            try:
                i = offset_position(days, anchor, s.entry_offset, cfg.day_basis)
                j = offset_position(days, anchor, s.exit_offset, cfg.day_basis)
                entry, end = days[i], days[j]
                row.update(entry_date=entry, exit_date=end)
                if j < i or (j == i and not (s.entry_field == "PX_OPEN" and s.exit_field == "PX_LAST")):
                    raise ValueError("No positive holding interval")
                if end > asof:
                    raise ValueError("Exit after as-of date")
                # Date-only filings are treated as usable from the NEXT day.
                if pd.isna(publication) or publication >= entry:
                    raise ValueError("Restriction not demonstrably public before entry")
                if g is None:
                    raise ValueError("No price history")
                if entry not in g.index or end not in g.index:
                    raise ValueError("Missing exact entry or exit session")
                a, z = g.loc[entry], g.loc[end]
                pa, pz = a[s.entry_field], z[s.exit_field]
                if not _positive(pa) or not _positive(pz):
                    raise ValueError("Missing exact entry or exit price")
                if not _positive(a.PX_VOLUME) or not _positive(z.PX_VOLUME):
                    raise ValueError("Entry/exit volume missing or zero")
                rate = cfg.borrow_rate
                if l is not None and (e.ticker, entry) in l.index:
                    locate = l.loc[(e.ticker, entry)]
                    known = (str(locate.shortable).lower() in {"true", "1", "yes"}
                             and str(locate.borrow_available).lower() in {"true", "1", "yes"})
                    row["locate_status"] = "Entry evidence available" if known else "Entry not shortable/borrowable"
                    if cfg.require_locate and not known:
                        raise ValueError("Entry not shortable/borrowable")
                    if known and pd.notna(locate.get("borrow_rate", np.nan)):
                        rate = float(locate.borrow_rate)
                        if not np.isfinite(rate) or rate < 0:
                            raise ValueError("Invalid event borrow rate")
                elif cfg.require_locate:
                    raise ValueError("Missing dated entry shortability/borrow evidence")
                held = (end - entry).days
                days_charged = max(held, cfg.intraday_borrow_days)
                borrow = rate * days_charged / cfg.borrow_day_count
                # Fees are charged on both traded notionals, not twice initial notional.
                fee = cfg.fee_bps_per_leg / 10000 * (1 + pz / pa)
                gross = 1 - pz / pa
                path = g.reindex(days[i:j + 1])
                missing = int((~path.PX_LAST.map(_positive) | ~path.PX_VOLUME.map(_positive)).sum())
                marks = 1 - path.PX_LAST / pa
                worst = min(0.0, float(marks.min())) if missing == 0 else np.nan
                best = max(0.0, float(marks.max())) if missing == 0 else np.nan
                divcost, net = np.nan, np.nan
                covered = coverage is not None and ((coverage.ticker == e.ticker)
                           & (coverage.start <= entry) & (coverage.end >= end)).any()
                if covered:
                    # Ex-date entitlement attaches before the open: (entry, exit] for
                    # either open/close entry. A T0 open sale does not owe a T0 ex dividend.
                    amount = 0.0 if div is None else div.loc[(div.ticker == e.ticker)
                             & (div.date > entry) & (div.date <= end), "cash_per_adjusted_share"].sum()
                    divcost = float(amount / pa)
                    net = gross - fee - borrow - divcost
                benchret = np.nan
                if b is not None and entry in b.index and end in b.index:
                    ba, bz = b.loc[entry].get(s.entry_field), b.loc[end].get(s.exit_field)
                    if _positive(ba) and _positive(bz):
                        benchret = bz / ba - 1
                row.update(status="Completed", reason="", entry_price=float(pa), exit_price=float(pz),
                           gross_return=float(gross), after_cost_price_return=float(gross-fee-borrow),
                           borrow_cost=float(borrow), transaction_cost=float(fee), dividend_cost=divcost,
                           net_return=net, calendar_days=held, sessions_held=j-i,
                           path_missing_sessions=missing, worst_mark_return=worst, best_mark_return=best,
                           benchmark_return=benchret, short_plus_hsi_return=gross+benchret,
                           borrow_rate_used=rate)
            except ValueError as exc:
                row["reason"] = str(exc)
            results.append(row)
    return pd.DataFrame(results)


def summarize(trades, settings=None, common_sample=False):
    """Equal issuer weight: average repeated events within each issuer first.

    Confidence intervals resample issuer averages, not correlated event rows.
    They do not correct market-date dependence or multiple-strategy selection.
    """
    cfg = settings or Settings()
    if trades.empty:
        return pd.DataFrame()
    good = trades[trades.status.eq("Completed")].copy()
    if good.empty:
        return pd.DataFrame()
    # Exclude only structural zero-hold strategies from the common-cohort definition.
    meaningful = trades.loc[~((trades.entry_offset == trades.exit_offset)
                       & (trades.entry_field == trades.exit_field)), "strategy"].unique()
    if common_sample:
        counts = good.groupby("event_id").strategy.nunique()
        good = good[good.event_id.isin(counts[counts == len(meaningful)].index)]
    rng, out = np.random.default_rng(cfg.random_seed), []
    for key, g in good.groupby("strategy", sort=False):
        issuer = g.groupby("ticker").agg(gross=("gross_return", "mean"),
                                         after=("after_cost_price_return", "mean"),
                                         win=("gross_return", lambda x: (x > 0).mean()))
        x = issuer.gross.to_numpy()
        if len(x) >= 2:
            # Bounded loop avoids a potentially large bootstrap matrix.
            boot = np.array([rng.choice(x, len(x), replace=True).mean()
                             for _ in range(cfg.bootstrap_draws)])
            lo, hi = np.quantile(boot, [.025, .975])
        else:
            lo, hi = np.nan, np.nan
        out.append(dict(strategy=key, strategy_label=g.strategy_label.iloc[0],
                        events=len(g), issuers=len(issuer), mean_gross=float(x.mean()),
                        mean_after_cost_price=float(issuer.after.mean()),
                        median_issuer_gross=float(np.median(x)),
                        win_rate=float(issuer.win.mean()), ci_low=lo, ci_high=hi,
                        worst_event=float(g.gross_return.min()), best_event=float(g.gross_return.max()),
                        p05_event=float(g.gross_return.quantile(.05)),
                        average_calendar_days=float(g.calendar_days.mean()),
                        events_with_path_gaps=int(g.path_missing_sessions.gt(0).sum()),
                        sample="Common completed cohort" if common_sample else "Available per strategy"))
    return pd.DataFrame(out)


def sensitivity(trades, rates=(0, .05, .10, .20), fees=(0, 10, 25), settings=None):
    cfg = settings or Settings()
    if trades.empty:
        return pd.DataFrame()
    g = trades[trades.status.eq("Completed")].copy()
    out = []
    for rate in rates:
        for bps in fees:
            g["scenario_return"] = (g.gross_return - rate * g.calendar_days.clip(lower=cfg.intraday_borrow_days)
                                     / cfg.borrow_day_count - bps / 10000 * (1+g.exit_price/g.entry_price))
            for key, rows in g.groupby("strategy", sort=False):
                out.append(dict(strategy=key, borrow_rate=rate, fee_bps_per_leg=bps,
                                mean_after_cost_price=rows.groupby("ticker").scenario_return.mean().mean(),
                                issuers=rows.ticker.nunique(), events=len(rows)))
    return pd.DataFrame(out)


def event_paths(events, prices, sessions, settings=None, half_window=21):
    cfg = settings or Settings()
    days, p = validate_sessions(sessions), validate_prices(prices)
    groups = {t: g.set_index("date") for t, g in p.groupby("ticker")}
    records = []
    for e in events.itertuples(index=False):
        g = groups.get(e.ticker)
        if g is None:
            continue
        try:
            k = anchor_position(days, e.last_day_locked, cfg.anchor)
        except ValueError:
            continue
        if k < half_window or k+half_window >= len(days):
            continue
        dates = days[k-half_window:k+half_window+1]
        path = g.reindex(dates)
        base = path.PX_LAST.iloc[0]
        if not _positive(base) or not _positive(path.PX_VOLUME.iloc[0]):
            continue
        for offset, d, close, vol in zip(range(-half_window, half_window+1), dates, path.PX_LAST, path.PX_VOLUME):
            value = 1-close/base if _positive(close) and _positive(vol) else np.nan
            records.append(dict(event_id=e.event_id, ticker=e.ticker, name=e.name,
                                last_day_locked=e.last_day_locked, offset=offset, date=d,
                                short_return_from_Tminus21=value))
    return pd.DataFrame(records)


def plot_results(summary, trades, paths, output_dir=None, settings=None):
    """Return four professional figures; save high-resolution PNGs when requested."""
    import matplotlib.pyplot as plt
    from matplotlib.ticker import PercentFormatter
    cfg = settings or Settings()
    plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 10,
                         "axes.spines.top": False, "axes.spines.right": False,
                         "axes.edgecolor": "#c8c8c8", "axes.labelcolor": "#333333",
                         "figure.facecolor": "white", "axes.facecolor": "white"})
    figures = {}
    if summary.empty:
        return figures
    s = summary.iloc[::-1].reset_index(drop=True)
    fig, ax = plt.subplots(figsize=(11, max(6, len(s)*.34)), layout="constrained")
    y = np.arange(len(s))
    ax.barh(y, s.mean_gross, height=.6, color=np.where(s.mean_gross >= 0, "#23745b", "#CB2125"))
    valid = s.ci_low.notna() & s.ci_high.notna()
    # Draw intervals by endpoints; no assumption that percentile CI straddles estimate.
    ax.hlines(y[valid], s.loc[valid, "ci_low"], s.loc[valid, "ci_high"], color="#333333", lw=1.2)
    ax.set_yticks(y, [f"{r.strategy_label}   (n={r.issuers})" for r in s.itertuples()])
    ax.axvline(0, color="#999999", lw=.7)
    ax.xaxis.set_major_formatter(PercentFormatter(1))
    ax.set_title("Short price returns by strategy", loc="left", fontweight="bold", pad=14)
    ax.set_xlabel("P&L / initial short notional · equal issuer weight · 95% issuer-bootstrap interval")
    ax.grid(axis="x", alpha=.15)
    figures["strategy_comparison"] = fig

    g = trades[trades.status.eq("Completed")]
    pre = g[(g.entry_offset < 0) & g.exit_offset.isin([-1, 0])]
    if not pre.empty:
        issuer = pre.groupby(["entry_offset", "exit_offset", "ticker"]).gross_return.mean().reset_index()
        grid = issuer.groupby(["entry_offset", "exit_offset"]).gross_return.mean().unstack().reindex([-21,-14,-10,-7,-3,-2,-1])
        counts = issuer.groupby(["entry_offset", "exit_offset"]).ticker.nunique().unstack().reindex(grid.index)
        fig, ax = plt.subplots(figsize=(7, 5), layout="constrained")
        vmax = max(.01, float(np.nanmax(np.abs(grid.to_numpy()))))
        im = ax.imshow(grid, cmap="RdYlGn", vmin=-vmax, vmax=vmax, aspect="auto")
        ax.set_yticks(range(len(grid)), ["T−1m close" if cfg.day_basis == "calendar" and n == -21 else f"T{n} close" for n in grid.index])
        ax.set_xticks(range(len(grid.columns)), [f"T{n:+d} close".replace("T+0", "T0") for n in grid.columns])
        for i in range(len(grid)):
            for j in range(len(grid.columns)):
                val = grid.iloc[i,j]
                label = f"{val:+.2%}\nn={int(counts.iloc[i,j])}" if pd.notna(val) else "No holding interval"
                ax.text(j, i, label, ha="center", va="center", fontsize=10, color="#202020")
        ax.set_title("Entry timing versus exit timing", loc="left", fontweight="bold", pad=14)
        ax.set_ylabel("Short entry")
        ax.set_xlabel("Cover · mean short price return, equal issuer weight")
        fig.colorbar(im, ax=ax, format=PercentFormatter(1), shrink=.75)
        figures["entry_exit_matrix"] = fig

    if not paths.empty:
        by_stock = paths.groupby(["offset", "ticker"]).short_return_from_Tminus21.mean().reset_index()
        curve = by_stock.groupby("offset").short_return_from_Tminus21.agg(["mean", "count"])
        fig, (ax, axn) = plt.subplots(2, 1, figsize=(10, 5.5), sharex=True,
                                     gridspec_kw={"height_ratios": [4,1]}, layout="constrained")
        ax.plot(curve.index, curve["mean"], color="#153954", lw=2, label="Selected average · one weight per issuer")
        ax.axvline(0, color="#CB2125", lw=1, linestyle="--", label="T0")
        ax.axhline(0, color="#aaaaaa", lw=.7)
        ax.yaxis.set_major_formatter(PercentFormatter(1))
        ax.set_ylabel("Short price P&L / initial notional")
        ax.set_title("Event study: enter at T−21 close", loc="left", fontweight="bold", pad=12)
        ax.legend(frameon=False, loc="best")
        ax.grid(alpha=.15)
        axn.bar(curve.index, curve["count"], color="#b5bbc2")
        axn.set_ylabel("Issuers")
        axn.set_xlabel("Hong Kong trading sessions from T0 · available observations at each offset")
        figures["event_study"] = fig

    selected = [key for key in summary.strategy if key in set(g.strategy)]
    if selected:
        fig, ax = plt.subplots(figsize=(11, 6), layout="constrained")
        values = [g.loc[g.strategy.eq(key), "gross_return"].to_numpy() for key in selected]
        import inspect
        orientation = {"orientation":"horizontal"} if "orientation" in inspect.signature(ax.boxplot).parameters else {"vert":False}
        ax.boxplot(values, **orientation, showfliers=True, flierprops={"markersize": 2, "alpha": .45})
        ax.set_yticks(range(1, len(selected)+1), [summary.set_index("strategy").loc[key,"strategy_label"] for key in selected])
        ax.axvline(0, color="#CB2125", lw=.8)
        ax.xaxis.set_major_formatter(PercentFormatter(1))
        ax.set_xlabel("Individual-event short price return · outliers retained")
        ax.set_title("Dispersion and adverse outcomes", loc="left", fontweight="bold", pad=12)
        figures["return_distribution"] = fig
    if output_dir:
        d = Path(output_dir)
        d.mkdir(parents=True, exist_ok=True)
        for name, fig in figures.items():
            fig.savefig(d / f"{name}.png", dpi=160, facecolor="white", bbox_inches="tight")
    return figures


def plot_stocks(paths, tickers, page=1):
    """Individual stock paths only, six event lines per page, complete legends."""
    import matplotlib.pyplot as plt
    from matplotlib.ticker import PercentFormatter
    chosen = {normalize_ticker(t) for t in tickers}
    p = paths[paths.ticker.isin(chosen)]
    keys = list(p.event_id.unique())
    if not keys:
        raise ValueError("No observed paths for selected tickers")
    pages = (len(keys)+5)//6
    if not 1 <= page <= pages:
        raise ValueError(f"Choose page 1 to {pages}")
    fig, ax = plt.subplots(figsize=(10,6), layout="constrained")
    for event_id in keys[(page-1)*6:page*6]:
        g = p[p.event_id.eq(event_id)].sort_values("offset")
        r = g.iloc[0]
        label = f"{r.ticker} · {r['name']} · locked through {pd.Timestamp(r.last_day_locked):%d %b %Y}"
        ax.plot(g.offset, g.short_return_from_Tminus21, lw=1.7, label=label)
    ax.axvline(0, color="#999999", ls="--", lw=.8)
    ax.axhline(0, color="#bbbbbb", lw=.7)
    ax.yaxis.set_major_formatter(PercentFormatter(1))
    ax.set(title=f"Selected stock events · page {page}/{pages}", xlabel="HK trading sessions from T0",
           ylabel="Short price return from T−21 close")
    ax.legend(loc="upper center", bbox_to_anchor=(.5,-.12), frameon=False, fontsize=9)
    return fig


class BloombergClient:
    """Synchronous Desktop API, one request at a time, bounded waits and disk cache.

    A running entitled Desktop API session is required. This is not a BQL client.
    No cache is used unless its exact request signature and data SHA256 match.
    """
    def __init__(self, cache_dir="cache/bloomberg", host="localhost", port=8194, timeout=60):
        self.cache_dir = Path(cache_dir)
        self.host, self.port, self.timeout = host, port, timeout
        self.session = None

    def __enter__(self):
        try:
            import blpapi
        except ImportError as exc:
            raise RuntimeError("blpapi is not installed in this Jupyter kernel. See setup instructions.") from exc
        self.api = blpapi
        options = blpapi.SessionOptions()
        options.setServerHost(self.host)
        options.setServerPort(self.port)
        options.setConnectTimeout(min(int(self.timeout*1000), 30000))
        self.session = blpapi.Session(options)
        if not self.session.start():
            self.session.stop()
            raise RuntimeError("Cannot start Bloomberg Desktop API. Log in on the entitled terminal workstation.")
        if not self.session.openService("//blp/refdata"):
            self.session.stop()
            raise RuntimeError("Cannot open //blp/refdata. Check API entitlements.")
        self.service = self.session.getService("//blp/refdata")
        self.cache_dir.mkdir(parents=True, exist_ok=True)
        return self

    def __exit__(self, *args):
        if self.session:
            self.session.stop()

    @staticmethod
    def parse_message(msg, fields):
        """Parse security-level historical responses without converting errors to zero."""
        if msg.hasElement("responseError"):
            raise RuntimeError("Bloomberg responseError: " + str(msg.getElement("responseError")))
        if not msg.hasElement("securityData"):
            return [], []
        sec = msg.getElement("securityData")
        security = sec.getElementAsString("security")
        errors, rows = [], []
        if sec.hasElement("securityError"):
            return [], [dict(ticker=security, field="*", error=str(sec.getElement("securityError")))]
        if sec.hasElement("fieldExceptions"):
            ex = sec.getElement("fieldExceptions")
            for i in range(ex.numValues()):
                item = ex.getValueAsElement(i)
                errors.append(dict(ticker=security, field=item.getElementAsString("fieldId"),
                                   error=str(item.getElement("errorInfo"))))
        if sec.hasElement("fieldData"):
            data = sec.getElement("fieldData")
            for i in range(data.numValues()):
                item = data.getValueAsElement(i)
                row = dict(ticker=security, date=pd.Timestamp(item.getElementAsDatetime("date")).date().isoformat())
                for field in fields:
                    row[field] = (item.getElementAsFloat(field)
                                  if item.hasElement(field) and not item.getElement(field).isNull() else np.nan)
                rows.append(row)
        return rows, errors

    def history(self, tickers, start, end, fields=PRICE_FIELDS, refresh=False):
        if not self.session:
            raise RuntimeError("Use BloombergClient inside a with block")
        securities = sorted(set(tickers))
        if not securities or len(securities) > 20:
            raise ValueError("Request between 1 and 20 securities at a time")
        start, end = pd.Timestamp(start).strftime("%Y%m%d"), pd.Timestamp(end).strftime("%Y%m%d")
        params = dict(periodicitySelection="DAILY", periodicityAdjustment="ACTUAL",
                      nonTradingDayFillOption="ACTIVE_DAYS_ONLY", nonTradingDayFillMethod="NIL_VALUE",
                      adjustmentFollowDPDF=False, adjustmentSplit=True,
                      adjustmentNormal=False, adjustmentAbnormal=False)
        signature = dict(version=VERSION, securities=securities, fields=list(fields), start=start, end=end, options=params)
        key = hashlib.sha256(json.dumps(signature, sort_keys=True).encode()).hexdigest()
        csv, meta = self.cache_dir / (key+".csv"), self.cache_dir / (key+".json")
        if not refresh and csv.exists() and meta.exists():
            info = json.loads(meta.read_text())
            if info.get("request") != signature or info.get("sha256") != hashlib.sha256(csv.read_bytes()).hexdigest():
                raise ValueError("Bloomberg cache metadata/hash mismatch")
            return pd.read_csv(csv), pd.DataFrame(info["errors"])
        req = self.service.createRequest("HistoricalDataRequest")
        for ticker in securities:
            req.getElement("securities").appendValue(ticker)
        for field in fields:
            req.getElement("fields").appendValue(field)
        req.set("startDate", start)
        req.set("endDate", end)
        for k, v in params.items():
            req.set(k, v)
        queue = self.api.EventQueue()
        cid = self.session.sendRequest(req, eventQueue=queue)
        deadline, records, errors = time.monotonic()+self.timeout, [], []
        while True:
            remaining = deadline-time.monotonic()
            if remaining <= 0:
                self.session.cancel(cid)
                raise TimeoutError("Bloomberg history timed out; incomplete response not cached")
            event = queue.nextEvent(min(1000, max(1, int(remaining*1000))))
            kind = event.eventType()
            if kind == self.api.Event.REQUEST_STATUS:
                for msg in event:
                    if str(msg.messageType()) == "RequestFailure":
                        raise RuntimeError("Bloomberg RequestFailure: "+str(msg))
            if kind in (self.api.Event.PARTIAL_RESPONSE, self.api.Event.RESPONSE):
                for msg in event:
                    rr, ee = self.parse_message(msg, fields)
                    records.extend(rr)
                    errors.extend(ee)
                if kind == self.api.Event.RESPONSE:
                    break
        frame = pd.DataFrame(records, columns=["ticker", "date", *fields])
        represented = set(frame.ticker) | {e['ticker'] for e in errors}
        errors.extend(dict(ticker=t, field="*", error="No rows returned in requested historical interval")
                      for t in securities if t not in represented)
        if not frame.empty and frame.duplicated(["ticker", "date"]).any():
            raise ValueError("Bloomberg returned duplicate security/date records")
        frame.to_csv(csv, index=False)
        meta.write_text(json.dumps(dict(request=signature, errors=errors,
                             downloaded_utc=pd.Timestamp.now(tz="UTC").isoformat(),
                             sha256=hashlib.sha256(csv.read_bytes()).hexdigest()), indent=2))
        return frame, pd.DataFrame(errors)


def completed_history_cutoff(asof, now=None):
    """Never request the current HK daily bar, even if a terminal returns a live last."""
    stamp=pd.Timestamp.now(tz="Asia/Hong_Kong") if now is None else pd.Timestamp(now)
    stamp=stamp.tz_localize("Asia/Hong_Kong") if stamp.tzinfo is None else stamp.tz_convert("Asia/Hong_Kong")
    previous_day=stamp.normalize().tz_localize(None)-pd.Timedelta(days=1)
    return min(pd.Timestamp(asof).normalize(),previous_day)


def download_bloomberg(events, settings, cache_dir, refresh=False):
    """Pull only reviewed historic events; one security/window at a time.

    HSI actual daily observations supply the common session calendar. Missing HSI
    data are not silently replaced by generic Monday–Friday business days.
    """
    if events.empty:
        raise ValueError("No reviewed historical events to download")
    asof = completed_history_cutoff(settings.asof)
    print(f"Bloomberg completed-history cutoff: {asof:%Y-%m-%d} (current HK day excluded)")
    oldest = pd.to_datetime(events.last_day_locked).min()-pd.Timedelta(days=130)
    latest = min(asof, pd.to_datetime(events.last_day_locked).max()+pd.Timedelta(days=130))
    frames, error_frames = [], []
    with BloombergClient(cache_dir) as client:
        bench, errors = client.history(["HSI Index"], oldest, latest, fields=("PX_OPEN","PX_LAST"), refresh=refresh)
        if not errors.empty or bench.empty:
            raise RuntimeError("HSI calendar request is incomplete: "+errors.to_string(index=False))
        bench["date"] = pd.to_datetime(bench.date)
        if not bench.PX_LAST.map(_positive).all():
            raise ValueError("HSI calendar has missing/invalid observations")
        sessions = validate_sessions(bench.sort_values("date").date)
        # Dated session source is exportable and can be compared with an approved HK calendar.
        for n, (ticker, g) in enumerate(events.groupby("ticker"), 1):
            first = pd.to_datetime(g.last_day_locked).min()-pd.Timedelta(days=130)
            last = min(asof, pd.to_datetime(g.last_day_locked).max()+pd.Timedelta(days=130))
            frame, errors = client.history([ticker], first, last, refresh=refresh)
            if not frame.empty:
                # Keep missing numeric fields as float NaN and omit empty responses.
                # This avoids pandas' ambiguous empty/all-NA concatenation behavior.
                frame=frame.copy()
                for field in PRICE_FIELDS:
                    frame[field]=pd.to_numeric(frame[field],errors="raise").astype(float)
                frames.append(frame)
            if not errors.empty:
                error_frames.append(errors)
            if n % 20 == 0:
                print(f"Bloomberg: processed {n}/{events.ticker.nunique()} reviewed tickers")
    prices = pd.concat(frames, ignore_index=True) if frames else pd.DataFrame(columns=["ticker","date",*PRICE_FIELDS])
    errors = pd.concat(error_frames, ignore_index=True) if error_frames else pd.DataFrame(columns=["ticker","field","error"])
    return validate_prices(prices), sessions, bench.drop(columns="ticker"), errors


def export_results(output_dir, audit, trades, summary, common_summary, paths, settings, metadata=None):
    out = Path(output_dir)
    out.mkdir(parents=True, exist_ok=True)
    files = dict(input_audit=audit, trades=trades, strategy_summary=summary,
                 common_sample_summary=common_summary, event_paths=paths)
    if not trades.empty:
        files["skipped_trades"] = trades[~trades.status.eq("Completed")]
        files["per_event_returns"] = trades.pivot(index=["event_id","ticker","last_day_locked"],
                                                 columns="strategy", values="gross_return").reset_index()
    for name, df in files.items():
        df = price_only_columns(df)
        df.to_csv(out / f"{name}.csv", index=False, float_format="%.12g")
    info = dict(version=VERSION, settings=asdict(settings), **(metadata or {}))
    info["csv_sha256"] = {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in sorted(out.glob("*.csv"))}
    (out / "run_manifest.json").write_text(json.dumps(info, indent=2, default=str))
    return out


def price_only_columns(frame):
    """The delivered report contains gross price results, with no assumed-cost view."""
    return frame.drop(columns=['after_cost_price_return','net_return','borrow_cost',
                              'transaction_cost','dividend_cost','borrow_rate_used',
                              'mean_after_cost_price','mean_net_if_covered'], errors='ignore')


def export_excel(path, audit, trades, summary, common_summary, prices, paths, settings=None, metadata=None):
    """Create the requested Python-generated Excel snapshot after each notebook run.

    Python owns the calculations; the XLSX is a dated results export, not a second
    calculation engine. Native charts reference the visible numeric result cells.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.chart import BarChart, LineChart, Reference
    from openpyxl.chart.shapes import GraphicalProperties
    from openpyxl.chart.label import DataLabelList
    from openpyxl.chart.marker import DataPoint
    from openpyxl.utils import get_column_letter
    cfg, meta = settings or Settings(), metadata or {}
    wb = Workbook()
    ws = wb.active
    ws.title = 'Summary'
    grey, ink, green, red = 'EDEFF1', '24282B', '23745B', 'CB2125'
    thin = Side(style='thin', color='D9DDE1')
    pct = '0.00%;[Red](0.00%);"—"'
    rate_columns = {'mean_gross','mean_after_cost_price','median_issuer_gross','win_rate',
                    'ci_low','ci_high','worst_event','best_event','p05_event',
                    'gross_return','after_cost_price_return','net_return','borrow_cost',
                    'transaction_cost','dividend_cost','benchmark_return','short_plus_hsi_return',
                    'worst_mark_return','best_mark_return','borrow_rate_used','borrow_rate',
                    'short_return_from_Tminus21','Mean return'}
    labels = {'strategy_label':'Strategy','events':'Events','issuers':'Issuers',
              'mean_gross':'Mean gross return','mean_after_cost_price':'After assumed costs',
              'win_rate':'Win rate','ci_low':'95% interval low','ci_high':'95% interval high',
              'worst_event':'Worst event','median_issuer_gross':'Median issuer return',
              'ticker':'Ticker','name':'Company','last_day_locked':'Last day locked',
              'input_date':'Supplied date','source_row':'Source row','t0':'T0 session',
              'entry_date':'Entry date','exit_date':'Exit date','entry_price':'Entry price',
              'exit_price':'Exit price','gross_return':'Gross return',
              'after_cost_price_return':'After-cost price return','net_return':'After dividends and costs',
              'status':'Status','reason':'Exclusion reason','date':'Date',
              'PX_OPEN':'Open','PX_LAST':'Close','PX_VOLUME':'Volume',
              'source_url':'Primary source','category':'Holder category','borrow_rate':'Assumed annual borrow',
              'fee_bps_per_leg':'Fee per leg (bps)'}
    if not trades.empty:
        labels.update(trades.drop_duplicates('strategy').set_index('strategy').strategy_label.to_dict())

    def value(x):
        if x is None or (not isinstance(x,(list,dict)) and pd.isna(x)):
            return None
        if isinstance(x,pd.Timestamp):
            return x.to_pydatetime().replace(tzinfo=None)
        if isinstance(x,np.generic):
            return x.item()
        return x

    def sheet_style(s):
        s.sheet_view.showGridLines = False
        s.freeze_panes = 'C2'
        s.sheet_properties.pageSetUpPr.fitToPage = True
        s.page_setup.orientation = 'landscape'
        s.page_setup.paperSize = s.PAPERSIZE_A3
        s.page_setup.fitToWidth = 1
        s.page_setup.fitToHeight = 0
        s.print_options.horizontalCentered = True
        s.sheet_properties.outlinePr.summaryRight = False

    def write_table(s, df, start_row=1, start_col=1, filters=True):
        if len(df.columns)==0:
            s.cell(start_row,start_col,'No observations')
            return start_row
        for j,col in enumerate(df.columns,start_col):
            cell=s.cell(start_row,j,labels.get(col,str(col).replace('_',' ').capitalize()))
            cell.font=Font(name='Arial',size=10,bold=True,color=ink)
            cell.fill=PatternFill('solid',fgColor=grey)
            cell.alignment=Alignment(horizontal='left',vertical='center',wrap_text=True)
            cell.border=Border(bottom=thin)
            width=42 if col in {'strategy_label','name','reason','status'} else 17
            if col.startswith(('pre_','post_','cross_','expiry_')):
                width=24
            if col in {'source_url','event_id','input_id'}:
                width=38
            s.column_dimensions[get_column_letter(j)].width=width
        s.row_dimensions[start_row].height=32
        for i,row in enumerate(df.itertuples(index=False,name=None),start_row+1):
            for j,(col,x) in enumerate(zip(df.columns,row),start_col):
                x=value(x)
                cell=s.cell(i,j,x)
                if isinstance(x,str):
                    cell.data_type='s'  # Literal strings cannot become executable formulas.
                cell.font=Font(name='Arial',size=10,color=ink)
                cell.alignment=Alignment(vertical='center',horizontal='left' if isinstance(x,str) else 'right')
                if isinstance(x,(datetime,date)):
                    cell.number_format='dd-mmm-yyyy'
                elif col in rate_columns or col.startswith(('pre_','post_','cross_','expiry_')):
                    cell.number_format=pct
                elif col=='PX_VOLUME':
                    cell.number_format='#,##0'
                elif isinstance(x,(float,np.floating)):
                    cell.number_format='#,##0.0000'
                elif isinstance(x,(int,np.integer)):
                    cell.number_format='#,##0'
                if col=='source_url' and isinstance(x,str) and x.startswith('https://'):
                    cell.hyperlink=x
                    cell.value='HKEX filing' if 'hkexnews.hk/' in x else 'Primary filing'
                    cell.font=Font(name='Arial',size=10,color='245577',underline='single')
            s.row_dimensions[i].height=18
        last=start_row+len(df)
        if filters and len(df):
            s.auto_filter.ref=f'{get_column_letter(start_col)}{start_row}:{get_column_letter(start_col+len(df.columns)-1)}{last}'
            s.print_title_rows=f'{start_row}:{start_row}'
        for j,col in enumerate(df.columns,start_col):
            if col in rate_columns or col.startswith(('pre_','post_','cross_','expiry_')):
                area=f'{get_column_letter(j)}{start_row+1}:{get_column_letter(j)}{max(last,start_row+1)}'
                s.conditional_formatting.add(area,CellIsRule(operator='lessThan',formula=['0'],font=Font(color=red)))
                s.conditional_formatting.add(area,CellIsRule(operator='greaterThan',formula=['0'],font=Font(color=green)))
        return last

    def dates(df):
        d=df.copy()
        for col in d.columns:
            if col in {'input_date','last_day_locked','t0','entry_date','exit_date','date'}:
                d[col]=pd.to_datetime(d[col],errors='coerce')
        return d

    sheet_style(ws)
    ws.freeze_panes='B6'
    # Compact result context, no title banner or Notes tab.
    for row,(key,x) in enumerate([
        ('Data',meta.get('data_source','')),
        ('As of',pd.Timestamp(cfg.asof)),
        ('T0', 'Last session on/before last day locked' if cfg.anchor=='last_locked' else 'First session after last day locked'),
        ('Return basis','Gross short price P&L / initial notional')],1):
        ws.cell(row,1,key).font=Font(name='Arial',size=10,bold=True,color=ink)
        c=ws.cell(row,2,value(x)); c.font=Font(name='Arial',size=10,color=ink)
        if row==2:c.number_format='dd-mmm-yyyy'
    ws['L2']='Days'; ws['M2']='HK sessions; month = 21' if cfg.day_basis=='sessions' else 'Calendar days / calendar month'
    columns=['strategy_label','events','issuers','mean_gross',
             'win_rate','median_issuer_gross','worst_event','ci_low','ci_high']
    available=summary[[c for c in columns if c in summary]]
    last=write_table(ws,available,start_row=5,filters=False)
    common_start=last+4
    ws.cell(common_start-1,1,'Common completed cohort').font=Font(name='Arial',size=10,bold=True,color=ink)
    write_table(ws,common_summary[[c for c in columns if c in common_summary]],start_row=common_start,filters=False)
    ws.column_dimensions['A'].width=38
    for col in 'BC': ws.column_dimensions[col].width=12
    if not summary.empty:
        chart=BarChart()
        chart.type='bar'; chart.style=2; chart.title='Mean short returns'
        chart.y_axis.title='Return / initial notional'
        chart.y_axis.numFmt='0%'
        chart.add_data(Reference(ws,min_col=4,max_col=4,min_row=5,max_row=last),titles_from_data=True)
        chart.set_categories(Reference(ws,min_col=1,min_row=6,max_row=last))
        chart.legend=None
        chart.width=24; chart.height=15
        for series,color in zip(chart.series,['23745B','153954']):
            series.graphicalProperties.solidFill=color
            series.graphicalProperties.line.solidFill=color
            series.data_points=[DataPoint(idx=i,spPr=GraphicalProperties(
                solidFill='23745B' if v>=0 else 'CB2125')) for i,v in enumerate(summary.mean_gross)]
        ws.add_chart(chart,'L5')
    if not paths.empty:
        issuer=paths.groupby(['offset','ticker']).short_return_from_Tminus21.mean().reset_index()
        curve=issuer.groupby('offset').short_return_from_Tminus21.agg(['mean','count']).reset_index()
        curve.columns=['Session offset','Mean return','Issuers observed']
        curve_start=common_start+len(common_summary)+4
        end=write_table(ws,curve,start_row=curve_start,filters=False)
        chart=LineChart(); chart.title='Event study from T−21 close'; chart.style=13
        chart.y_axis.title='Short price return'; chart.y_axis.numFmt='0%'
        chart.x_axis.title='HK sessions from T0'; chart.legend=None
        chart.add_data(Reference(ws,min_col=2,min_row=curve_start,max_row=end),titles_from_data=True)
        chart.set_categories(Reference(ws,min_col=1,min_row=curve_start+1,max_row=end))
        chart.series[0].graphicalProperties.line.solidFill='153954'
        chart.series[0].graphicalProperties.line.width=24000
        chart.smooth=False
        chart.series[0].smooth=False
        chart.width=24; chart.height=12; chart.display_blanks='gap'
        ws.add_chart(chart,'L35')

    if not trades.empty:
        event=trades.pivot(index=['event_id','ticker','name','input_date','last_day_locked','t0'],
                          columns='strategy',values='gross_return').reset_index()
        # Omit events with no completed strategy from the return grid; all remain in Trades/Inputs.
        event=event[event.event_id.isin(trades.loc[trades.status.eq('Completed'),'event_id'])]
    else:
        event=pd.DataFrame()
    input_cols=['source_row','ticker','name','input_date','last_day_locked','category','status','source_url']
    trade_cols=['ticker','name','last_day_locked','t0','strategy_label','status','reason',
                'entry_date','exit_date','entry_price','exit_price','gross_return',
                'after_cost_price_return','net_return','borrow_cost','transaction_cost','dividend_cost',
                'calendar_days','sessions_held','path_missing_sessions','worst_mark_return','best_mark_return',
                'benchmark_return','short_plus_hsi_return','locate_status','borrow_rate_used',
                'input_date','entry_field','exit_field','event_id','source_url']
    datasets=[('By event',event),('Trades',trades[[c for c in trade_cols if c in trades]]),
              ('Inputs',audit[[c for c in input_cols if c in audit]]),
              ('Prices',prices)]
    for name,df in datasets:
        df=price_only_columns(df)
        s=wb.create_sheet(name); sheet_style(s); write_table(s,dates(df))
        s.sheet_properties.tabColor={'By event':'23745B','Trades':'153954','Inputs':'888888',
                                    'Prices':'888888','Cost scenarios':'BA9846'}[name]
        if name=='By event':
            s.column_dimensions['A'].hidden=True
            s.freeze_panes='D2'
    ws.sheet_properties.tabColor=red
    wb.active=0
    path=Path(path); path.parent.mkdir(parents=True,exist_ok=True)
    wb.save(path)
    wb.close()
    return path


def write_html_report(output, audit, trades, summary, common, figures, metadata, settings=None):
    """Standalone readable report with embedded PNGs. No remote scripts or fonts."""
    import base64
    import io
    cfg = settings or Settings()
    anchor_text = ("the last Hong Kong session on or before the confirmed last day locked"
                   if cfg.anchor == "last_locked" else "the first Hong Kong session strictly after the confirmed last day locked")
    basis_text = ("Hong Kong trading sessions; one month is 21 sessions" if cfg.day_basis == "sessions"
                  else "calendar days, mapped to the preceding Hong Kong session; one month is a calendar month")
    def table(df, percent_cols=()):
        d = df.copy()
        for col in percent_cols:
            if col in d:
                d[col] = d[col].map(lambda x: f"{x:+.2%}" if pd.notna(x) else "Unavailable")
        return d.to_html(index=False, border=0, escape=True, na_rep="Unavailable")
    n = 0 if trades.empty else int(trades.status.eq("Completed").sum())
    s = summary.rename(columns={"strategy_label":"Strategy", "events":"Events", "issuers":"Issuers",
                               "mean_gross":"Mean price return", "mean_after_cost_price":"After assumed costs",
                               "win_rate":"Win rate", "worst_event":"Worst event"})
    cols = [c for c in ["Strategy","Events","Issuers","Mean price return","Win rate","Worst event"] if c in s]
    parts = ["<!doctype html><html><head><meta charset='utf-8'><meta name='viewport' content='width=device-width,initial-scale=1'>",
             "<title>HK IPO lock-up backtest</title><style>body{font:15px/1.55 Arial,sans-serif;color:#282828;max-width:1180px;margin:36px auto;padding:0 24px}h1{font-size:25px;border-bottom:3px solid #CB2125;padding-bottom:10px}h2{font-size:18px;margin-top:32px}.meta{color:#5a5a5a}.table{overflow:auto}table{border-collapse:collapse;width:100%;font-size:13px}th{background:#eceef0;text-align:right}td,th{padding:9px 10px;border-bottom:1px solid #ddd;white-space:nowrap}td:first-child,th:first-child{text-align:left}td{text-align:right}img{display:block;max-width:100%;height:auto;margin:28px 0}a{color:#234f70}code{background:#eee;padding:2px 4px}details{margin:16px 0}summary{cursor:pointer;font-weight:bold}</style></head><body>",
             "<h1>HK IPO lock-up backtest</h1>",
             f"<p class='meta'>{html.escape(str(metadata.get('data_source','')))} · {n:,} completed event–strategy observations</p>",
             f"<p>T0 is {anchor_text}. Offsets are {basis_text}. Positive returns mean the short made money.</p>",
             "<p>Gross short price returns only. No assumed costs are deducted. Dividends, borrow, fees and execution constraints are outside this price study. Returns use initial short notional, not margin capital.</p>",
             "<h2>Strategy comparison</h2><div class='table'>",
             table(s[cols], ["Mean price return","After assumed costs","Win rate","Worst event"]), "</div>",
             "<p>Each issuer has one weight; repeated expiry events are averaged within issuer. Strategies can have different samples. The common-cohort CSV is provided for a like-for-like comparison. Testing many timings is exploratory and does not establish an out-of-sample trading edge.</p>"]
    for name, fig in figures.items():
        buffer = io.BytesIO()
        fig.savefig(buffer, format="png", dpi=130, bbox_inches="tight")
        parts.append(f"<img alt='{html.escape(name)}' src='data:image/png;base64,{base64.b64encode(buffer.getvalue()).decode()}'>")
    counts = audit.status.value_counts().rename_axis("Input review status").reset_index(name="Dates")
    parts += ["<h2>Coverage</h2><div class='table'>", table(counts), "</div>",
              "<p>Input rows are a supplied snapshot. The admitted subset has primary-reviewed cornerstone dates and completed ±30-session endpoints. This retrospective selection can introduce survivorship and look-ahead selection effects; it is not every IPO.</p>"]
    if not trades.empty:
        skipped = trades.loc[~trades.status.eq("Completed"),"reason"].value_counts().rename_axis("Skipped reason").reset_index(name="Event–strategy rows")
        parts += ["<details><summary>Skipped observations</summary><div class='table'>",table(skipped),"</div></details>"]
    parts += ["<p>Open the companion notebook to inspect each event, change settings, pull Bloomberg data, and export the full audit and trade ledger.</p></body></html>"]
    Path(output).write_text("\n".join(parts), encoding="utf-8")

# Embedded source snapshot: date/ticker rows and primary-review decisions only.
# No market prices are embedded. SHA256 verifies the snapshot before it is read.
_EMBEDDED_INPUTS_B64 = (
    'eNrsveuOXEeSJvgqif41BaRKbmZu5u7T/3YumMHs7hRmdv8VUPBrF9EUqSGpri7sn3mH3SecJ1nzc4k4zPAgmalzMpgpCRJlzJMR'
    'x83NL3b97P/5h/ov9d2nv3yo//Km/u2P+eO//MO/vfuHT2/yP9cP92/e/fzLp7+U+Knev40fO/X3v7x9r4/KfdYf/tP7D3+/z+/f'
    'tTcfftIfxZ/qu/JT/7K0/NLH9798yPUvv3x4u5If/xqRZf3bz/Gf6v3Pv6S3b/Tr3rx/N79qHspn9Lv3Sn+MP/38tv4l/7X2L/8L'
    'GhQw6HUo78qb/tt/aW/e6oNOftz8dDOK8xDnH3788zv9Cnv3n/7L3X/4H7+8+fT3e/1W/MHov/b+/v9+98/v3v/t3f1/jG8/1uXP'
    '/k9/8w9gfgB3/3++v6v/GvOnu58/vPkpfvj7D/OAa7nrr777KX7Kf/3Hu/L+Tjm4e/Ou1Q938e7jm3/94af37z79VT/785sPf//j'
    '+t1//fTp54//9scf//a3v8Ef//rP9V/f1b99VOLHt28+fqolv58InXH98Y8fqz7QocCPBo2dqE4YY8j+8a+ffrof8QY/GP86eCPH'
    '5nPe+AeAHyBsyX/3/sO7+uHjJ/3W+//rwy8ro08aC/9oGMJEdcIAsvvjz6Xdm8JeBxNLTLla8cmYxpIiUs4t9T/B60otAQywSU5K'
    'TaWRCz5ZihbvQeYxG16GP0lAV/a/P81z/Xj36a/1blrgd30vqgz+fjdvs39cBdQfv3n3T3fLDrl7/+4O/d1/zZ/eJ5VN/94/zrNw'
    'nqG7f7OZoj/sPDEqJHMhI0M/nPidyD1lZH80AWSiOmFMYJiGcg/8v/7n/6uLuj/6wYRlDOtE/+nzCdSlq0fSm09v/363HG8f70Du'
    '/o/4If+1zyP/491pL0QdwB/v9Atyvfvbm3fl/d/u3nxctsQks1h+evNJR3z3Mf6L/pmjiucf7/63t+/f/6Ri+Sf9Jp2dD7UsstmZ'
    '8T+/Q+TwuRCkzzyGLbnzRgm4rAclDDnvp7GEmmpjX0LQxZ8w+MjWgoPQJEKRHGqDYrMxhVwmLFCCiyEmZleKINyvm1vlh8+1UZYZ'
    '6hvl3acP79++1d/4X//z/+sLo/9VD8VPep/o5XX3c/3w5n35wz/eff6p8/Zan4RHfd/OQtAFIeAvbjzu0/hybgVGN98KSuitAOut'
    'cMkb9IPudfBmnTzcyH35my253chbTp80GPnRWKCJ6oQOxvC0iDAHRqziJRdwsQasNet/hQtkbAQoLnGkbEMl7yk4iE4vyVqFQ7NB'
    'd7JfNpf9AWgrg//0/m3Rec26dz++aYtiqIfkp/jmnQqifXj/k/7t5/cf33xS9fMfJ1HooM9HdT9JZ57P07LHNXd1MlTP8v7BPUdd'
    'z8IXo2fhj8bjpGdhJ5Q1POmQI97wdfAWvFyehO4HIy/otHBG5tNCia4f45d4g9fBmxiCC5XG/oBuS+6r0qj9MZ3K3AnV471Mez+R'
    'qSXW6rIOMpZasZUSISUB1V7IODTGoiExHL0NQWx1qs1ESS7bhKr706KNm9PwZwn8i97+7Y1OuErg3adl4rcH28TlPvr7FeZ0EQV9'
    'xcOJ5hdkQCpvMBuQ3AndBB5XNUgZvdggXrWFF7RBPMC8QZToB5tZVYURb/g6eAsS5OJCcj8APj9v87tVw8Q99qFeVA5QNyNR0LNi'
    'vrH0J8q82HVDjpjvA9iSe558Xx9UTj6JV2NOTHYNpVjvAlTRnSY5B/2RA67UCop+Q2nZcPMYiwNXkyS16uB08858XByBqz2uAslv'
    'f/n45l/q3af64aebSUE5shfHoteTcUvuKQXpfjCedU8lupU/H2M1N1M5sbESml5GEiuyTrFOL+v0V+tzZTQ65xWa/oaKSd8BuVCI'
    'tsTV9ySTy5Mfd/90Lh9vFq+femAWu3kER5vFVydSFV18aF5NPjLELbmreTU5znh1nHG39cLsKCnRcovNku6nUHKK1QAiVdMghqzW'
    'VcomNwyUTAhUdXsZyCQRS3Wp4L0lOrsU+TjzapmW3byIo8noWshDE2ReR7gl9zZ8dZLnZaJEt/Vo9vVCjamk4go30WHGgKlAs82Z'
    '6jA3dp5Trrm4YER/IaQUBJovzVBJ6O+J1g2nEydHGr7TtOxk+I4no+sZfCkZWfbMSp539B6uRUEzLxIl9GimWVVFyDVm74q3esa5'
    '1ggskgojNgPA4Am83khZz/KUXOIgmFRND1YildLiSQ/vQzZb1+Kf/+E/Tr7EST34UN/24NLdp/frwaRcffyx/qsOtTsU/zrJ8OP9'
    'JBzVG+rP8UP/XD4LYVUh/vwPW188Lrto/c75UNw8uPQiysOP7DylI+kuA9qSO3uOjxKvvSLeY0Ms36VYCb0fiBX8yzHqjIBfeAOv'
    'FwXB6apw9BlvZCZ/u72JcTBHAexO12OYvVjdn2/7QYzXeF5iDL85nvt7t+TuQa3RWJoLnhnIo6uu62fQEjvEii161cH1JLLGO9cq'
    'FM5Y+/ElrgbV4ApY5+8Bz0Et+zgN/BkmmsJDVxv14b0cN2I35WaNHzvRPQl+TSEY8GZst4NePm8hPIz64sTbSwqE6b6ZPUBKKG9B'
    'rvIG05p8Fbx5fym3cJMA5vxuPWP8HmcMzBkKaq76sAY4p0wF41ZrYsh7f/+W3NXO++qgavWFS/GxFIylZmdji2qD6zGu+mh1tRH2'
    'CEAIEUylqBZfYBNVQSU2aOCelqyrzzJOdrf2nldO4k0Y2eN+S+7s/gIyiwU6RV7NnGGUoBSdcxFVAdUqEFOiGui+YJMoKqRajPep'
    'craFq2mVo4ccmVVIKEHKPTo4meNknuZ8XDjey9QeMtonHeHh5oB+NrycQw/QTP5U6EQPn6+H3og3I6+Dtz7QUQiNt+TusUq/hvP6'
    '5eKXzdKIsTYIJUgs5KqlTHpuhRQbV5uilGq4h/xSJcglcCHBaGt1anlHuLfenYOV/smbhZ/kN+ax31i6B/J4v/EUR+Q1jjgd7Ws6'
    'lZ4jdKEh4wvSRrAvXj9rkThpI3bVIke89Y35GnhDMzhOMdxKz5KLdMGnH0Wz4wc60f8JJ3/e4JCdEylXcmft6spQGAlbwpRbJEPB'
    '+JrVeGHvTLeKVXfSmfYlFJvJRgIsjqP+hmTO4LP398SL1o9bd94RatUziEXgYZLFFKxHuclSnCoEkHdSJVFPTj1DEC0uWfw4mal4'
    'lff5/b8J3o3Hz3l3U4YWbMmdA46ABPMtpoS+xc86PVZMviWTev4ss1S96jGxJefBluBNtSbaUiVYL9bqn6CKQxNopSX92D3yyY9M'
    'cNyGXGZlHD+en/yX+ve7P+mj9+8+3v2b+PGu1Da96Zd3fRz/7Ze39Q78v/sj2Lv3bXL4/+9L7KY/+viHP+w8syMxy5RkCDdZ4i9x'
    'BpX0AxUaaUvurEIDLmNRwoDFOcCQUvQJa6EKKSRyNgaXhbuft7Bg5iYu+ZQzFHahSY49K1C3lK26TWK9P1cnITzO19uZvNBy37x7'
    '8+lNfKs76OOnD29y31sbhXn+wEhh/vyr/rDzLPU0fIMXemO4lXZF+3lHsHtHlqQgJT7TrkY8r94rOsJ7dXUwtrTQTIqRipBwiAjs'
    'bCE9352JBYFcq4nBJqocE4DLqoghJzUNdfHWeyY4uR43JWm7n+fPI5ngB1bMzXyqO/M8K5XYia5c0Jd4Pnw1DgeTKjgO4jgHYMKa'
    'eiTfhxZVvZAGUhkKkm9eWgu2VjKBmhFJsQWJptyTDWdHuHkxq/GKZCTQRZ5ZP6xlS+5bFAlL/o7thAGgWQX1EH20GCg3QwaSz3pi'
    '2BKJneilFfRuQwyFfa2M3hsVkQFp1fbcQrXL5tvMTheKPOI2W5h8pA9o86kHV9p8n8rRPqCr89hLs+QyU8acM2XM53lQXxzdHvLu'
    'C3AaZ19/YHhee5SiKxRDzdV2J7lVg9o1Vvs7h6I/DHpthCCheN+YSqXsDGT0ep72JVLul/LZE2fPmiPV6x3w8WvGjjNtFqE8w5oZ'
    'y2K8ZqbRbsmdz4GD1kW4si4OTa56beuB0JjbJpVjMPO9YH0qUqoHNWSyS1VFrVasbzFGn2v2tachu4oJQjTJeK7A/c5IVX+v8j2Y'
    'X51UvkMkbcxcR8V4MM9+snz9lvyVNvA+Gciz41o6YYDsfNlQpaS7U61QjGp1lVKMqH1soBZvSIWUXM2gsmg1pKSXeApZJASPGFXp'
    'qvdrue2J58MSkN1FVHTjaNAnf4qf/qrTWta8SL8+eJgweXpwG9fEVVkM1pI7z+tKntn8rpeF/X1ZHLcsFra35M4Oq8+cZ2puzLXp'
    'zrMa9lCtlMgOLFpsesO74K0NPutqaFR1/IWjy7ouiq1SBIunmPTap3trzBXP7qNivr8BmYsAfYPQv80G2KO2yJklR9r1gYKf77/o'
    'uWBJmI3e7C2wCz2PnnJNlBw5VCuc9O8BwRWfcwy6SLpjyDj9p5FbLnfe8nbQQfFwBXxbogCMl9pmER6MuyK05t1PGTyn3PThAjmf'
    'tjO5d4HZeDAMOVmprfoQfSKorBoed2SEUJkw6N8seIzBZKzJZ71NKLSkxoCa/zWpTuflVFBB5vdV8O2rIMjYVUBbcseSqcnioNXi'
    '0LEEO1scwQRP7Hx2hTlzMw6pRS+YLUfVFlCFrmZfzwly3RooElGqM9la27zHe3+2+ugm3gC6asftEd64OnFjGU4D2pK7m+6HCHE+'
    'zQdSPNx2v5n0KHgZXNRr7flMHhp31+OAl0JfzkUPWKkec23JBtR7liAjcMxNQtQzN8YAVZecZD2InWHUXwI9jV1296p+PEvgfZmW'
    'neBGxpPRs0FEbovrImFO60s2paIGUi5gQk+KNL4ZLGwgxpjUZuIIBkpoUFG1qNhMMbrZWjaVpaokAL8/XJeVuT+/Uz5GOwC25N4J'
    '3HN+bFebQ4eYme/DFjpaGHWQHCpejyhT+pYIFhFUACV7XyRgjA7RenbcxJMz0lAaM5XPvBnhcfPcuXy8YrF+amTJPOb7dp7IHldy'
    'D7Ly7bJjXkiuJv2oq2CKr1AnOuCGOeWhPuSNlp318nlTdXsI64cvhbcOjDf7tKUTRq+wU1wM+ErO6W3S+6YM38HuPT3504f6w3/+'
    '03+9+/HuoYb69DzApVYZplrlbcLIxeQsw9iSe5c7jQdTKCdVixxyCchiHeSWYq1BrRdnxUTEkKwvrgTVRqwPphW2NbbmqFqSVQ/p'
    '+HK9kPwwPeQ7E6EIy4X/1/ZSrxvl9um7+YApmLAvecW+VO2Bnb06BTN7v6EpCJ794JQDd7tTDtxOScxilhpAJT7D5LngeXnvltz9'
    '8BoOhn3utZjBq6kE2YTikfQk4yipxBxdqWgMWz2sippMzqth5aA/sr7mnGTNL5oOL+MOPrwOl4xq+zSKSbkt+X3EN92qyKpx4hbj'
    'JLClKNVY7uW1zMmV1LJ3HcKMEIIwCpHvmbhuhhQGihxssAzJGbqHc0XnkcJcZvIyQnF6cMMIxWhiRytjdoe6G2aiv6gJtG6UoXHK'
    'dpjIvTM0YCkH7ggPoAt/dk9EV6vHjvcGQi2jjcnE2HcGkVBgMpW4FcToorERVZMzWFxk42y15f7k7MLT8L89QwP2wtMeM9c144AD'
    'XGawtyrO05Pb7nVsz3EBmOMC4N2XeO7v3ZK7X6jDwWQGZ22F0IscCvc8Aj16KaOHynrb6NoLNun9CW3K+zctYNCN0pKe2zW1e/Lu'
    'dKHSsfV5zyCZh4cmLmy9HLARmTtBQCd6bfqKo/2Qs0lBeRWcISMP8KFWLMyZ3Bsfao23h6nQGOb1I73SPlSO2CqAiT74LOi6MkpT'
    '3nsl1WnAh9pU6WnKRompCTiAngsP92p+nwCizNOr7mkyix7fxAT4WhMT4OdoYrJigoZuc6lSeHae+PFxeaPCUbfTQWRO4REzh0cA'
    'wlWel/duyV2viKuDadkaqdaCbYm5Nb0DvGruLVcPbAOElKWmknK2xUeoKsDorbfc0Rkd1HuaK0bN0YGr55GMiDcjhYW25M6w3ADU'
    '4b+8d4HmQhOYIGNgdr4a02GL1F5yzagcfAbhHHTcITjBZtjVHCmoZqjWcsguhyjWt1r1KOVW86IizruKnnbsLJzvVFDzFYb1OnMP'
    'W15MpgS6LblJrdgHJng1FXqEzcEcZPYeQQJz9aYWDsS6L9RCLdVlKdlLavooWwwIxVMF9MmqYWsK6maSRGf9HB6EDw8I6z6MNp6N'
    'MLxmnfUHR3jYxrM5kuwy8i25d8DyIMnyFcl+S8DyBYgKeQDWTQuqwUrurnnBqnmBIVkKx1vhGKrkYvUiMnpquKonnEdLNko3hnMU'
    '7mDe5HupdvEulkYd3KBKMRTvz93mDDwWmNPAo4q1aVyTL/PL98EDHs9SbxD4cHNNJX038uDbCbNe9rgzqIM2TauUOtGzouELPM9Y'
    '+a+YZxPcaGfilrx+VOySjooLMlkn9NTA2UbyKdXYkm0FUyCpWfV8Va4EJdvI3trWek/IGHtD29JKZuCi+mRS1QWQ7ZyKyBNqjj+y'
    '20HfjVdwu2l6svtJe3XORvJcBrgldz5pn1l+Xy8k/p4EImL8AL+U4FY2adgNvMoQLI4WmrAVEK7yPL33Fk6kZ+MZGC4yrW6UjXTK'
    '8trr8jBzF2HqxDbTYcTzjbLLno3nEMyVSOqNgPXdEi/bA2b2sxCXN3iV5zVGuCF3Tw0djaUKJDWuSu/rYb2zQCxN1XaXoMVqs83S'
    'er91p5dPjo59A+uiJLEUwPaeH3IlEPwtltbhEw0e/agaSLbkzre3W1ohdkKPM5krwsR3zNqshmvLQhYyRL2tAzqHGb3jYDEqqRaS'
    '8d3Zl4Fiht7F27aOdXUv54KwRzYw6Ew+qbZHrtb2yDM4pcfz2K0Kz68WVBTcgl4Jbtv7eMTz8bCiw8EUNMTFNi6VanKswtEDIlWI'
    '0ZZApYkX0cPCcKrUnFjXdJlzMaqt2uDc6pTeZDK9CFzRa5IBGUXRjXtRHYDd2gG488anXugj3uCV8MYuDOAOgbfk3nCHswrEnTg3'
    'eEtNbwLoeXVBNwuRcVApeEhUMXs0kNWmy4GS7XUZsSZbHIBag2KtKViWXCxe3EmPxDuEJ6GEw1WUcJDnQAkfTmSvYAjPjQL/ALAc'
    '59CpTWSC6z1BPbjSoNc5kZMQi4RSSksuYfARbcrBB24ptxxS7+RaE4nnU/33hLv+WJm+BuR39Kcc2DCy919SHxrV/NbLY3YWnipu'
    'LhbsnNP8ktrHy5w70LP2p36i13mbb/xXzJuhV8AbBB8Gje5WgNKZ3Nk2XXxAMvmAgHluopjV7jFqBWUfnBVQsqNlqFHKDByiiw0F'
    '9ZqMUFiNVVVMk5PsiJHJMuZz2SI9wFf9+iHaubyM9W0eXIYH6eFHdp6O6SgMo+408KKOQliPwl6VbE8N5/3wKHwVvOlxMeqjaG5j'
    'sc7+/LBTS0OarzbuhOnZYyeX9UOel/duyZ09L+OxNMotGDVURUxLRMR6clCq0htKdlw9n9SOTYFT4mCsjZhMI5Ntz/v2VE9xE3qc'
    'h+u55tnAQD2Cl5RfSmYNAUxdnRBP++Yhb3OPt9fAW/AyakH4krpxKUMLb2HiDcyXeKPXwZsT8INm7KfyFNm7PKUjecwVHAxTBYdd'
    'SiKb2FiCnm1JT7CUkSi72iKAtUGkeidG5ZDASkZbgSvVUFz/RRYqsSyu+7kt9SPLU2QpT3mcjSkXRS2fP/lCGHk1RvFRL955xnub'
    'kUskSlwutZXcOfcRVySK3vIO0M12cdHryyKjZZeSK0b/L60KmdZc7/ckNXMHenGuOKsKMhkq1plmQivMeqlt0WPDobmP/uISXDVm'
    'XJ5cJNTND45IqBvP5kCwy8C35O6wwMcIlq8I9ptggb9/SRG6Bwk5YUED2JA770E7+xt7TXovo1vqI1pEi4aL82i95BQRSlMBgZ6v'
    'qeNH6d8M5Jpb5NpxnU1DlNYjfcm3oqLyJzCiI/vWLlABY8naBx7bBw++diSHwTfsPM0jkbvznLnfRf6bEPkKeLEhdz6QjxJ5uCLy'
    'rx7IL1yGqOb0V7F6D06FfQDPi7MrzRfPzrlo1fjPGBm9UZsfMJScgLPPHe6AcsyCUbzTq9c4a13Ra7iamuEUcns+eN5rEKtHpF6O'
    'J+2KQL+ErbtLNs0xArRXBPhN6TTfj0QgjErizA1r+M1eleLO4I/BY/BwiuD3Hk9kvsQ7LNlRK7nnavzqmAq2mHKwEBuwXgPV+lgq'
    'x9odkLbpUk1CmfRHpTFDbT5HML4nAEhJRRajbKqmNPjkasBnlIGQH0GJwpbcHUoUVrRNmEpsZniiht7q1mgIai21Bk5cApKKWUpw'
    'akEkVjG0zMmXRjE0dKH37MvQzKm7zowkCo9GEoVhqGh9cBlSNw8/svNsTKk7ZtDu+BYRyrXlr6GdFmXvD6aycyhLy9/eJ8xAgKu8'
    'L+/fknuDfHx5UFgkNtTVR3oH1Qi1qIYIoquQku+NPS13iz4ZgqxKJbhKas+H0EKrKaR0v2JnPWwCt3sd9zPKKeixMkL6wS25O9IP'
    'rmA4aMAGXEJEtlT9VK0tMNkK4KMvNuiZLOwK1WAloXf6K76I7TnsATAGaTk3PWfuwxnoBx8N9IO7Af2MeOv7gdy4bv4FpfHBsoaU'
    '6BeQnFIUB7zBK+FNwsXtChsEZvjc23x4r5O15pFhRuwys6EXsUTLSS9TL6pzI3lMwelOUsXbJQqhYiqQeqPzlmMous2CkJ4L+htZ'
    'LeUlTWPOdTu2DJ8fuDG/OTUuXE2Ne+z3uctP/WFngYxXzsTGltw9EfKYxcFXFsc3JUL+JqQN9DAkOcX01lQL2D0qZXuqS5jbPkze'
    'eTFLiWpjUSNV9emYyUf2aE1i39MqYoedJ+uaMVzRFFc6Ckpt+gObHEwJGe3U+/TYoNTUUGco5WW2LtX504PdrfurkzkW7DTyLblz'
    'B5SDBBuuCPbrOTQvQlIkA3j5FWV8JY91dQLOKG3cCYMU5goVJsoBxdrYorfGF6LCrrpU9VeD7fZvU5kVVtU89XI5sNgw5OBN91R0'
    'Bk7HCB4bKJYx5vjpyTelBxwCTn51dseSP1VmyaAya498kGMkvfpELyT9LfkgL1B0qEbche3AN0QUh53gSdSmYLA/BtdxAxa4Tf1B'
    'L7q2V3mf30+/Ad6BjP2ee0jyPNCQRFQDJuNclN5WjAKiycVna9AXT4nJVSCrd6zzBEyZslCyvU002997SD5x3q8skOeNc+22BvBX'
    'x7leiVCRmUftBHFL7mslrR177NyxB5YWbNHmItxyFnGRa7YlFEQslKOJyas9bA1gKwA5O6eSnn4zQZAUqPeBkTM6Bx5qJfUXjDTs'
    '5cEBGvaVKZvAP90o1GS25HM6wzpU6RJ5MWoK4LzOEve6XO+AhDLYHJvrIMMZsrjiuCbTcePVgIAoHCAlNZ+Sb3pZRa+GkpwLP82J'
    't4MO7T5jj97f07CeIQF3PLvdswUj4EPk25XSoH3kLG4+NQLUevT3nT51ROYBwVp3Ap83QLkQxMLWlty7z/J4MLmmGqHnzAq6nFoD'
    'CgET1xJLrXpdOhDbMneYlSlNvpiOAl7UFGriBO7p1FOWDkWk+E0vFzsE4nkxcZpJs6BVsyC1Fe0p3qyLaOACOKE5yd5oTt2mXavl'
    'lTDg7AKDb5vlqPa+/kvRx6g3jt46Wa8g4zq6BGO1phQqItkJRegditFzaqp+SDk3VMVHww7ITmhOV5nrE02XAWPstb0vZxH5ubaX'
    'O9FhDc8wB+4S5sD3Nn83yuTyOzXR6pjiM6A7dmKLnROCDAHd5YYdaHivPief98wg9yWe4ZXzLMO69ZcUpAc1RVb8J/qs1vEKhMdr'
    '4E2MjJJL8YbJpbDbUoV1qcJnvUYHLMMrZ/lhO/gpTHeji8fOPT2/Hdbdnts9fu6lOD3YXc2lvmN4BjPFqUnpqXfdYC7nXr2/z+Xj'
    '5zLQqNHI7MZcySNjqhOuqF9xRb1qLzyn1IHavinnEEKxUVXq2HpJZu/hXLEneLpqW3LBt2xCTtQsUU/2rt2Xy4zplC4IRzvm3dgb'
    '61Y/7SE9K4ZzNpLnMOyyM+jtIfLbtBd5rFf9e5IIcuAR1JPfkgcXaAVY+7lBh3nnWWfsSc/WSglkoktsVAqmSPTsqepvZ69C6/D3'
    'IWSJzKFVXzJkC1RiFvLn+p5waKf0GVvKX23edlA50HjSrgiUzgKdyL3bwhwjwLltz0CA39IX5juSCEIY6dA3NPeGsM6nJ7tPzqR5'
    '2lXz7IH4cxej0eSs6Nx4ADr31bGYUj1aXaws1QVnHZQkqTZn0GfxnA1WtfSsYfCVrC7Z5h2DIPuke87ckzenZqxof03p1ncjHSCx'
    'I28bbMnnTJNQA3XpCALQ0znmmJDLVSKpuJyrxvlWbOEMwML6t9oFiKFE/ZdJco+uFmhs9S53WAxg2bRFATjU9+/1BY/3/a+fush/'
    'XJ58PbVq4uwZIu9j+YwW0jL4Lbl3C51j1soppebhYvl6HutvVfpdaR5B198syw4vGvz9iqOTfwwemKxbj1DuTa3Nl3i/WW+T5+W9'
    '4+SPKlDdDStQ3W6Vja5XNnpLbq1snMDl/VXel/dvyf0rUL80KBeaat62WXB6vrkmHKKqOQlLNTY0iXo0qlHcmFxsRtQU5p67X5KP'
    'PpS2dA+Do3ulPrecwkX1IMBN4BOe7LyHuU0KdOLzyshL3mYE/NfL26uQG9JlextzkyD3PK/2wpx+Ms96PfzYO2MTL7z3+2IKOK0G'
    '4SXv/f2/Bd7DBZYGTX0Nza1436vHH/5oaE7qxE5sTcxLnvGGfQ2fh2eUgAOnq/Fbcm98fbMgdClhgN1SMASJSst6mCYu3ug9WY3e'
    'lULBx1qsLQ3Eh1YrEkKHsQEpXArH1nzjKGcMUTp1VPt2fP2nJJeHC+f5zlPy53dED9PEJ2h/xFu1IwxLDvYePM9JtdKJ7qQ0V3le'
    '3rsl916Sw7EQVckZKrZM2RffvWse0TGGqKvPq2Kq+mlIKVtDIbpgomkxm/6/1LuJb5YkmscuycMnOlykF8ywgLfKXu7fa3cCqLcr'
    'QL2detyc+3te8Ly8d0vu7P0ZjwVLwpwL6lLBkLlFl730chnI1qm140hiA+DcCyP1fIucigQ1lYo4UBvinulcLHIFjOCrTt7nmXQZ'
    '4Bnfrkei7NeJb8Lm9is2t++1Ae4LPB9eiTscC+RCBZPkVKxB1fMaYFbbE6ljFhpXoYM4tGIyBoe2lBRNQL19TQTfk9GXaAIviThP'
    'WmjPM+moB/IAewLtltx5dy8BC54CFqDq1Qy+51rB2ncuxN5tQfUX6sj4aKcyot6ytufbVvG+Nr1FKoOqOvqDYgX0zrl3J8ABtI9y'
    '7U5MPtq1e/rUCNjgK+n5m+Jq+wyu3fGEd3ONceQ+uVm2HO+WLbcC68EErKcrHr7Ac3/va+a5/32kE4ctuTP+NqFbVShn9Dycwwm2'
    '1/qAMJMp2TaIHdZVd3eg0FIKAaHva2Y1apJFKqYFU13KKaWaKIb7U/SSDnVjysiUPDcKwzGY/vzgv9S/3/1JP/P+3ce7fxM/3pXa'
    'piH88q4P8L/98rbegf93f1Qd4n27+/TXeve/L6dDf/TxDw++7U/x01/18/2j//0//+ePD58ekG00ltxoFS2ztCV3tzKOWUX+yir6'
    'Jivj92WxXRbIAwCVHqp6QUU9sLSL6YTppb/ny0JG/ky8oT9zL2jgQAu6Rug1A6panTRC9gPT40aQ1HNhmt2tMM2uhWm2g+r4L/A8'
    'vXdL7l4BOBoLBNcbEhBm4ti7DItt0oxt1XS4VZ/JlSQmFdWQvZqNLmafvZFY9bdS82cwRjwN/xEVgEdPNOElhoM7A6K5A5oKglkK'
    '9norMONxqTlOHVqpWsO2iLQYGvmcDeacYlGbQ406KVVNDiYS17EMbam5pZZDrdX4cwcBeDQimhvhEF5N6t984DKJxH0NKW0D67EP'
    'iuGV+ZzOShgZE3xDxXqns3KK9YX+zxwTWGN+Z1SDAe/9/e43wTvIIAaEdCveVQXaCwE7LIV1neh3JH+J5/7eLbl3fsh4MHoxpOxR'
    '/6jiMntXTHS+78xeI8HiolSqpjckS8ljqax6srQKCDZLzPfWwbkzJx1nUj2PZIJqMxcVXHwrfBHqJyDushPpxy5DtfGdVRV8cuZR'
    '/0nvU+yvMr8MYEvueb1+fVCM0VibkppghSplNdbAqi0WWzQRPURVvaSkxkZIkg3d26c3r9NbFmvSXTSX6tAkxCfmZT+zFEQG3VP9'
    'DQ9E3G/b+RXq3E9Q5+C+xDOeD0Q84kAcD4ZVFVYzX9XiDOQjSBW0mIIN1jTV7YrxtkAsRX9ka40t+2hFnwYs2Xt3TwFOmWyAhx6I'
    'zyCZIACDqr1Tixe3d4uXXq82+yGlE2pS2/lwLj5AMI31NGBp3gYj1VbqLW1j6ABatTnbXGohJ2MaAgY9KRq3YCIxrvimkzL72B4v'
    '7kHDlm9LV3DjBjBuHsHRjvurE6lmlB+AqcmpPd1M7h4qkzWCpOOhMMfErYBaRhkEShQhqhbBQ07Q2584HaYe5SQWmZuNIUdwragO'
    'wsbUvhL4ZEf1SJk8uomxPKmJsYzx4fFR37fzRHal0n0vKBsnFAzcyYhAs0QmlOjJkvAFnm+EvvFsPOvmhYGDDW4Y279SbzlFr0et'
    'sZYHu7i0T982dGmfnx4BLL0NiuvdZL8goR6rNb9L6DuRkOoLNAA9mT2eK7lnRJNmn6bam4GcCwvgxeTcdPMB3vNAehFNQc41R3Ec'
    'aoKOBVt9SV0NJTJd53S1cgd4zxBzS+C6iqrX5D0ttspDx+2+WucyOVcwQow9CCNk9gevUyb+hL50KccO4S43Alz5nuZGQrh0ZfSG'
    'NLcCowH96p2MaAfmx26qGCczeEz/yTabfsg8/EaY7xCVF2ES+GETcti54Nb2QcGM2+x6ZYOqpXPfGDXTXDLROh+zi3ovBlSzLZoK'
    'Qe0560qx1gqhtd1FbHOWaKCqCe6kSeQT1LV7dL0tPKneFq7W27pnKKO9Oo8qUwd827IBWho3pqgWeFMDvFpIyTtIgiYU9i2W6nLm'
    'nBuIVHKYVcqcnCkBkvVqshnO/nsoGzinOxyPSX91IvWUAhnFjZG35J5qCM91u0syH6vJIrPbMXmpgB0TyUbnGqNrtqZU1Aw3zrli'
    'q3M2SCaLzkPU3WmCr2yszSaqyE8oCufRH4Oi0JMj+WqrGR424p0fHAH7Mp7MkWCXkW/J3dNiDxHs6dZ4KNlvyov9/kUVRr4wdwoW'
    'zOTBIFhuTbB1PcHWw5xgi8wpOyOBfbJUyJpgYsBmmv6Snre1n7i5YHIWyarkStXHVUqSasXZUy7Hhp3DGsBcAyh3h0GNjyftikDh'
    'LFA4ICX9KAHCFQF+S/eW70ciauO60WkYtuSRW6x37pvTD20ndGxLDrvlTIadjcI25Wg7KkwvNqsADZra10YPxmRaUbMqWJFcvAq1'
    'SBCXGrj1dLQzssbR995RTQ2HUzMW21xFc6NqtRtMgZeLhLNTJ9qV3Pks4bWmqzca6E1w5jyNEHsMBIt36CulBJZCdcE2Ex3lCqSX'
    'O5SUOOaUDaTqrS5RCzX7ml2aLSle0ooedY3Dbsj+V3hTfekyZdR8Fs+RvW0bC4tKbnsURfzsGYSYoHlP6INTK9RR7xol2eQiUiJL'
    'Eev6s55bCWg7EhxxU31Jzw9fDJ6RQe1jw1ETl2Mn7fzkiKTs8SxMLS34+bFa3YpS6rqv1ixAksXUDD3QbqLpMFwZ2SNUMLH5mHUT'
    'xO4uYFVds0T9VatrSwpUfcgEn2G1un2xWge5+keCuI6mpwcbBrnH59toJnc+ojwu21gJ3cZhzndv6KJekap/5ZZQsk8gzourxgPr'
    'dQquNBvJcqmp6pkUsm6rKqS8lO4YOqOr+UeXVazAVo9zDFzCYX3+5Locv+37/ZNG5X/Vu3cW6lRXS1/V3na3Zd1qrPVikSXGXawr'
    'scUUEUoJ6JyJYmLsidYACYKqaHqHMzus1VeMhDkiNynd86Jn9NZJ8bg7cCnb+Vqi9ElHeUIt6NUm19/WP/dU7Ht4LehIMuNVcqru'
    'WsmDzeiDVk24smqOaA72+yL70iJDsCM8RXhBmGY9DWvBopu6EpM/YZo95A2moMOr4E3YjeEoNuSxh8ODLuCAS7anrVwlJvHUMWG9'
    'C2qaxFotEeth0YoRh5lSjMlKttEUSTFQa9lLS0Lnfp0Pm6bv7mO72mP9yNbpn8E+IJ5b3rF7FNbGLhIcDqYkTxal4wahyzEGDtSR'
    'LcFktQ+i9YlLEdU+rcuGeytl7haFjY1ZfwPuSeQK2MarFqGu+MskJb4hFtFejcgmN4Nf3Qx+i712yfPy3i25d9PS8WBqacnalFuJ'
    'zutBkiSyd+S95VQyUUt6AKklZdTyBcqJQqypFap6IwYXdJ0Knz06/lDX4zNIJsggX4N+wFsl8kw9hXfK1wjQE798YI8052uE+Xo8'
    'eVvgITICTW5691Ku/t7jEWXp8ahqswoWT2rNgDd4rbzh65Gbg0uf7M2wfSZnHJmdqsvN3I+UO9GRovCcYzPqUEVmS+5cFjEeS0tA'
    'wK3pdWA9I7iAFCxWTtXZEFKzUaRVguqNxxwglOibhGhThzTJ2zby9Ohal8MnmnCQ9xJOhfAzuXsyU1hzcIIB1Rhnh79OF0WDjU0F'
    '3QI5tgCgKj8TNLUGUG9o1wzUFjMzVBGxlpmdqo+EIFtAoUfV8S9cPimZ6SqAjAnPksw0mkjdPRfgDFPVJbyk05Dn9uXYiX5Dn07D'
    'C94mcIJXwZuIGdWb3g50u9ey7FbVaNaqRrMt8x7yPNXQbMj9601Hg2klSKoGoOQowYkpkHNmVfg5qqZfWS+B1pqAE0C1GSB0C7dI'
    '9jb22Ne93bSVAHNwvenhklGdir8TF9/JBbdL6bdRU3y+Dk0n1BQ/w0Rd8Ly8d0vuuhqvDgYgGO8ic+YcUrSUxbgOyZyzXo6ltV6C'
    'YpxpYDDENgVlpbLFjiMkodyTn8xS010aBIc2CnkGyUBwl3VC5ibt+55cH2LM0l9Xib7V1kYTI946sPFr4E3E4wVv4VapS7M7we3m'
    'TvCrF8F/jvk84Fnf618xz6im0aBZyk3W8JXh3GpLTXguA/jBbhPeaHbUDhwgwRNsyWdrKGlP5qmdzFPwS8s6KB0+JpcY0HryvWLD'
    'Jo4Wqm3iqWfXGd96Twwm77lFqlyqTy3YEpxbgKTtpem9e7zYLpb640qc7IV9v4aB+4CPL3G6Mu3j9TGNdkvuW7Z21BIIV5bA1xNN'
    'XplMSejBERTOkIcruatqe9nvfCn1KKVhNB3/HXwgtCmhii1Cbg0yRIjsLKTis6/S9LOuCnKjQLY1apNufKXr++5xQnfVxzNOPAxn'
    'gMYdupeMp28kS3eei5XcGQb8IFGG3yU5Su8NW3L39N7n3ZTflN77HcqmI9hduFF9F8htEAhPhc87qPEAvcsA97T6GQKj/6Tr819k'
    'fq39XsmdEQi/MiiXHaZCrmTmFqjXrrtgUjAWq4gNvZtL81Ak1yLeZROZg0HvrKSsNtn9jAFH5wr8pyAQPqsUiMD9NroHGXRf4Pl5'
    'uwetY1Elk6P3ybIed6bV0BcdsqqcCVswwOypoikEvnkrhlWUCXwsUazL1t6T9d9196DzpKu+6Efb3d7qrIOdXIr4o5EZHx870cMt'
    'X+QZFmfvSu6qEV8djJViPVSqIJJJ1xVYdpxs6t2gORoPUziUXK/SwOKlu18o5x4y7RiZ9/YUMJFTlPwIReqZJCMYRpU0tCW/mqU6'
    'nQtPTo+aQbK5E7rTaUG+7JCXLYv+lxzayqD6W6hck95AkKzaXC52RPVA1PSWMpFsrZBaLXpKJNjU/uFngKT//Zf0sf6PX9QG1Gk/'
    'zezd2/f5n3/45WflUM3Djyq4f3v3/sObf3rzLr7V3/vp5/ju7z+8rf9S384bLcd3XXwqzg/1Y/+u+PbtOil/vPsP/6onTOl6WH1X'
    'ztP1w8efa+4rQz+n85DfvH0zr5H4rtzFn/R3a/lh+nrd3bne/e3Nu/Je5/BBYtpD/NNN/Q4dU+Q8FtFTT0aYcai5E8ZY5/6S+/eN'
    'FuPC75a8uI++x8UXfl98L27xuUF+jtogtCV3tg/R0GwPKTH9b159mVHXUZCQvfRizxoKUK415RaU1nWI2UdXVdtO4Jzv6ES1OMRo'
    'Q8qbhrB4Gv632oedy7F9iMuTX9kN4yrXvfDch0GG1CkVLcy+0Ae+ll+1/cNSINEJQ7LUwBvOmNG0AmoHlehVFQgoucUMThxYalhT'
    '866RT1jUYmexMegDI85I9ae2iuFBfcSL3v4yt2Yfb//wwIX74MER58JYdptve1SCF63diyf9yPjTsTBalHRelDR00H+PixLw91X5'
    'alalDX6Ac3mjxkHzu3dqFrHiktCES7IxnEUeXtBTlhTKSyoWnAG5oBNdoPaUu/GQtynX51XwFsSbZ21g1hEl53zdnrUVeoPCuQCM'
    'oImg85ZrE0wRrPUNUrFWAqnGxV7P1F4YWb0L0CDY6D3ZmKILqcS6QeZ8bP+yi15iD/uUXX3wa3fU1dm4Ihk4S2YJVO4X3Hp+yeye'
    'f/BdClKNUhkU5cDN+sjtVT0HJ88VLJ4rDl/iGW7lR30enoP3ZgDhjn5L7g3hjn69jr0BWepZvClJh+cCeWghF+uNrcxVuFpTnS26'
    'NSGbWgLF3oWTjRQbmivgvRradTaR6Dz6o5Dbu0t/BGi0zNYV3HL0B+GWDydzLNhp5Fty3xggIKyQ/B1ZySxF0q5kqFkPrpbUtrDs'
    'cmt9KXbYqaTGBLiq9kc1zIWwtkxVjyHOKXOqNjq6Zwsn0UJ4WiTmpYhNwMMILdDdDop/n/RTe8IjsRMeCdKpcPeS5w0S5ErurM6N'
    'x8JJbMJoVFso0fgcnE2uFkpmriPUBVtKqI6IYoPKGAqUxBh1vQehegKyfCw20vNMNDJdogv5Ew7jTP7KZiW7qHTeLAwoYfSwnxkI'
    'U7mmGsGltBo5WjEdToT1dsixdExHahISt9ioh+T01GlVhGI2ZMDgPS/C8Q8w7HZX6fAC63FV3eYn5xYuHzc9XM6Pv4pdNLHw6xvL'
    '7CygKysMzysM90b6PG6xuCuL5euW2W9W+r313Kh3mtywjxjv1kdM1j5i8hkYyoDnufPf6+W5o/sOwiprYHIm9y48x9Vxib3OcVlz'
    'QGrL5xZ0RALgmuub3JUIvVDeZLX9Q8PeOoWqRAuZnCnSfdzgnEtOTims1PuKPbLuHHGnJMYxb73yxF/Ugvf9xC+nXro3uJ6oTvSS'
    'o5MPb8DbjfbN/rxdJAbbBdngRglZYSeHEf1oVsOlE/1MOHv+HvJMZzAKGoBR7OBJuDIYkZwgol76AsWmXnMbE5cAkiP2nGcqFNW8'
    'JBvEd0xLRyElTi7WIsbGe5K13fWxvoTnkQyJu8zJtKcu8zO5c3scO/df4k5083bOsAg19hlPxAQWAthqyGdJjZtaKS14sskSOaqu'
    'qOaWWqjSMon+yc6WcIqC2wc5wLtjdem3XslKscuTiwDg/OCIAOB4MkeCXUa+JXeGFD5IsHPe5kCyX8fVfxGiQuEwyAU3YUvuuwdV'
    'aVuhOieEmaWYrB98NZbqA5foTEbbbDPFBpt8sLUVPT0FIQj6aFPIElxK0fjqoukNo9c9uB39UTCP5irM4ygwcnpwCP7jcDLHgp1G'
    'viX3Tng/RrD+imC/pQX4C5CUiLtUUKin9t+mjFvIjbp1hC2584lgZuwi7oSeTjxPjZgIxrTovZ7TTCmpEtS8l6Y2ExUpHHLIkKv3'
    'PX8VCorNeog3Dh6tC2bNi54n6sgTwS7AOZ9V4b559+bTm/hWv+njpw9vcn/H9qAPV5uBjBfm/OCIhTme/fFKmEa+JffGyztmJbgr'
    'K+HrR8grFa0qQ6MG2bwld2+QzWsjY+4+ujnSm1SSuZUUK5RkOBVnfHRqJrUeZ2O1iJoqYsm43HyKudRo1GbCYoU5hpjv7bkvNh8c'
    'XTV8LUxn+FpXaH3w7998qFlf8vHBz4/pFj2a5LHAJ4625N5R12ME7q4I/KuRrJctQefRjfxSN0J7cd7AwK0C8h0Nh242nBAegjVN'
    '0DwEN4K569Use0H78VztCZ3o9XV0dgrjw8ROfS+ZV8wzIo46PANuye8hcC2AS+RdzxVwi/3PnAPZAizZildFKnnwWFLoNV+2ppBq'
    'bR2c3rsUvcmNkmBuIsWF4uvSccluWT7m8u3BBbx2dPcn39Tcpv/izWKTYwmMlhAv83mzxJrfzlQvTG3JvfN5jtl4c93bYOd9Sz7P'
    'q5Uv8CW4LN1kKz09H3YtLZCptEDsqRkBX4LIhtfCmx2FkxFfUsgVYQm5Ygdbsf4LvPUkvNfAm3g/bB3JW/LZEB+nxnKyNpaTKZt8'
    'rrhQk7MaW73YhiwmZLJARgqamgzkYLNVYzMHqRGib7F2bO5I1StVrDvhK0ivUjq4CTg/vmWAPOpTO09uT+h2F1AGp7SVldz5Vl0S'
    'U+yamDIXzmXMuVAG9gApxOSDit0EsVxd9tlQ1p+RZ4elBS4hcGtZbKmFGlRbT3mYj0y6mcGHH4vruPnUZUIePf77Np/aIVV3PMMj'
    'aS98bMmXIW35XdpfkzaiGyYRmtslEZqdbPoOijEnEcJUeYQbLfLS0TaBA92IZ79TQh/2VBzsQYHe/51nkCT9Sc/JsVeZnwfw22Be'
    'HurYsqQ1bshdUTlA53uFTmMDsmD0BajVWqFmTM0uUTWeUNQCTA5TjDUWm3z0NlKM0Fg1GfABTY619Q/GTcgc+DXhH1zmEJ/j9sA7'
    'YLZcFclTU2gXW1UmWxUxyIr/dLnUJu625K5YG0cttU2H3N/X2ve51mjQCQRuVQkwQ0DybmCnNIGdMntewU7ps+abF8zPA3C/BeaD'
    'GmqDwnEjt+u2ZHivnj649vSZnBJnc/SC5/m9r5lnZBsGCbc36uY5Zcsa3Mk0CbSYJkro/8R+gefpvVtyZ0N0PBbATORaqJBrUWPT'
    '2BAoJu4ZjsQhA8RIlJuU4LM3LulPXeHIbGMRF08VfeEByuq3tJM4fJ5DCKO6b7Mln7Ory1LLb6dafiBYsG6jZS7NioFUjeouIakG'
    'E2KJzWQJwRkfSjYlqbbvairGlgQ99bTkTKm6+3PmExyLqiJPsvXnT329ArO7EB//9cunvinyM/j6UTrczvKdQFBG3Zhvkib71ACI'
    'g7lusRPdkX7GxBrwBq+Dt97deJCZc6NUGJwxAXfyKHTQut4Px7BdMKdn9Dp/anhwwTwux/xvgHlRVWwAM4u8JXcuA/JziJE7oYuR'
    '51qzgHrvJvQ1gq0d2sVHSr3O1gF7D5kylYyp9FOnBv2wZyxZDBXTa/PKycjdDP/bYWZxbEYuT46o9hnPwp/feRhVGtyi/PN+NBpc'
    'qjFvU/dgCUfBTNyS30Eql+r+Zrkgu68+yIzQkzuuRCsRfOQOF6daTzKhuCRohKN1rpjccnSSkuqlutbJZRf1/4YZHKzRkEkVODrI'
    'iSP0uOXBr80q+fzbzpAWnyFaXHn6h52FMl5VNwr//xZnn4K5OPDsDwgvqK+7XZIvOtHDVObUsz6YQW/UV8vbjXrW78+b3scjO+aW'
    'YI52pzRpB/bHDoBoya+2gO39tuHcz5fHQJYbck+N8KtjqpapF59W76qHSB5KdgGJkdA74VSiUO5OEWe9l1xjk9JszcVb4yItwGqb'
    'hu6PhwB8XhmEi3LTqZHpLdbf0/uCz9XsvRdp31tgTvgml7zNlfsvnzcS9oOSErQ3LK/Ya82yXuLBObuEw6D/oPOOJ0AJ9oPSErRb'
    'cudz48tjMuhCCBUpS3POpCo+m4YOgiVDkEpJlDFmIlXDKatRGaJq5tK3ZUllbVEAiwifdG48rwyQ8cKa7+/HF7S3AJe9BdOdDCcf'
    '1QVv9tXwpivy4twAWLADV/K6E+Jgh37PwzK9ULlrSB6XhCzD57QcFHZAWI2Egg59yKmHyoMShkPtqQtTZyDwenkTNgqp5NyAevaf'
    'K2mBcJk2y5EVwtNpPMRXPD154iyv3zKJ6unf8mQBGVzQ9syU+BtWTWK0siZOt+SuyFpXBxNTFpZmE5leJKznWGh9EWX2JicqqdlC'
    '0UaDNXoGcCUaPZaJW9Z1hf6e7HIedxNKfl8l+62S7oUcdV6VG3pA+EkRMh4m7y5Pvi129jx5+rzm6bMuCJGrghhWa+zqjh+Pxfso'
    'NnEpofVub1zRF8AUXG8/qFpUZuMw2FBMdT7W3mc5q9KUXTQ5suc1vXGa0ceBcv1mha936qjE6lbIjydo+x3SoAS9Hj/Oq6G7pEHJ'
    'BDR4RjthgUFjMzRb8jnrlJZ0e57T7W0wc1GoOCguN1Vy1FYQ4dJCiS0Ho0aGcBCuCdAUm/WB3m2pP89BDQs1KHyOJxRCOtiF3+sO'
    'vw0POjwp2eFKoYQ8tlBiZ/mMFhKfi1/4gOKX49aKv7JWvp5z9LvwPxM+uC351f7R3+ViCL8fHMeuHTHBjG4g3JI7HxxhbSYcejPh'
    '4BcoNahOSou1puKwtgbVq+kkxltnqUQyESiJlECqy1cxOfe89KrrJBGLP6dChAcNn7+eCkEPYo/fVt5KFxVAnz/5BuUrPOrFO894'
    'DwIzjFzusCWfTf9YfMwwJXoas3U206KfW1UcAaUFNjW53hvJN2moC0BVrMwp9NRJk601Ysihrp+AkYUxZmOjvYcFpvrE5GGgbN1P'
    '9/hKzPVTlxAY5lHfd5QAxkumD9vcquvpq51odjyCO7Vbcne4U7uiMtqeyz+72EPovqgUSxCv5q/ezqznc8kpcbU1WZ+dntwRY2s2'
    'Zqt7kKoJFrPTtcLnVtwPRn8M3Km9CnFpr0Fc2qMgLkeTORbsNPItuTt66SGCnc/QgWS/Cb70JYjKyyBJyvCW3LeUY4VZtBPMIoRl'
    'LFUPh1RZiquYC3lUcTUbKLBeayjGmphCEIoJmzNgvDgCFacLgJ5q3mTRmcd5pk7glI+tCzBXPVPmcIfT1Xnsyb8mfCdpLr8qleKc'
    'QqHCX9P1R7zB6+ANxZnRfchb8uZpsNNhwuthoguPgl+iha5xkexjwVQgc1NbJpSQoQGLlZKd6rLGN+psk5iEavKQU0MZsyU8Vc6b'
    'B9XMe1+i0wuunb+8UyLm6duGiZjnp4ec9SPxXMFom3q+vk6Mths52ffnTR4mMrvJAL9N+tFs04/aq8yjOqKLytQPjdZ+aGRAyJyR'
    '8nHUd4625O5950ZjMZVa8c5R1T+TajGpBT3oGpZcCrNPVmrvZ1wkdmtdL2yKNXlQ7caFaty9P/tq6NF9574fgfDDW4zPOtlKHumr'
    'HWhGs8UZq3Cv03c+xNpdbanoxYM1B4c5Q3ewmJwM1ZBt9iQqLqgYk7Oqiuqtdu+uaJgH1GkcAnF+dWbGUpuV/5ulCjz7FNjLbIk+'
    'DLsld+/HZNfGNLZXiy0L1ZoMvd80ETrXyBi1VJ3qTJxS9uBFGG2UUEliUTNWqpqrkoPuQkxqvm77Mdmj+zHZq11+7LUuP/aoLj+j'
    'yRwLdhr5lty9H9Mhgg1XBPtN/Zi+f0nBBUr0q8W82zYLHhVpv6hiIloLbmhKgA3n4noa1WC/Ct6QAwxKrE9OTr+/99r4xYvXCdW5'
    'Fo+PapwJ0DofQ1cxY+BWwXOzqmG6fogERxxNcMI2NxcCq10Ye6amoZJTPIHs+qO913zNJTq9etiX0R/lEr0ymVcEy2fB8v7e66ME'
    '668I9uvXxcuQlDVu1AsbbwUwJhcxtycDfy/XBc7XBYj7As+361nxXDwTjNJnX0xf8KlO2q510rqG4ZzaNOCtO+Xca+ANrYwA+G+F'
    'aNvXDO0EwSJLn/dOdH0uXOV5ee+W3LUy4+pgpLbsGTwBJqhArUbp6UU+W70q0Cd93HvVNiISMbYlRGcwsjOgv+fvyZuzDm6PrMx4'
    'FsnYwBc7zS2NbVdyT4XN/mjc3HjWdsKAsfNYkJtzaVIms6OMJqv1RxgsCVdvsVYbK+i/Ek0irwIyxatZyN6mYFI5wae7Q/st8wRY'
    'Gca46W7Yxff04ADHzZXJvCJYOAsW9u63fJxg5102kOzXg9gvQlR67ly2ZLkNKuqTzxdnFghzN9dK4tkg9KNmHG5L7ny+ELq1eYDq'
    '/EZmp0IsVQicKZacay1bx/oslhZNrrVZn8Sw7YtRvGE9/G2SmMlYDyUkF7btGdyhOYRBX3C1L4PbAej56hyN5LUMaEvu3k7jEHmt'
    '6awXAvumfho3kwCgs6OmcGFL7t0Keca4pBnjEtyadsvWgUXbmghVdj44lugqQTPVO8ik2lRSs1oPbiuCNkfnCmHMyeYkm1bIR2KY'
    'zk3nwtUWcuFKI93+YNRI95De11cneSzwiaMtuXcr5GMEHq4I/FtaIb9gCfZrb5RbdqNmv8PhwA2HEy5aIc8oDzdyzlwWvD/dOWPc'
    'j12TExeWlnT6k89QSgbMnzAJaIBJ8Gs399cHlaL3Nkui5IQRSvVksqDkVgy27k10UVXxorZvNFClpmQDVVekYMcqXS7WDWrH46Fi'
    'nlsKwDTaEXQ7DH7AvTD4Z/+g6UTn2V3leX7va+ZZryQ/aLj6ohLeAi44aWEqGsdT4040Q5w+flFJvrwm+XYr0Zsv8Qavg7cgXkYW'
    'VXg5WE66FN0MZKBEL1b6Im/0gpJnv8Cb3nQ84s29IN5o5Y2+xhtNa1JeAW9I4kbqJr2gBsdgaNVj+gWH9gu8dUXqVfDGHka4TLAl'
    '905Imw06ng06MbMPJrkspSaWGktzEm0KHl2InLhApcKEJrtUShaLZFvmRGxDb/fCVpXkTabho8zghcknNVyGa2lsCMcXml+Zx75e'
    '4fKc8bfqtTG3TfU7RdT8zDN2oq9j+BLP/b1bcu9Y53gwmF2oDjAYrs7Z6HNsJninBp2wF7LGYimpWLX0qiPfOmZoSSRYs6NC95b8'
    'Fuz6MBfeM0lGzd1BKdnGhjVu99JqI0vZkREDdikIwMbQaktosbJtQWrHZ6gNIEPOzbuWoZXgWkcky9FKKsRBbOPu/wZ7T+eeanJw'
    'VZhxV6rCvozq94edp2wKTVwWtbilwdpK7toN1zizAGgoYcj50/2QEuoNVjKTzwVyctwqQ7ZokjEZOdaYXe/3zTYi+JKyDdQbmer/'
    'ywkX6zz419GhFJbubxdZcHDRFu7zJ98AmeIefvfOQt1822MKfsAsLU6V0PUaaGlxOlqrcF6rMGwG+D2u1bXW//fF+noXqwUeYLvg'
    'rfrxTpC3O2ExWlisRCU+z3u61NJOULsruauWdnUwEoL3pXm0JZPvUA/F+1pDiSlHl11xanvUagIjq56WQWzxHUjN+Giq5Xvr+NRK'
    'BtyRWtqzSEZkZAeunUhncm87cAHK6oQBB7LUrqaCXNUABHH6X4nGZgCVh+sxUZdtUmHVHINykK1REtQspAItV1EtPzzo2vUoO5Ce'
    'YgeSuWoH0nPYgcN5vCrTUz2yzPmfzwZ4epTA5YrADygxfFXrQ5eHGZzGt+g3ej5Id7JGAy7WqBK9wo2+wPPpAF/Jvf0E48FA8uhz'
    'FGeSI+bEDTCllqL+1ZrEPlSy4krQywmLmqFROBvuPYxr9Lbcs7Gbfp6H+gmeQTKqG1xGNW7T2ebJtyvP1Qe0dE2Rk3fcXrbcnNvV'
    'vHzeSPXjQa366SjGGdR031r1z0+0pW7eh1yY0RN7x715nLGxUXYeMkabRPdbZR2TSzlhyJXJdkggCFSsr64960WCV68EvFIB/RWU'
    '050ncyzYuWnlhjxWJdxLsP7XqITfv6REBkGcc5rQTO5c/9uzfyYTuSf9gMjsL4w5QZBgvd5gzFwguNyazXqJBUNeReOzyipzJU6S'
    'qZaUjWWoQX9b8LQFzcMkp93LSdy4/8oyW/u4EcZzNPXgkFERYdiS+6bn65Jxa2GcjkXlM1vBKXix5MFGEwsnFYQxzYH+ExIhu2Kg'
    'Nd9clGZLcyVSza14kGR9C23bL/hYeclFy/JN842wS3L4lTkay2sa0JbcG5nyeeX1LciUNxSAhUtIoDU4tpI717Ms8Sw7xbNA7NLx'
    'L5JYB7F56H1VWpUQM0MLvXOoR1f1ZBOrJ1yszQdhSCCmRR/IR/Sm3q9VVceG9riH9uBKEZYf5nafHhxRhDWezCuCxbNgce8s/OME'
    'C3RFsl/fWi9CVKpkhEH2krG3yW8fDoduNhxkubzS4VRxN5MHQ83BenpCPz3Bzfm5jGh9SaGJ7rWIRu232HIzVlLzuWcgqL7c9NrA'
    'DN4IFX2D6s+qTecgqq/d87kT5rFXPF+r/5pefQwC23jOxvKcBrgld77yD5KfXJHf18+l70kgQQ/di/3ubgI28vRyXcClXBdwQhVZ'
    'E3pHvMHr4A3tw4KItQz+Ri0s1PDzwysVHjz4Nuj1+VNHXMRg/HwRK2F6Ws65KQhdOErd7zP6K2ZUDacRGvULKmBRpTXMmH7Y3ep6'
    'O5wQmx+i1K+wvrdrvEpP6r1JY3eNPEsQzcja00t6Ty9/Bha6mF8+g+zwAGRnl36qo7FQyaq0iKktxOZaA5IEtXegdw3VTKm2NP1Q'
    'LZSNi8n2XshA6MSK2j4l32/cGebR/VRflUxBRsBmN2qTimcPyx6QRCuGpEwYkucEkAue8eyTxJNPcmewqOFgjDNRrTtfM6VKaoa7'
    'lNHnXIvJUdeytIi2Bgk5dGSlZtn4plpycRgppXBPAc+HOBwNFnW4ZKxcYgrZs/fRHuB9NPNtYjthgMzct4FrbBwspYBcMMdWU/BW'
    '/wIOHJLNkzu/RUlZXBLyqZXiDKYyYQ+tXQOnvl/4SB+JfeBk/DZdwl51WppHfd/OE9nty4fYbNOxB35L7l5349d6EVV/eOkKTUAF'
    'Sk0VvXUJbMnVukYueygRDGVrSe+IgIn0SmlcChGHoM9cC7nlbd2Nf2xwDfy14Br4XWIwV5ie/EcjLRpwSx7psKHZZlJVFDDwDHGx'
    'Gk9rGypX9KiLuu1Yd1NtYNkX8Vg9BfZRiRSlFHGG9caHVGuj6l3zOTlf6B5kaycfCQUDU5X513J07VqOfgBMyGqAnubw1KZpJGg4'
    'Cxp+F/TrEDSqKjvQYF5SObseUrxYcVM5O9gv8NaPz9fAmwSgEcQCvigYAlxhCPqaRPwSb7eBQNmdN2TPo5CH35I7ZzHADNPRvcN6'
    'mzs368iqlhCqQkq9gUX1lYwqoIWQ1K6Vhi2kAq0X15Ti1aYoXvRpoYTZVpPbZyEOf3SIw1/1qA+daMuDY1zto8kcC3Ya+ZbcPfZx'
    'iGDdFcF+U+zjJUgKcIRaSFvy2TL8bQc8ptnDqoQBEJ6bB0DszQHQ9rT91rK3Em1wnkX1nCS9My5grK6qFVJDd1DF0ko1/Zui83bd'
    'n/3GoyPVm2nCntJ+Feh4W/LK5E4uSBodxLIld3ZBMs7lwp0wOpQ5i9tLdCGnlkqMIalkC0lIplJIltBZyk33ZwFWE9PoPm3O5hqj'
    'zw16Q/l6vzli5HEuSFhqpEa7cpfiqatM//kdeW9GmbJhS+6cfynoVt9lBzINuOCsV8N6KHLT3RXRCvU6glg8eTRV1MTo5fgOuVIK'
    'DcXa2Ns+dcSN1lyspyJpOTjYP3CcbZy/YVgiOj84oqvHeDJHgh0mau7s3D9EsP6KYL/Juf/9SwrB88h9jVvySCsf5vTsCT4BLa6x'
    'uI2VDx6kgq86cIvoU4TCwfYa49ysHujVNWEu+hPUXyq2JY7Bpqi/15KLsHhP4ej6BDwshR1OKeynqTkZ7wP5zRUGtwu5PPcUkJGR'
    'o+olQdbB7G3t6kmHUbR4lbfZWeVfAW8iJIOeyGRu1BNZly7B1Ra8Yaem4qdve9hU/Fd08XVrF1/XO/meE09JRm2Vw+/zu8v8drNy'
    'dO7wrcqA3dSOfheM4p4VA4bZkszpaXN6jG7gq8wvA9iSO4e9vzYqKj5VTGJUDbChFBMcSpNQfWJq1rLqdg2iVy2iqvqAORjVU23M'
    'plnvZSl324TvDysLfk5JBXAXro7VyLUHGLn2VPJqp5JX8HYO3ppiskSbfQtJsDUovrWk1q2qpOxLSDkCQ5TgMmOMrYpzPoI1KXPO'
    'co6Cz9XNjwHs71x+U/jlwhz+BkeGHdnKq4uDH/V9O8/4VeHjljw2u/6oxeCurIXdXVu/1aWD4F4bgjedwoIj3uB18CYBw6CjKr4k'
    '1N3F1MPF1Psib3Ou2svnDRnCwBttXlIX0aVw204dLfSPcyHwJW9zM5FXwJv1btQQCbbkg1P8zbs3n97Et3oRffz04U3u5/gu+TJi'
    'ujLoxDta8mX0J91XMh/qiDEAx+Kzq6V3TUkTJkKJDTxV58UnIFshhB5fshRqlBZN7S3Mcsj2/qTey6FdK+eeTI+Flt586qIv06nF'
    '8dfE8PT0Giboc89kRRbAFwKd+uCuLpTbtrp+lXPsBw0c+VRVMJPPGdhlpFW7ot73fEZ3ltgMl2ZK0w0pHdG0pBodFDVcY+RgJUOF'
    '6iklm7Le8SWmVDNX60USllNgl0+odMdpv/QkpZbG/dHmJ1/TphdJPYv2O5LPaCEtg9+Su5vNh6wVd2WtfJPV/JsUfvBiBlf6jTyn'
    'tF+1RFcRJn+R8wFpQTSVyV90AkW6ZJ7OHjE6IHb79UEh5+RKyE2k+RZTBdOq07Uecw/8Ne9MyRyy6ijCWTCJ80XVGG7RBgp8Pyf3'
    '0ja54PE94J5VCsiWBuWHxm/JnSPoYS7z5E50gEG7tmUw7Il7xUr0FJpjE03u/Wsdq2oolFPj4Ng6AU4SWs3ofSVU9RGgncrjwmn0'
    '31rnQEsd7A4pKGPeepDEjmNQG/LUtPPXawQ9nDB3g5NOGAKc/bQ6rT63wKkk20yIGXNOaBOmSA2ctaahC83GTK42C1i96NJOEIx1'
    'hdO9nOaMzKEYfGEMy7pM1kWL09PPdw8JX53LgVSXcW/JPTfPcWIFuCLXr++elyAo5IfVk1OJ2WxnruSOnQd6ZdnsN7OdMCaE2TQO'
    'EvTWaCX3FjTZF5OiHhbVkdUnNYlKy6rmVY0vSDZFE3pHR2O4o4qZ6k/9u0+jfxWdBxYRDBMEzUUQ6/MnRyT0joX35N4t87fx/G3E'
    'PHcYEHmo+69qz0vyMvrVy+g/K6wY8PaiIBv+f/betDly40oX/isV/jA33gjoKvdl3k+9qSW5tVjdtmWHIyYSQKJZFlnVU2RJav/6'
    'e04mtiokmgBYRZp0e0alFIvMg1xw9vOcT6wNmAlJwXma/vDkcJ6mgao04Z9wIQtTEFtZzNcuDOh2ec4K5eGCMSmJkdoWjhbcIuaq'
    'QEi0Shkl4L+Uc6XjhevDeZpzw3maST6bxjTTA52sh/9pxlqUmPNkNaZ3P3UT6ifvD0+sPZ/pJuiRm3B7JcZTPVqdwuyloj889UtO'
    'Rf0sVBBiaqljwA7yuSBWVdpqkytZFbmTLOe6ZFZKUVRlbjnXOXwhqQYDVWq4AJIIMJMq33vJqTjzS07F2MHGb06A2Zvco9R51Q/U'
    'H576VTzPeZmR85ryKj7gASimBy+M+YKS/vDEL4xpXl4TrG0Vn0Vw2ONCeEUZr5xysMVVZYwvQAPjwmDgqpDe86Iy1kupqBKFpUUp'
    '8kJawvNMtrC155SKKgCLkumss/cHwwM2R1MdfXEO1pne/ZGbwLqbwOJNOOmreKabQNXIVbjdLH6aZ0uZEKngonhMen3kYAwHmBnR'
    'QLGl1obBxqewNmukTvQcJvbheg4nXo/bgsLDN2F5O1zStMMNmQuf3qiGefGWed1TZLd+UEoQ4kewgyeOr2Qu4do6RSrLifc6N6j1'
    'C+98mTtqsXjTsKJipTFWC+kEzRXBqmwNTJFQa7LYWonVt+G8zYz/PY8cXhWbqNJm5oEAOwP4/8Bv2v78DJgtDbA/j8D+wvQK2FOJ'
    'vYQ+nrogCtc8ro0ipijmKLWJcDrR0IqwJ7E2qye4qE6qgulGZdCoMlAZ61eM5poCRzLC6ryspM4ZdbSytvRCGi8x0MdJzhxwJM4I'
    'phuQAnQv4UhZldZl5sitMl0Do6cL66XXhnVzKlmeQh9KrNJB1tQd8A5Ng3doOnkzsmZaJyI2w9MjUaYexpWmLIQpcgOSj5S88MZV'
    '3BTw/y6n8I/wRFdwn5iCu2TB+DbclXDRcm+k4mUmmOxqWc1Zpd99nIwgA0GGVWoP1jDSDkBaFmcycFq3t+NBse2hb5JU24t4mM3w'
    '1C2L0w/jER+FmJwT4HW2VJaqQgu8jLkofZUbgXan5RUBBc0zbBSUO12KwltSeFZkos0f4Wd11d3XyfBkU8XHVCBAY+YWFoPz+KY1'
    'qPnDtcXEt6ewNm1T3bQOoEdO3Mypho0QATYC1dAakcFWYKUIUihpSgFKQkFJbp2VVFpjDC8L+J4Jr60sq0JIIpWS0gvm86oqyqa1'
    'xpnRUlrNPFVdpepvBuD76lwtOkY2M3WwvUZjzfDEuaZnOti2RHMmXMqjOSolByFmWGDPQCenRiwC+8HWWQpgHlNW67i0KI2oSiOr'
    '3NOcCl4YDf/nRKkw69FJDWq6qJQpBFWM+FwwwypelQR0MKu6pq3krO6Fek+SDT/jF+do+JnesnB8KUgEqvrDe+zsTkAVqdMSqQL7'
    'XscHdQY049wKa5iWFF5BZxzDClmMFRs4Q8dpUVnnQFUpCk1JUZaKVJiTqbUiZRtZ7JZ2pkAJaA5qfmd3e4e/OodTPX0KI9eFd9eF'
    'f74un69Ld11AQqqEyfMQcYrWWuGn6j5P6iYLNjq0+KfWzGu4yWZ4aqdD+mGM0tQzJXJXGloxloMwFN44bPpacHgBWAmCj3NmVVFi'
    '6wUqrNLOOixtZLbIuOS108GiJ/WcZt49nAzjx8hg4qGao53B0WIbR0uIGOjRNT9Ys7uHXfNTP2dpWKJWvkkqoidPAsP67+hOETTk'
    '14BFFH1LDBPCjXbKCSJyS0lZwbPawnoFqnhRSOAxQhupBeGWOCkLxpiQ1BjCpal6SMtnTgKT6RykerdO04c6vUfp8woP1B+eGhn7'
    'POdlRs5rCjL2Qx6AHJaTPVg8RAcAATY0CtsvbqsRtRN+cXGxiiIsFqsoRNJWQnxqE+nnTZy5iXyI1mm/YLI/PLGzi7NYNYKDrs4X'
    'g52+oFzowuSGVFxzpa1QGAHVFlmEkV6ZShQCeIblhZI2r7xTjAgED26L8Hn99JN9XWGR03PQen8wBF2I30yoqeZHNP+/E28n2AJK'
    'mwlw3udDfWZWx2vmckQ3EawoKyuEKqniwsD/qKVOKCmqnNKi5M6AzVx6UhA4X2YKlYONTF3O66b0d4J9HgF3/sQXY0DRZ8eDbvct'
    'dYZJ5O7zQbKf7gzVfUOyP84Tp1alcN6IeSAEaUwPSbj7e18c7Vp8XH0C3kZxl2pUNRhgeaf+1C7Rz7t0uEtMMzIop2d1gKEZntQb'
    'BJoGi85+1DPQBWIjhIJkhee2yo20ZaWFFLnh2mnruHGFEKUriHJ5ZfPcwpfeeOy0UQru4ZcqTer2jKp7/LM5Rc1RfGSCe7PeyuV/'
    'dVfbfHTbU3dA9XNY4/DenObnux/q8/1YdD+sGfqrHiir9w7wLS1qi0b+J9ok9eTaKHsCa2PCkFRarn4oUKShoFu8ZhHlGccB+hm7'
    'VClDEqlSRPeHJ8ZCSj9LYZxj3hZUV44rIj0Du1VyyZyRjFdekFLYkle5pSXzyLjAtJW8Uh74WEXyTLAmJCPGzNUpGEj3senyuGdn'
    'zFJ4TGChpAELJQgWqruIMtUsFaAQD+isFyd01lPCGeIx9rz2XJBPLb7xo/K2dPvUCcqffCqS24qaglEvcmvy0pXKlmVlZAFvlGOV'
    'd1oxLwxDxzC8c6D0WqUqW5XMSVfaTHR5yvbcecr3dlJKmWFLYPIgns7Fi6zbr7Km/WqTRkmloAkHJBH94Zn7WluCPMJoC7py5IHw'
    'E8zRrcHxq8oWucA+Fi4n2MPYCkcN/LsCPYpyqaRlhlLhqaJlAVqrMDy3+GeWlMR1OL3Nus7W1xrlSdoxGb+Z0PHa3vaLi7fZsDoD'
    'Ggb98GHyCmCi8AMVlj3OXVRK6gkgKyfNR67rfUSs9yERsSkzwLkE8HGVF1VRukorVZYFmHVU5wjtKYQEvUkSbrATrFA6BytPMl0U'
    'jDirbVajug5qmU4ehKVpiJUEWsvRF+dIck1v5sjB0u5gT16kdr6DtSMHe3u09lGcFBN8iFJMOpRicgaUYtGA5ooAmmvr3kuqMkqU'
    'RBSGlpblBfeew7cSe0qDdDJWeyOIzCtXgUgjnNKKaustpdzkzDdxdbROZqIUk0UoxWQUpZhMRSkW94NSnN5wzE6l/IHhzhpkHetE'
    'pTSWAlTel6YqjWVOUSwhUMIVFpRmUchCS2kJ9ZXJqaagalfOGOMKfv9wZxMgrk68R6nzundQslOd191ByR7iAJRKFwXzx2vV9Jpp'
    'pot/n8LaGE1GKB+PNYrdmqI3HHvHYCYjI59aG31EHuNPrI2rIZI8GjT0URVjmqYYE0u6aYtNIo0dCF9Ze4ab4Tm9CMDnJKkj+gh/'
    'AK9NjOiXlWSuJJaDViWdVbZgvqqUr5jWUhFRaqKEYUqbwvpCK15poQtCrKdlqYqyZe7das6UNzESqa537zxZEOk9S52n7Dz98gye'
    '/vs/v9tTl/7NDoSnuxf3hifWbm1MN5c4QLdbfSBEgo0J2lvltM0V8AAtyoKCJZNbrzzChPNS5Z4pqySXhnBPuHUu96A7SUla2EJ7'
    'VhdxvSenUaHSO4EsXZJ/k6KSxf4sTSM2Dw5QzTCtuEqu7TFhKo2uTcHdTAUyHwyBRJ4IAwhENGlENIkiunNj2RQWBDH94albVKcf'
    'RlDhFS+5lhzsK2DsIqfG597R3DHB4V+sqoqCc2HhBxUnypaqosDfGalUkWfc0jYKc2Y8nHs4GSqsTgAePFDJUq9z3dAZdNwjb/DF'
    '2J+kUOLw52dBiavfehwQ2snQ1D7jWvnnfT7pPhsiU3UhDyEUUXOC920QvKUPBSnKAkc5SVIZ9n6MegnFAUIatYmniTUHuv3hSXn9'
    '6MMUwK8pB5XdasOsAx0+VxqbqxbwEw3fSFp454pcc22F91rpEn6uRG6o1Mpm3AR3KA28/owK4/2cjLLSpoJJ+gFjrSrFadL9Ej4d'
    'R5BDt+WZggC0iQAFnyYlZnR740rs5+09xfYyqWTCzcPU8fYGd0B/d0FLe4ydn8I6Gm8AUyeC6QTruTbyGdjPnNo7tW2i9Ww4IAgO'
    'jrP9/6vedHPylEVUUhUOcFLzP0VsAzU8+3pT+sMT9iYb3SdPMe1TCm8rI4nCyLfUzpaEGit4JQzHjDXQsJkqitxSMC6oRDPDUuc5'
    'o1nTTi3e3M839OncUD6lt8r9dbuuIzgiRHBAS467SAwXWudClsJTXXj4V8mKXJYcYYhUQTkhhS4FXNcSlCeRMwq/XehCi9KWtCu0'
    'DUge503POSzEnyqjRgr7274vZ5Ze6W0fuR8dTMYw4fYUiTznuQJq5ApMSeR5UmcKNj1JBQwfJoOcpnobLVszAfua0y+tVNzawD4J'
    '/oAQ3vdnHK2ddvyOnqGX1K3PJJy2XgoQyIUrbZ5X1BWls85QL8G8I5bCDfZciKICYa18oRmRGq50bq3PSQ3/QtBpwemyooz7PQKQ'
    'kCzlurMP6FI6UVNueMlY494JGeP0E2sOdPvD01YBjTyLJ5Zo4AiCSy4lMFRQBEnpubEV/DfPQSW02inhfFGUuZGGFboyTjJqDKsU'
    'OqJtvypt0YW7n02n6jhHW9dOtd7wlPE/hS3UaswSijyXk/gsRJuq4E6CtKo8s6UtMNexVHnuTcVzmuellw5+RKhU8GJXzPLSVBUR'
    'ugKhJ9pKYnpW9SUC+rB0x2dafzNAiolfnAMAJr2ZqYNVfbEehyduzn2mgzUjB3t7C4xHcVLANfRA1QhgN48nN4k3SAWIVYMAiW3e'
    '1fHaaI099ATWBnZnCk+c9oenllesCY0wDI3UOZe5EDY3hbcG1B4BL5qmmqF8ktqxohQgqIwFMQbXTrmKK44pY9Iz5UrCNWhEsqsP'
    'p/MgliT8xaTCGTXxF8MznEroJXcrfXL4gA8WGH2cuygSPnGKO/R4EjloxBjmOIiFgU0VvrSpHOansTZjH5hv8YhS4Z2QHtRsmReg'
    'GjiHNSrG85yrwsM/VOYmB25WaUIEp0bmuizhZyR3FWWa1FXBi/iWqt+PedUs7E5/dfp3kasOQcH+W3G0p7e/NBX/e6geiRiwO5Fp'
    '2LTT4nU7LWY/tWagq57ymiVjieY/RPWHp+4RQ1TdBAVBPCmNIX1RWuYdBcbImCmFEz4vBGyvZ14akYuCFU4oMI+qnCmqDKtEUcrK'
    'Mql8abloK/2ap5/TIoaosRYx5CSRoNFVp06gfqL+8PMJnPEErLEkBR1rHi7XIdWYKD7UKZD2ECmiQWOyAY1JdO2QE3vB/zP3AuSB'
    'SiBTUfKIdGJJayAQ7FkEfF+Mri0kbj+RtfFUW0z+mBLuOa/XxrFXH+vq+gZrC2gsnDyFtUmpUmj+sj+8z7QDWaPe4wCV45hnTog3'
    'uRNWa50XrNQ0t04Q77ilJYV/ECOtoqokucoJ0T7XXOjSFpLKgrNctGkH8qytGWIvATnfumj+aiiL2dF87RdyFqETn8fIxWHdxWGH'
    'F+ff+W6oz3fj3HejXkZ/eGL1+kx3w4zcjdvV6/+Yw2bDqvdHpdVQbLVCY8I7DaV2tJWOibU9Km30E2ujQg8TygwCp/WGJ2bghtTd'
    'JmAAeynq5h2Ga2pcrgtfkMpYmTvrqXee0hyhomxObFE5UXpDtSh56RC6kfncwGcJT1n3XWqf/mw5hXj603t99P5g+Dqbo6mOvjgH'
    '8lR690duAutuAos34aTs+kw3oe1ZfHwVbs8tfJpny9hxIXRIiGTmAUFq04hrgy8WY8/UjbMYCY2zwNppubkkCeyZ/8y9UEQmesMQ'
    '2x+eGfNE1Z2dcUAo13W8ToBSVtoqZ7Li8FaXHG50LjTcbVshMIOgFROKsQr0tMq7SnhvSufLnBQF9aLFzFBnBfOv9+hM/V2SOxOy'
    'F4cFeLzOXmyGJ5bZPOIKCxzAs9Rt6Cj3SjGh84phK2JtdclyrY3KXV5RomlRUuZAgyaI3qjySpu8MiWh3toyVyYTrY+I2rOi2tpB'
    'fmPXRYva5OsXvzgHm05vZupg6yfvD0/dSe08B9uK4OOTndJK7REclVVqeFTiccGmiDpRFwd9m2C4Nh6e+ymsjXLOUyhfqj88MQiR'
    'JE0lGka/jI42cwHii3ANr5fBjuNGM+mrnBBDy9KWurBgqeV5Lrn0ghXwL0cq6nNJvSwxfyXrnLPkvD3eR0KE9W6dBp4ovUcj58W6'
    '8zp5jPZ859V2IJTzgrQPfAJMJsr9mw58zfA+YwN1OxsR2tlgGUQMoRfMqpKYMjdeC8MqUmBmNzw9kXBwXpUe20NxahTV2gvONOXE'
    '00oaKUom29jAeZsLxhwivSjzSE/CMFbzp2//6hwSNH1cqXtVr7I//HyvPt+r2fdK8USzIU76wxN3kicxripwABoxifktcC2cldpX'
    'wnGdm8I6xrRmlUQMdWZKSUtKjRZEuLI0poK/lT7XzIPSjQXrzdVpn/5sHk0BBNKNzEn9zbHKXX9xhoMd2cz0wYYn7w9P3HL+TAdL'
    '6cjJ3u6gfBxHJZhNhP4I6w/HH+jMvJ5/iZhG8PRWcmZjBhL+BFuhiFrXK0xeUOK8U0rmlS8dlrJUWlsnDGh/hao8cYR4CQKBgcoH'
    'Z58TIagW0hYsi9VjvFv02Rwb8pZ+71MDjneZ5U7H0N/91i+ZuD+4UvFwVTNPc5dTvRAfAoW9Aaqk5FRAlZzWQJU8xDt5x5lIAkI0'
    'uqua4ckhRJMPAyoj2InAgqTSTjnmtWOGlCVn8F8OrEadW8or4R3RJbOFNQ5U0ZwL4UlumM+46SBED7EGzgAhev6TYTpVA0DlowLr'
    '1Q1YL+azkjafdbC2WOP1JNYmjmG9Y1a56A/vzY7Dwhoi6sIaIggzqm63SouKaOEsZ0wZTypZGqIKC3ZakTNSOlVp56tC5YbnhXFF'
    'JQrqS1bqAvQ5rXuFq2duwKfq1nFz7ThydhCb0c1N3wLkGurhqsIe4S4qm67z5of1ku3TfQkbFpWJJiJ78gJKTqNK7EuZg75csrwk'
    'hua5EQ5kqSgEKTQoy8qCMiy0l5UUutCegrgSLNdgCOWCmUJmPFVAyUz21+3ul3y7/WW1v/bX8Gb8frO62bmAK1c6eF+q9SWOg0F0'
    'vbq58Ct4027wu4BqB2/P6lUDcJJv95sS7g9i3yGIMJwVvnCr4sJt3vvVzbZDvfutJtvBz/VKCVPQxGwMzZidC5s4eQ7pS4JPLnqL'
    'OG9R7cnuhPp8J05zJ6zROhGyeyDQ2NCS/DQFlAyLhHhMXoEBKoz6E2sOdPvDUzefTj+MoUwJolzlBPy84kap0mnOtDGUF8qBjqEU'
    'aE3A6ctC5zkT8JuKK4fhIleUGWekzkSyR90ST67K38PJMCqHbcENLuyBbiMbtIZcvGaDryHhVioh4+JNWDwjn1p8I8Ob4SnZ8+0P'
    'hRAsJXog4WY6boE/G+PKykstKLOCcyMEA6Ysc0Zt6SttmdZayFwx6+F0a38lqw9xEcbYfZ+CPM5sC/6XNrpEY3TptMUHTURCYkRC'
    '2DqVJie5FbkxGot2vRAlhRNgnhUsrzwtFeHeUUUcIlSVVZlLQmCgq0IqJXjZdryVRwGkKcUHSyJCzV8N82/krPlOvI+YoUxU4s2i'
    '5hF1RGTRisZUfH3Qw36wtvC2PYm1MSVT+HBtW1t9hra2uukYqiP8e1SUrC49A/3TakMrywtaOCFEYaq8FLIgpZUG2aQ1nDpM7ZDK'
    'y9xgPFcqnVdd46/zthVXYy2p1dlaUo9uWWhEOtSzgpHxMN1HTgbCAldWxUJmhoPDTlNqAJTXWiksAVx0Ct0y/TAWxIbj8EAl0Qb0'
    'RO+9BbOqsEVRUplTS4Sm2lbKWkZNkRNgmlVpfWFBjHOhMhmvKgsgn+Sc3Ufu42QYtyzRFjc2RmyGJ1avNNGobXBhVM3h4CddtBDE'
    'Ohi3rBTAKErDVO4cZ5gWYpWhpSeV4NZybkuiuLQGuIoptSklyz3+qa2jhaEMYaTB4xT1atjr8VwLDk1K5MDeZA9SDN/yoxNZNYjM'
    'HQMUMMCrZz+xZqD7xNfM2L/JOS9fG2vWhn2MKG3AKlJre6DzPPnaGNfDIif7qNbWgFWwCFYBVnWbng4aUAIv3zwgXr4+IVh7B9Le'
    '9dwarjnSfcprZjzhSBEPUqi3+A7XjV5YaPRCiKbt+3m8tqYR6BNYm+Qk0TyxKcOLw3sLCEtsBVbX0VGsoyMyCreQwkuKihSukiRn'
    'RV5W3jqfW1VVngntiaxA+eRSOM+8Nc47R4Thec6qXJatLUjPXGEYmr/NdOOoDsLwnM6a0c0NiTYpDZHZRytdie0lTqU0waexNilk'
    'ogNVzGJshvfagYo0bXNARnBpojeQ6yqnwhgBRomxxCldFJZ5bhCe3OtcUQw65qTQ1lLltQAz2mimrZMFGNP9DlTk3B2oCFvUJZFN'
    'SsvXs6Y/8SmMXBfaXRf6+bp8vi6966KoHugGATy6N7xP3UAxUtfYMwTE09HBgeEkX3EHH3kuifaWlVIoW8iyEL7EnLCqUl6YnIlc'
    'SVXSSkkjC5GXyjvOe+gDjJwZfYDNvy7tXw3rCdX8+dq/OgsGQvJ8UhepXlZ/eOK60DPdlRYs7/iy3FoM8p97+lLzBBtpTQx1aGKc'
    'AaeEyojGUeWVyV2ZK05VxTkvCyK5LgSXvsSz5pwIp4nJS1owVhJWMU0rX2juNckrlan7xSlJnzqxyViTuidok2YzUwebAKQ59Wt9'
    'noM1Iwc75a1+DCclVKolCzGPqW0JadqWEHUA+CdUqiXL011bkDxPYG1cq0d9bnXorRdx6yCaFaOpvvUPk+EZW3Xbk7XqZk2r7lCg'
    '1lkLjKb6tdv+8KRR+NGH4aWX1JTCKGzEyUpniJZeWZZ7oiulCJiLShvHSOEKWzFBc6lylQtsKZTD28VNB//BzmpC3svJSKMToCBt'
    'LbY6Qy02q9PCcEColrE5mCHWC+zp6CriSUUKVmCDZV2Kwgsw52XlHZj0BeeyYFbmIMl96Sta5coQqVWmjqAnppdiq6OC66mICCMF'
    '3PU3txvxIQB7fiN+ZMP/sVHEqJSL0PaH9+vzMY23IWQu1a114ZLSHMSd01o5r41VQhWwmJJLrFMgimtRYvm+0bw0lmCWbIFAHvDm'
    '2oL3fT7m3D4fO/8i6Xtw8I9u7sgtoN0toJ9vwX/ALWCKHqqTnLTFSO3w1CJaMdG4BQS2HIyMibrcwCELkMCiFJVnlbKlKhRmcnIl'
    'SVUIENCVQJQGp5TIK80KA9KgLLmpctczxZk4a3InEhiacab74lA2NPt4Ipme3r3EUTaP2h+e2qlynpOU/+EnKaxKtdUS/eE5YX1D'
    'a6eAsEKksrLp8RTawUcHglRVpbmUnpIcnjYX1LsqB+ZsvTe2ZLryihZw6px4xVjhnReVEZUUcB9IkWdUtS142HkRVnjyKLtvJvTS'
    'tbf94rm2eeQu8O4u8GMGffdq0fOfvRg5++n5so/+XDkVqfQK/qhSEFiTghAwWpt6kcTSkGs+haVZxW0CnO7xeKgQ9KtWBXFAQBNU'
    'bY7icG2Pyvv2qbUZQ9MBp0eytmNPPrF2fG1NoOXxr03Z4zvJHldDQtrmQdM6D1qYNn/9eG30ceV4f2Jt6HFK1F89BJzM4pZLjKrY'
    'cgkGodviJ9aGpVlPYW1UCZ7KzZeP6E7WzVJo3SzlE2ujjwvE6xNrY8BMUvKN9YenjbtTFnPEJA6ACI8AvoWrsEm0MFKCyutloS3l'
    'llTwUyO44awsWJ6DVmwsKQtniMyxPLJQoiSu9GXjxUEJxuahrKs6N+7ulvnI2kJdKUnhtKhHpNuS2GSK4aAPBJdaG6KN6CewNqb0'
    'QP+L2ey94Wl9Y02GtgwZ2vBIsQJWCsWFKHJDmS2ctMJKIkowo0PmCeWVZyxnnBmiGVjLyimvmfDGlMJrX7LGN9Z/+rP5xtJpLIH0'
    'SZoWjOxR+rxihn5veGp+dp7zUiPnNYWfPeABCJtEgLYPiOB7kqB4wNO16JtRTBnaAOtiDYUUn1h8eID+8MQIbbc8FFMux/J0q4kv'
    'SpFTA/9ZmdyXquSlKRUWqLscBKmQuc0dk8KzooIvpdZc0BoCiNeHuNTndq+nkAKUxbrFJwkoG8Fyn8TaRKLmnDwUlGjQYag6FWJa'
    'U9Ns0d4lXZOdxJoD3f7w5Fh2yYfxJREqlwJYgiqLghJZcZ4zwqkkijtGrXQFd1aXOfyJ0zn8updlWcDfIKBYJrTsF5yfF8vuHk7G'
    'knQs7fG8aZaxfgxBdm9aYm38iaxN0lRh2gOB24dWHwtqNbDefdlfifN03RGsLuwS2E1Wq645C5UpYcs+b/c5tpsKolOtG+QjsrUp'
    'rW1tGBBCVdeSfrg2hBJ/CmsDQ4+mIQx7w1ODg5IapwEGXcMFXXknykqBRaoVUdYZKa3mCgQ9E7x0ufMeftt7hW1dtfXUeUWVKzzV'
    'Xiib8Ra1yhxlxs0EBY1Nrk8BCppcKIo5MWwcJb9g8qF4E/2CqfPwC1k3ZMIBbEAHhzrYgvox+sNzZgSNPlthjKElL4jUREpHvZHG'
    'l4IUiKRcFHADiYGLakufM/i34qWlPodxhaaeb5ukwYGeU9l8mFOj0qaiKfZRRRxME3GoX8qaEw7WFqMpT2Ft1DKdiDo/pggfmnt1'
    'ZBYlWJfBMlxbiDo/ibUxZY6UWI3vG1OPJxOCUhbb+OIAQdPatpiDtQUw2ae7NvlE1sasTuGdPuKaxp6NTobtbgiC1D4a/wPwfhK9'
    'zyTi1H5qbeivok9hbSwR6aWP6NwYyu7a30fD2rgYXVvA5X0Sa6NCD7NzHhOAPkWQ+VrnYlHnUqNro48LQP8TawODnKf8tPYR+Wl5'
    'DBRzHMDaBBtdW8zrfxJro1Y+rexTWOPo2mjIZHoSa+M21a6TiEelc4lG5xIHTX9TawNWyZ/C2qRM4UUQ3R/eY0kyiN+6gh4GhDIR'
    'lXpniaDeq8L4quJS5lVlhPEce40Qhy0PCqNlLpwV1AjiitJRq1TlXVHgOOu8doftAk5ekizrlgKDlvL06ItpwGPxr87RiD690ekb'
    'EVbVH54yb+98hy4/H/rUQxfD3l8odh8TxFEtmngUTcS2LqrB2vjjKvr4xNqokTLh7iaPybVISe1apCGBmI+vjdbtXp7A2pgRqdi/'
    'eVTJaKJJRhMHZmVqbfRprI1JQQbCkT+iOymAl9QyAAfASnibtkWFTp0be0wwd7RObYJBqI5rTS+hE+bJk1gbFZSnQmePyqykjVkZ'
    '3abja4uhs6ewtmS6nahhJprhmRMM0slSNteV4VSWiNKnqpIUhdXSMEWNw7SIwlqmac6lKpmteEWL0J1NYx2aZJVoEwzEmYFG/gMz'
    'BMOS+8P7xAs703Xhn6/LSa6LYsOYLm9zfOLwxGhUnNUdULEci3ETUQRK6qkprVWkylVVsEoZZsFWrxzjwgrFC4eZ8tg9tnDaKuG5'
    'MYroghtXVaBNyVa1YudFkLPpgi8ZLcT5p378Vy3cND/+4g5bTpstp/3Gg8Pjl109QTO8z0YRn+/Gv83dAI1RpPRf+Zhc2LRxYcdC'
    '/M7eTKyNPo21McmGKREKwyq94b290xyr9KNhjAO4Zywaj4xSYwtaCFaUrHIK3lpZUklKxYUnuaME7GZXUVoKTS0rae41Izm876WT'
    'IofXuIXjIuKsyHS4YTPfXBEf69yIoaObG1ofklTA9xFBIBALC4zFYmDBgjXUBdcSa+NPZG3wk0RizANVdEfP8EkqBygCadaJFyq0'
    'smStEppaM33aa6aciwm4zqeEdic6wnchujA2kSTxWRzNhaRUe0bLUtnCeeacKGSulQVTTOuylAZsNODXpXeFBdZsmDaKgyJWcGDE'
    'uoNunteCZSEWM003bmm/mdaf7X7gnFMbHupEWAq54DG5u2UsA+I4QLWxdVEN1hbRDJ7C2ij8L1Hf86gyaetOsg2OQ6cSJ9aG9YBP'
    'YW2M8aQrXz8qdV826r48zBLmqcz1J7E2rihN1ag+hCs/JpvztMx+ILQdXesph96B9osJgpAezXAmQUibvgY0AL6SFkJ9sKO9FhDN'
    '8LSJQyPPwgmm7HlFiXfYnZwWhPC8FNJzqzUoQ7ryloANKoQ0zJXCec5zAsYpfKNkkYl+c7ZzmqNP4NCVTficY5VsM7xfp6NqPGCK'
    'UK5jYah3lVS09JVjpqgqyyr0kClaFHkuK1IaTrk1tJRMCyV5yYrceeqZdkRxQVUmOqejOrfTUS1xLd5SaNu5GudMf+JTSF2X+uH7'
    'w8/X5fN1aa4LS0HiUdUfntiwlhGnV+AAnoVHF3ppia00N64gXHAwrysnKmaMKylDpFXGYEqmc8cLZkRBdUULq6zkjpeFLspezT1V'
    'MwzrepEpGdF+MbSd5fGfnHg30DKUMpUs9GCIY/pkuFaU1bhWOMCkPTm65ppuf3hixLGRh6EVzaWVTOVCGFlq4FLOWQdaj5OyUEJj'
    '5w/gXY5VijPupBbwS5W2nqlSVjRjVrcQ4WdGHLuHk6FCkBTqA3tUyAisQUbA9C7SpooO1kYfV+raJ9YGbxRJcJHH1VmX2iYtzx4g'
    'I4hEYEN/wegTWBu1w3RK7DXwqPrNUF7zEhpK2T61Nuw18BTWxkCKD/QphneyN7zXcDKrgWgRPxw+ollZOS+YZUr5QpWgalGpCkpz'
    'I1hhOHcEFESOHbFoWVSl8UZ4DUJQVSANCbO2bkDJu6WdMZzM6OxwcvtXgzZa4YHvI9Cc3Pb0/QhP2x+eGIX5TFdAjVyBW/XtJ3em'
    'lCURsfjj1R9Yl9bPkohYT2Jt0vJEv/CIe9gMT90vnNa9vBBNgVIV9WtTakmcNfCyce5zsCOKnBuwcInUlSicK7zw3lastL7UHhsC'
    'l6UGC9lKKa1wbR1mg9o4p184NUv6hR8DJHZtwKm5jzbgyX3EZBfNUn367COK76iI3MBwQNAd3CbyHK+N1R0YHv/aKJci1aeP9Ycn'
    'T/BgTb4BI1zpuje3sF44b0oOkq4opatkTkVemKqUhSfGU0kqV8LraokvCclLafOqMErIXLLKNoXw7dPPSvBgixI8WNpJpRfNp0/U'
    'cWh0h0dOm3anTT+f9hM6bbDTk7asfChPJkwtTuQvA/XoS5DSjIAcrsGVI9yO+cTi8QH+QxavdQrvnvWH91leRSmtg6w0VIzHB9XC'
    'VQRTp5WUOXeFBFbgXUEovP2+gHVpsHt0VSgjZUFLysFuYoIXiuW51GUH9xug/M8bzaaLeMacvzrx5qZuQQPD3xt+vgVP/hYYlgBA'
    'eaB0oJDef5q0Zf4l3o3oxcU2qggmOrrmSPdJr1lymQZw6A1PnKR0DLgQQfaJyY3z1HLL81wY6ovcs1zCm5uLSjBWSOeFcJSBBWl4'
    'Ea6syovScePhrS7bqll+VnSj0LonjW4UKJ8mzJzeIgx+GZtq/mkfVYNM3TTI1H2ow9TasPnnU1gb4+kgymOChK0bpLPYIL0PByh0'
    'qtH2k1jbABL2wXpQ3iFx3DSJ40eQsMO1xb6ST2FtVifEGtP94X0qsZw1DB3kL8jgmDfkq0q7XHAjMKk/51XpqaIcbDQtcoIVSpXQ'
    'yhtQZjWVrpSVzwtpratAqTW56Mk8dm6Zx/T8IFD8qwlZu3zW9Cc+hdR1qR++P/x8XT5fl7bVZspFxszDeYmYOVVKWdN2jMVyFf2J'
    'NQe6/eHJk/2SD1NUFPtcSyMsNTZXyhassJIUtgIbgRbGVianprKsKIkwQmgplXaUC1IYoWTGle3UFHnWZL97OBmjWCqHnMiHKSRK'
    'Pg5/sMdhxtBEig/V/eF9snZFa6YCAzBhZUSBIYJW1JYCDODCSG14SUvjiKHGc2IYBUOYcK6Eccx7TZWgRUE5JWVZWltI0rqzVL20'
    's2FG4YYtCkHfB8dOb24o3CWpoAbtD8cFz5lvBb7coRW55JrbJlUNLRF40+P1UK4Shbc5lbmqKBZuSaqMkIYU6BEBVmiJK4Al5q4o'
    'y7wsXGkxNqYVNrvjWdNpEaE77Jk5HqW3C/F6zz/1i+e4IvVO9zZYdPmpwysSFtMfnhN+8PMt+He4BVKSRJYRUf3hmUEoFalrTGBA'
    'GK2z+gvHgdtb+DCVUjp3lfKUOesIZd6AnuMLkAtMwUWAoweDwFCqlaOWu9wJZnsSgqgzSwiSLtCJX0zQ79Vtv3jivU2de/28/eHn'
    'c3965w7M89+kbusOVQm8qUrA7Hbdq7gw6fqsx782k4LJeqAKp9TjUPNgRUnBBqMk0anz4R6HpxJLyANBXRh23L+IPRiqSOpx6rtj'
    'HuhxUjXG2OTvwdwJ/N+lmWLyceiD9a3Ex0nFTB+omCH5OA9XfwCPY40ZpmA/EBJq8nFIKE1/KM8YVTTRyvPxtKjDdpcRIZTiAH2i'
    'XdvE47WFVp5PYm1GiiHEzAPlSYSt1jrFoB9TS4oay5AGLMOQEdtco8Ta6NNY2x+y1R/Wmw/7m//xv/rNzfX/La5//cN/Nz9bl9n1'
    'dr8r/P/str9lG3fls5t18YvfZfH7kB+/c7+FwXVW7j9cogXp/yd8/Y/N2wvvb+h/i/+mmcj+uHfr64vtfvXOFxeb7eX2/ccMrrVI'
    'c59//IF9WXdnzFYrU4/pP/4Qd6I3N5s8d/P2T51b/jd2y0NT1l992O7g8FbPrj74nf+9o5NxPfQVoXixGaVfxkpxeTSv+m9svfjd'
    'uvRu9Xq33X9Yvdiu3tyUGRg3iaAJmGr8y4g4cjyThpl09vf1P9eb1WuwjlffwLPuNsGId5dx2owxmUSjsjBtLO1RR9MamNZk375c'
    'vdm+R9u7uIaJiwx7VwxOS2LI7x9/kF9GGFjYUUwmNekdNXBakybG12vOxBae2GZ/WRfopVi9XrvNTe8urL7er/91kTGh1fFGYPcL'
    'koVmNKmdoAQmBr3uxcV641bv4GVYX1/Vh/Zyf/Nx9dXO+4waM7RDDXaq+ccf4IEjfjVcNdaMB88PdNhCOmwWHYrrodnPP/rN+7D3'
    '1iiTgmNWMKf+MjpM8DWhzTg1J5swJ501J8PnZNlbt4HT8+7XjzBxub++gdOtXxZFOE1IQKYzWLtOHibHSXn2Fl7h68M3D5uxDCbD'
    'fiWZrNnDYDLkalRkb9arZ/ubbbzHhA4bDxqUFHg+McmahfOpx8N1Iz+bMiubNStyMtAw38F78b/r1Zv1zcV6fwVr333IrLIq3QQM'
    'z6rurhXOqh6nZmezZqezZkduCSzwu/3Hvdusvtpuy6OzY5qKwdkZjDUDa5fJs0O+CbL5W3ji99d7uGCb97v9evXjhdtducLvgTO5'
    'S7h/x/wisHaGt/3L2G7mmCFTZJ1gtL9ycFW3xeo5iNWde++Pb5shqSdm+MQs+cTI4ECm/P3C/xMfegUvxsXeBX5/ebl+D7IbJoeZ'
    '5HBihU/8jz+o+pEVMFLWPP5wuy0e5gJCoY54BiGGnBVY71fb6ws416/d+maN53vpfkXuB5T/Cw4748yYBCHYYpg7FhQfnwFDHoeQ'
    'DBceJTa8RlfbfH3pu90/SnnkJGQdIRe1kXNyAneyHuM6VrwZD9eB3G8ONTySu1Djs6nxxdSQC2M3rWdvnh9pKpbTVAo6cngQppGt'
    '44st6leQpyZn0ydH20bOmBxZPfzOm/Xml0vUNaJcsnKI+iAQzvQff4C5Wj2D1GOampdNmTf2sJ8xL0oTOKWX283798jnftrvggJ3'
    'deV3xRqGz93mFxCsZrgAGxUl2ylKeKQmLQoYCpiFhPgsQihz4N38DtRnZIA9bawR4IbYFBc0yAVNigsylAXw02fXH6+KC3+1euNy'
    '0MtB51uDQEf5A4owzMtoIhUdTGO4PtiYonl41YyHD69wl5YQwujVHEIoiUBl+XELes6zb2p5z1RKF5GoiyTlGUOpA2fzdr+5BgXS'
    'wzTXcKyb4kDuZCCLeSL/JVxTfF7TvFeqGQ+fFzX4ZZRwZ+ZQQokHNspbv7n279ZXDYG4Q2lMhTA56207G9t2lHKTJ49gFJMn5yjZ'
    'gMt+u11vNoc3J6r39R5RphNqHVPBPmWqs0/jmKYIsUWEgiE8hxBKVE6ztyCp31+49er5euc3x680aOTHrk1dR3hg774kQTnVQEjX'
    'Y5UixBYRCpGtOYRQsnEQPqDbbN7/DuZzbV7g5mFMHBSQa3gxj5WP8CbCzYXXkKfeRI5yh/Psh43/rdbeFUnVE4YLZes7xHucdPgm'
    'cJQ5t8/Zcecpc6K84aJT89549+FqC7en70IBO46nnr0vBW6lwxbS4bPooLThMnv7VXN6rZSxfKDE4xkqPMOkU4WjlOGqtsV/8tEF'
    'dt3p8oECCFE1VIFJVIHRn1DbB6iljOnAHMXMMkqxLeh0SihnuIYXy2/+Bf+Avr35P8Amvn/RPwfOCJlnSXGUPNxk36036yv3e4+N'
    '0uOpTI0cH17RcIwmugBM73U1I68rip1pZPTdyPCJZFQN3b+IDIo2brM3IH1SepGifIxYO2mPWJIAm0uAziEgUMQJkn0N0uZfW1Ai'
    't/kWvXlg76dfCl5fVd67qjx9VQVKtalz42swZ24UZIJmr9fv3Uv/6xrUlrf+al1sgfmj6zCcNQd2lM6baDwWx2+BQGkiQJp883rV'
    'iq5ms6lS89xUAmWIAH5/tVn/ZX2N+Vhofb8H4e7L1Yv1roCZrjNJaOopKT4lTc4b/PDAi79987J3tZXVg1B/fEA0Xuril2C81NXj'
    'w10NTvgJEwcbcc7EyNWFzL7eu/8FEd2Yh8wm3bgsa5nhYOnI0oXKvncf97+sWwEf3UdUjmp8qlP4WuUvEYNANj55dtoqqxNnR9Yt'
    'gHWj88Xv4DVOvdFSDdigwMgBbnWMHeiw1XGsUlTYfCoxyDWDCooLYWqB97ctmg6rl26NPvtwb4IPw0iTOBCqw5bF3P64ZXFMU2TY'
    'AjK0LgSZTAb5uLDZ39Fya5yR8pi3RoGko6TQPemgR/YImfets0aRMH1WiRxbkuyv+9/XqISVBx49nPfDld/ADUgpAJSAAkCTARKJ'
    'HFXS7DtfXq5/TZn5jFqWCg2L6AIWXTCgHg+fHa2C+TToLBrIwiXLvv7zuxdff/fqZW1B4eSUJ24jmF14TWro9XBN4pimpmYTp6Y1'
    '8vzkqVFUSJ597a9B4oS7nWGL1IQLkEqMIiaVOImCQYrsb+syRlCjkUGlGdtT2tvTZn8TT4dyYcK8tL0PE+cNgVmZfftm9Z17v/rJ'
    '7fwXr9zu5qJjVIakrhxH2RhrKo6nRAEhVfZ8f/1x9R2aDwceSaqPA1rhrWbx/WO994+NvX8oI+YQCC/4HAIoJqTOnu1/X7269AWY'
    'skXHPhLxBgzEZLyxkI9nQ0YtDbCLn9ern1++WL1wH6/cJh4gUybdfwSelda8HwT7Stb3TQwfFhn0jOljgtT06ZExSwt6/O9737wW'
    'xOrUHqARxduoC3rGyYgNJZEx3zprLLucPqtCxqxI9hzE0xauRe0EBF3XJBywnIaJeRv0s6O8QaEePWViGpuwT58YWb6i2bf/9W71'
    '+nKbu8vVq98/7Px11HiopMPU0Vrta3U9EdV0nT4/hQx/HoVGY51KAdm9Yp2j6ZvL/a/uxu0wbOl/rq2CzFpyR2mukPkvITRXwCsU'
    'BYpnX/210cEP9FBqmEnZfCprgvDHLEChVFAihJp+BRvkUK1lRt7RB6tQOsyZf6YbVqGUUDL7o/94BXbT8/X2ulj7DfpXohhSKbUk'
    '7Hizy3WOgh4lwOYSoLMIhEQhlb34/vVPq2flr+hmL0Hk3fjdusnuQeaubcJJSTFckJb3CkWF0tn3sN35dvXt9uM17H8UGttNEBu7'
    'D2C+2XlRCIUyQwFT92uYNYb/jU1xMWJQ8JiUIFbIupXNnm3e+0sH6umB9zcyMGtT5oGKervq6e3RnR3GhI8wM2TqM+lFO2ExPb6I'
    'HuGL6GkUL5pkL/3lzfaA1mv0fP8Lk+a+2WTUGpuIUlOb2TqJ7fi0NUoBTbPv/c0rd+1XLy63e7ic+2u4PzEHaHhMuAxah4BoLzg2'
    'kraiUQzMJIFRpjkkUA5ojAOsK8xe+nZ9ibIy5NKhFzDqaXgmqbB6iPTxOrhXi3mbFvMa5cASQrJGH5xMCOWA5shW/e43zAtsLV4Q'
    '/TTxNtIo+inpRD8dEf0aQxFTJw96xZzJUeBoAcbMJfotD1MZ25iMNSoVAicxntLFXPlozFWj4FlGh8+igwJIg5nyzfevX/7w/es2'
    'xAS/Fe6vHqQ9qhoPTo1Y2xolglbZs9/9zh25Ltu0uKFDvA5+NyFpHbNebVqR0GipzKQRwwQzaITEVZ29/rD+3YMkQ+WouNhv4FUo'
    'dltfS6KMs+PMaztwa9mg501xPml0cS0hqu9GlC8iOnStzSCKklib7LnHzGDgLN7/8uFyfxQLAM3BkrsFGjTacQsJzQo4aFQLtM3+'
    '/vFy+/sX77abfLcu3/vVd77EVL2aHvqBUoY/iX4g0vMDkRE/kEZ9YCEhOoeQQcFsSPYVKEvF9mr11/UObsN15AyKm1TMgGLMIGmv'
    'G5TGhmbhVn3Y7m7Aot6V6211uf0Nn/0GvS8vUCclqXx7Ht9c3suJH9EnDMrkJYRi8v0MQiiZDesiln/d+h2mAt988fZit9784vJL'
    'j5mUNOWwxEzKdKzGoHw0PHu+W68PdPboR+Q67TgKfkTe8yPykWNFCTlrejprepSRRmRvt9VN4XY+Tqvs4LpEfMrgMem0BjnKkA3K'
    'xCnzRmzI6fOiDDQy+85tgTl8iHH56ysfUv9Rmrx7mVFOUymeh2lUspdGJVNk2BIyg2ytW8ig+DUq+2Zzvb4EsRi5AqZB1f6rd9sP'
    'wO64SgSojo1j1TNeVYoSW0hpYCbfQglFsdGYOfGTR0fL9Q3WJX2z+dVf31xhcODdbn+NDI8maj9CtET2oiXN1aUpQmwRoVALEt6R'
    'iYRC+YrJvgLutD5INFm1zAT2jnIi7ib9TChnWUZolvQzKP2Mzd6uC8xXWj13MHD4r9JtW7USxKyU6TwRMZKublEIWZJ9+/Zd64/e'
    'Nfq80SkZRHCypHJqUQZZMNfcJqR0P7vcX4Fls786ULGx2CfljGYyGApMdiZOTEsevoYWZdASQrLNvJ5ICGWQZWCs/3P7MVl6oETS'
    'q8bRq5ZMELMofizP3u6D6f399tftNSYhHumDTOnUi000mAYkmTVgUTJY4ODfvHjRPV6iRIahVc9GrHqL/NpixH19A/Zp6VZgfP+2'
    '3f3SJpgbbVIlLpi0S5MeHYtM06rsj+tNeY2Oz+jRxTDzQQ6YVcmYTawZor2aITpSM2SRZy4kRGcRQpZpdfZ140ABlgY3b+e+eAEz'
    'Nu8j1r6kPCCBk6H3hnfuiTFxb5FnLqMUu7NOp4RM05rs7YtXR7UebWL4d24Dlp9K6dexnCeGOGq1l43SYQvp0Fl0QkWgzX569eO7'
    '1fNn7969+umH1auY2dBGCVNtX2mdo1w3q4ubZtJRBIvGwmwisUnhZCKUxBpE0vCMn9C9+mZd+dXbqFy2Sdfc3JYL9+nENCTFlpEa'
    '5sPdRipUIhKavfabry7hZ6t3F37nPoTiq+s2jQptR5nyTMbcUGq6Khb6CVJsMSk+j1SoW8ROaBdu+8/96tVH1Nh2/ji6Qe0day4J'
    'C2taQmheISYJRZMg0V6jkH1/FFXiSqQEDA1mXpOGhmZeHMvk/GzG/LItqJw6f6jTBMP8b3v/+xo4qN9h3t/1geuSqVSwKWYCsQmR'
    'GqTCllCp40FTqYQ6TiKb1/O5u/bBWmu8E1F7Ax1aJ818Ha1v3bO+R1KbkBRbRioa+nNIhQpPuJMvLre/go0PVvOHg0LMxni2OoHj'
    'ElwulPZ8Lk18qh6Lsc1UYYVLyMY42GKy/A5kF682lLzCsQQ9MyTjNCfZKhfoqWN8GNclqDGG0Dpt3d6iGSdJsYWk0M86i1QotwWx'
    '85Xf7fzNzXr19sMzUIe4SZQTNCkOpCknaFIc0hOzKROzfu7EpIlDIS+xnV/r78DAf/NrUHsKrKjafVz91+oV+l5YQuzBE9mR4jga'
    'sQooyb5Cn8cVXqMvQ6pZxvmxYaiDyRWzVKJpqKMlxEYEWwQomDZ5LHCdM3lEJaDZy727RMUWXoqbvogO0b6BThgM0pAXJNpId3CL'
    'j5hztEYqmEsnpLCH/LSpdIIWQBmW7DTp632XQBMwI3f0EiEhtojQTCcRjdAJYFC8cNvCNVVIgqdqwAWWZtcm9WCeII1hL9t8mMAk'
    'fkRHMjCJ59ubIDERsyQF2hUMAEp7FoActQBoxFJYRgtZ7SxaQTLDXQxR9RJDu1/v3UfXlS/EyjXQ2GjqTYzxXWZ7L8voUQTJvISU'
    'agvPppIKkhks47+tc8y/x1fmqFK4Du+ohBcEriUdccLTiMBAdfYsz9fXv2wbt2ZUWqVKCQZ0Y1GWcjHQCLsAinq3J1ik+3F/GL+E'
    'KytN4p2LUpW1SB/BByzGNiWIhmW00IUzi1aQFiASn19uf+krk5zJVOttm+nWkX0MqBKkAyNtrO7Zeldui0FtJew+Tb15jb7TVzbG'
    '9FQWhMUyWvjm6Tm0guxgNHu+226K/XWv/m/3AQ0glSqY57GykPfq2Pk4ATaTAJ9HIGLdsLbi6cV2c32zg1uEpT8v/a/+cvvhCmNd'
    'nCQSY4PjVNb+TBFRksbQfwLmwiJSPCCfiBmkgsSAfXhWbG98ER0INwgxfWjiKZXIszBxA01vA9knNjCYkAtJsXmkgvxiPZnyZosB'
    'fe8PC1oVTRVnMSzOSgdwacBLgM/sa3d98Uf/8fDRuUl440PSW5MSNpwwMG/4+d/gCS86w+3Qgc30XX2UNIAlLCE000VJA1gCxSZc'
    '37yMyczcDMO/lNU2ROfjDqoNGeGxLJgrt82KChOZM2uQSIiUsN5sv3j7MToFv/5Y7rbvQTWrnYTvsKiRp4B2D6REne2uRhyEAZRh'
    'IS2M4ak5tIJEYl0s6o/+cr/5AhVpNPQaaxazR0mqgTSNXnbakKLNOEmKLSZFZ5EKkA3wmb1Y73bA4kFeFTtX3dQJtiTZEZ1mzYwD'
    'YRsAE+Aze3V9s99EnJ6IrN9hVsmZ5l2ALIDP7O9oQL9Z71df7Vff+t88oux/7JWbprSOICRUF/RCPUqOWC88CIn5hKL2P4dQxGPj'
    'GMgLagLWpV3jwV66/LqLOCVYn46ySPdgjOgojBENCArLCNF5hIJ44CL7ZrPZNi7z2oPWEw9Map3auwOLTPYsMpkkxRaSGhh/t5EK'
    'oomjaAKDH/SE3XoTIoO/tbwFCWbMKJnASwN1jjZOzeGbEqQUV9l3oEIX253vCsWGFzlmxHWJb4PJgoDowxy86udutVFRLdPBCz36'
    '7gVuzg0ixkWwkH7OaAvcAJIjofOz0GdxFJMwMFRus2++ev51Y6/zSbfDfuLIAu+8bdbGHJw6ayj7h8/sp+362u9+DaGh65uPl75z'
    'eKM9SEXK0x0ruGoPSl3hJ0d8wgEEYAmlWEs4h1Jg1aKHOfMV3MJrzGsLNm+dRoTmnL6ru0YEU2IZrbkem4BJAJ/Z650PEDquH4sw'
    'icSAusZO9nxczR7K5PxsxvyyxlyfPn8QDoJn7/YXewxt1c65hKzniObCazQXMZgoAnWK7Blyyl+jGO7XQUQ1NcOWigkgTBbrdWqI'
    'gsj+23Hz8xSQZ0TyXEaW3oUsX0aW3Wm1QUQImcVahkt39WFIFbGINU2mKh742lSgxcdUaxEcYMtowRKDt34yrSCeBIgnv369envl'
    'djeJ8nPOEkpXqEmwbZbk8cRBVAmd/W2/+cW15cBBPg2cO3yQasfjyzN6HMGmmTh7bIcwZ/YgDYXp2NjXHjbmRT88CJuEIidl2gfr'
    'zPQCPGw8wBOQHJaSYvNIBUEswNy42bnfciAD2g2IeeyEFdIT64pfOty/EeEmxkROEM4LKB0Uyk+hFFAf4DNg867eIR/9sL7pVCFq'
    'E53tj2E8YyUtGbMIZRDV82gg6OYsGkFIS5r9CNeg+HjoWlDqGA1ODBDWRA89LD0/mzf/EdrabfMHYSxZ9s5drb7FKuAhTnvI1tCJ'
    '4OvhTrHeTrEkJbaM0uBMbqMUBLTk2bPvX/Xr961NuO8xhIb1qCTpvg8AFPCZvXsLit6myZaXNIUlFUrgYSJOO7ynyL8TWkSAoJg0'
    'Myw/SoapM0dMbYksFo2ib9eb38HU70DcGtmgjvHnbEj7sDFVyjaVPUdpU2NqnYxg2wvIHiSDLSDLl5FVdWBiGdkgfhE2ow4koIfr'
    '4iM6eZubLBJQiHUtWlN/FnWX6KrhSSJsARHdILBPIxLkveyZpj/v0WI/DolwTu9sZsgg/ZfRmmtmBKAQ+Mxeeo8FoSgEfl+7budo'
    'svtgH5iH94B5eJICm0uBzqMQJH9A9tj97q+PESTZoJe5asMGrMUBDcEQPsYsgsSfQSECzc+gEGBE4DP70xrTtrer19tNCigphTMT'
    '6rLMSMI6DXAf8Jn9+Ozdqx9WL374/vtXL94dFjuAbpYxkywhD0y1SYiJ2ZOj2nbA/VhKis8jFSSyYhkGzQMMcGq3hDUJJzaR6CmS'
    'SWszwHDAZ/bqi7cR5aOf1+ze+6DgvcBYVApdKYDy0h4orx4F5UVSbBmpoEfOIhUEdATrwB4Pva0KntMB3pJsub7uVVbR0coqGtBA'
    'ZlKg8ygEga1kViPKoPuyxdLotCREx0zluAcFX/fU7nEsMRqQQZaS6njXJFKxPYbK/vw8OE57OSC7BufOpJoaszoFnfVS0Ef1VxVE'
    '5FwqKEtmUQkyUunsjdt+2EdshDbaYGjKjd1PnpUx2qDHZ2eTZ49bNGf2IAeVyZ5fuuKX1Vt/7a78EFsAVE/OExnNIeXBtE5yNFTr'
    'cZIUW0yKzSMVJKOyQaoU69VroFRXKCvOU4HjOpbYldGo0TIaGrBXJs/eYLxMnT0gn8Bn9v01itqWoawbfB5NaDpv0Y6giNKAeQKf'
    'CHET4Dte+st1sd7urxuYqgjYbK2RCaD96CNg9YvNe4+fUEoC+MlCWv2tmkQrCEPNsm+/ij6uWJR5mIJgB6q9GGCgiB40SZoOW0CH'
    'DyBQbqMThLDmbYOfv9eYps/dzQ3GMOsqfq50qsYxJCWIOg+hTvUc07sCIsoiUjGrdA6pIIR1L/fkOeh5Px4lj7/A6IFNJbkdFDqr'
    'XqGzStJiS2kNiqpvoxUEs5bZ31/Xni81qH7kdSFOU3rDB5MEOajhld0j0Dz2E/r9uKFTiBgNulCrYYK76mWdpx44SMNltAYZ7rfR'
    'iq2qNFg/v2xj+6vVsw388363LvaXN/tdgCJEMyjZsLlmOY22Gg5iBDWRBoCUhbTw0GfRCvJSm6xGuH6JJeEw57MCJr1e17kSu0yb'
    '4+YV/AgjZ3gVgtDSNgsiceeO5/wAkxKaYDGBZ9qGTcZCr1GWGWTXbCJ8FpGAEQKf2QuPRZ+XNXJsX7sP5XggA1IZzMFTJjtMR7Af'
    'xzAdkRRbRiqU5wfTayqpIEUNzX6+cZc/HuNlM5bERIqNxSjr+AllnVoWxyJJi82gJe9Ii8+gFatNl9IKEtuw7LsttpW9ud4ncawQ'
    'bFemIL9icQZlPVQuNnZcQWgvIRUBwOaQCnLb8E7qwC0ssBNDExqv84apSOlXsSFG03Iu6jxsrK1gAGpZSKtpqTeZVpDcRmSvwr5d'
    'rp57h009D7NTjeHphKouh2owcRCdCHyyfo8+nyHupzXaJnrKsliyFNMX66wPMnbXgjU7j0iMwMwhEuS3Udnr/fWF37y/cce1oIzp'
    'pE3QSYJBXCFgm8Bn9saDxddW2nWo4VKn2qn23W48NkYaS0kNoCZzKHTJzVMpxI6SWG6wRbdnwkulhgVYA2uj5i1yRM0IUCazqeD2'
    'z6ISJLOxWX3/v4K9asO4qGl8DVrAzUXGeCJk2CBM9xCgdROvTpJiy0hFB7WeTiqAqMBn9r3/7XK9+eUY0NMO6q+6Imzai7U1Bdk0'
    'SYLNI0HbgvKpJIJUtjR74XbbmxoAatBkPhgtpG48RHqNh8hoP9EggW+fF8OPs+YNEtCy7Oc1lqStQ3fYoWAChcWmHEfINRrGdDxz'
    'kEKI0eJCkKvlQv1abprKnqER0oeanuNupD0iUmFLqNB5VILksaI1T3/yRb7eDitx6LBXyzHc+C2VrTaYjEtIsQEA+W2kgtizEmR3'
    'qMHyx3D6aoDxEsIcDBtZsXRL4ABTA5/w+JsrtHV32w9+B5yjl2+ICrChKdP0wDvBe96J1NMH83EBJTbwg9xGKcg/2wsEvkTHbee1'
    'ZYKzsTa+Xe/ennIjk0TYbCIDDeo2IkEKWpP9/WJ9sY/+O25Joksso7G9UHOPUDWgHQrBaPpPAKa5jULIx7gDBT6BQpMhtYRC7BJt'
    'wYDHLOUECAVjiV7is5qxUBs7RM8lMqsbCyOxOzSBSwWc3a9eum1oD33YRnvgJhJtpXZbnh0zi8YqP1hAvFlCCB05dA6h0JWaIJr1'
    'b8DsP7SJzEeFXirRdUW2lXeis6/ZSNIXC3A3SyjRfo3fJEqhPzSJ3Q1+x1yjazB9L2PSabeD2JKGJLI+WZ312bULb8YiSYstoyXm'
    '0wrdownP/vzPfRurj7lswqSKY5uCVX17wSoLgDdTJ+9Xw06aPPSRJiJ7ewNGJmxP/7WUJFkaFIHsmbq9YocFmJvJ0x8UBE2aPvSM'
    'JjJ74/wvG1e3USBJW99mbYoKH8wT+kSD9vuX9TXIoyO3NTypoikEwr65qnrmajRdqR576tAsegY1OTCO51Hj86kFe2YJtdCfGq7g'
    '8/3NttGyMkVsEsKRZXwkdMYCUgt8Zm/2QVM/8CJ01f+gLEia4ukNq6U9VkvGXt/QpHoZLWTrZA6t0KYalLKfgdK/LvZDPJ+GLQ2L'
    'jgayqo4FjFSXIS22lBauaw6tACYDn9nLK+Cy8U2URqUAdGXW+PnlYJYg9ODwnnuMhgUkuT7LU5bYu6kjLGC6zCUxTxkJcC7w2SYI'
    '/m2/+ed64CPBwjKSOmMZStbSndsDsAp8LnLk1vkVTUrFp/tIIim2jNRBy8pJpIIYop0tiKpIHekJxbCgZP3qMUDGUmgoNUKr6EGU'
    '8LF7GiTSAkqqxYGYSikIJyo7h/FPr7765qdXByJQkxSS9UFgTPYCYzJJhi0gM4i/3UYmiEiqsq4NDDpSVt9vt+UlmqAhXWmDqa93'
    'RUxBWmwprZmIKSwAy8Bn9sI10P3HUevQz3kQjrnNmcwCxgx8Zs92O7e+PPDT8JTzhEXnCet5NBgdYzFBZE2bPCKYzJk8yCgwKvst'
    'fzYBG6C/N4jjkwTTF7GIRnRw7+NiI4iohaTkLFIB0AY+s9dufx0h5gIuQJ2wrVMJOLXLgfaCNnQkaMMCiM2M+ZvI4eT5g2xkNJSx'
    '+Y/p/NNEGSWrMzdZ274+eE7GNLgAXzOfTJNcN5lMkJHw1rzevt/+fhy5ESoVQosdrdooCB3HXWEBuGby9LEX3ZzpgxSG3/rLFuzk'
    'Ynt1lc4GlnfLcmUBqWYRHTqPTpC/IEexZdbLXb/rsZUykS7f6wM4sKYCMg18Zm9AxF6ic+Tl+v065IG2LeisSQUZa6ND9+J/I/mT'
    'SITNJRKDjHOIBKEHr883321vDjWEPmYCZVqkYvUHqDI1aLMdJ8UWkkJGMotUEHmwDz9tN9cIIXns72nUU0b03RDWWcDKWUqKziMV'
    'xC3ot29e/fjnCOl0ACVabyL8iqUJqKpDrwnveU14khZbSmvgobmNVhDGoEy/c5j4CFvWNN++aTczIBtIPtYqhvbgkvRoGwwW8HOW'
    'UKJtz5uJlAJ8Dnxm73Zuc10AAXdogymWagMJj8bbVp7HUwbpyGn2R9Aj/he1xeaR+35mIidm6suRW8aDdFxAhs4jE6Qjx2BlKMeO'
    'nQy/3l7HStmu6Zwy6m6lrCyg9ywlNauqlQX8HvjMXv74YvXSXUfdiIrj7RIDYRDzd9lI3hsLeD0TJg5Ze0EbmjpxkItcZK9Bwtyg'
    '/yR0h+gxr8tQIq90iqXw+jXnvdecj1NiyyghQ5lFKYhnLrPnfvOn1fPvjh1dEaZQpiHNKIbBk4B0LCDzwGcGz158vBno8GhKPfdr'
    'fHtkogKTx+LILp7bFIGklhCE5UJSdSnLVFJBWHKwD7+KJLq+9EqyhPcG0YvISM8VFnCB4DPoWWW+8+6X+sFrj1kCxPmoXYbqtbFI'
    'WDk8iKZ5FOg8CkEgcZuBtAPRVyfGFMNeAoLKVNft2KKwAV8JscrR1y9IpCWkeAi8ziAVoIrgM9SWvA0guYFY/ToolXodQg6C6hUQ'
    'qnF3UEAomkugtkmmEQgSUNDsq8vtb+uWkx8VxqlBk6OwVbHBCGmz5YOvd0wlCbBEiyihV1nPoRTkoGDZqzJA5m6a41YpYNN+JPLT'
    'aKMsoBFNmJgPQpy3TRxknODYgufnFFh/AIzkOuUkO8xwUL0MB5UkxBYRGqRS3EYoyEAhsuc7t9482++2O3es6aJpolRCZEQ0tzF2'
    'GFCA4DP7eb+p/MH73TdzMVSo7oh/xwIK0EJaMwHwWEABgs/s3d/+2Hbiq6vIRAJK4gjjRvaAZ9Kzszmzs3mzB4EnQOBdbFHFDWm1'
    'm/f/RFjHHczaQzmxqbIrEfsdi65cgYy+K8E6XEiKzyMVBK8wXWesb9d1sA/ULEzeDm8QhkNVqrK+ttg6r4Ecd20FDKKFtJCNzaIV'
    'hLEACfnN99+8/eEodBlzra1KJaxGid95Hum45zHgD80nQmcRCdBD8Jm9+OH7Fz/89HL19atnb959/eIZhjJ++uHPP65e/LB68w5f'
    'YT7m7aI9b1fj+RJJUmwZKTEAJ7mNVBDKkmZv17uN215hDXSNWa9T4TmbqTELN2ACwWf23H38l28SkuvDiAopSaEcNDEr3kuoGCkZ'
    'ZAENaB4N2fpOp9IIElLy7N03teIQ0zZUykPeBBFtzyk7ks2GE7MpE7N+dHLSxEECyi5m+Gb7oZZ9/bwv5MnJfpoWW2AmUfZZQAWC'
    'T7C+/4XWd0j1w8YTmz7chxSp5o6gs6m2y9jxvEECSZUFpIcfscX2Gw9mUnmMcG5SN/zgvRW99zZ1w4M0WkaJzqMUJFOAqWkqcLYb'
    'NJZ/+HBzgNfKdMq/f1h5wHqVByxJiy2lNahyuI1WkE3SZG+3e7QHQiHdX7ZFE2Z/Ve5xnClLecp91d/CT3eDZgEbZymprvH0JFJB'
    'MiGEzfYXuAU7BIFqMaEP8u9MCgQ4AkHQ2jCvpaAZYylBQC2jhRJ3Dq0AnAOf0cBefeWwecfHvg4Z6+yB6bNUUIBnLXy9GMwdpIWi'
    'WKNw5RE162LbhyLoEkcSiL8BDIbSNvx2PHcQH4qhLX2zLn2DXD/a6kd1gE6d3pvYjyAybp1XDpjubfMGMaF69XagNX1xgWiOAYh7'
    '+x7fQ4yByxR7P7BwWM/CYUlSbCGpgTF1G6kgSpTI/vZ1GsORE5Vik1Hl6MI7TTkHT5Jg80jwthZlKokgtVTMOQkdbV79frPzV77J'
    'pKlBDFQKVO6osdqnWwqxAI2ziFS/e9E0UkFmKtUVJvzo3m9L1zbMjj0q0MiiqTPqI3qKHog0T5Jiy0jxAV71baSC1FQ6+9N+4/Lt'
    '7/3Y7jGcm6i5SOSGqufuGJPJASRn1uxszuxBMiqTvdutLy/xtBujoxdtofquQZ0AkLOEzLygTgDHgc/sh12+v974dS9i2DOpQPgy'
    'm8pgrlufdFWgjSueJ2mxpbSin386rQCcA5/Zs32cumsot9teX3/xfLsr/Q4Fok2j/feCobIXDJVJUmwhqUHc9TZSQRZrmr3Y32CB'
    'a6/tHxr3OhWFp7E4v8kD7jJ7ktOz6dPzfu3/pOmDtNcse7vfXGxbKVZ7iUwi/ULGCyYbw60HHyiS87MZ84sWnnDq/EH6a549K7FF'
    '2nH5uEpih/W1C9HTLnhyfjZjfj7QXm6bP8h5LbKX7uIQWPGg2yKadjxhsxNErkn3lmAB/gY+O/fSOxCL1xfrar96EVBX4B3/LxCI'
    'YproHYM8REJsCSE16HF6G6EgeLXC0Mh6X5u8Kgl/1EuNHmi4AfoGPqMl+uK7YwPUmrFD7VWGdH7d1KkGaTeHwsBzfBuFIPG0yZ5t'
    '0A8ARicy7gYKsL1GuD8k5WzvN1RXvYbqMkmKLSMlB73bbyMVJJ+2HWKuW19tE3aN1KkQQoNgLjpSY61UWcDZWURK9sHSp5AKaDvw'
    'mT3L/+t5MBkO0s8U06nkWBJd8W1EjzXjJAE2kwCbRyDIOEMz/vIokECZECkRVKdWdzC5ahQmlwUsncmzo/ycNXuQcIZlb8C6ARaL'
    'FFzNvhMwHQmI52Zj0pOzGZO3uz5t8iDbDM+eu3f+jcs78A9h08CbPUDM47mCnDECcTj+F0MMAYSr57Rq8MtTujihtT5Eu0ZuZEwX'
    'N8GsXEKp3xt3EqUg4IxsX+Gf9yAxt8EfenW136yjuwodPjrNA+mY7AyoMfAZOt9dROf2secGk50YV3QMwLsF7e61iEudcjDxFpKi'
    '80gFwWdALHmHyszbi/WHD5350pESxI6lcPEe5KodT4cJ+DWLSMXcqzmkgig0Jnvzc6pdS0hDT0br62Qo1kuGYnycCJtNpNGVJxMJ'
    '4s/Y0PHin/vVy/VFsvWa0PZumLUsoNgsoTQXv5YFEBv47Hxmf79YlwdAQNgeFl03qRo3YMaYJZXsdcMCtAx8Ij55E4I6rDyrAYZS'
    'Hp8DvDzVw8tLyL+ANbOMEJ9HKAgry7KXAW29S/45vHOIipzCMqUx0k17kW46dtlsEF0LSfHo4p9KKggyy7OfvCvz7ccmfNBzoEQj'
    'hxs5FvuhvdjPxEhxQMJZRLaLfC8jyxeTvcNqg4TvweSA/AUD7KhbK4Y2WKqDmoh8qlZiexVSMkmKLSTV1GJNJhVkvO1k/FtgUe/x'
    'tr5xm3+5UDiFfZqYlDyREXkgHVVPOqokKbaYFJ1HKqgXVmU/ukt3DXL/KE1R8mS2tkQE6WR7ABYQeeAz+/Nmfb3dI7jqN4kSIKvS'
    'eV4z4i0BlmcJpdmRnYDNA58Z7H95HfoAYkfAUIndqzOgVqUcuySiWZEemhUZQbNiAaNnISU6j1IQ8dZmP4US0h/2N+V2uxsIE3SZ'
    'kYFbpsZabuCPdQDtGYM/ZgFPZxEp1ZpB00jxgKoDn9nf3j6PuHEmCVR+iNrPP+ma5AFBZ8KkTXHktElpeFLavuRv9jXgwU3/CqNr'
    'laTsk9jNjuguKDFWkMMDXM4iUl2h0VRSLKwKOyt7RI5+v3rn1ld+fZR+ic4MlcLd66Wiy8HcPMzNs+/WxW7743Z3s/re73exNfDN'
    'ugp4ppiZTDW7W0ceHhBslpKi80iJsCqRfb91F4MMFJVA58K23A12BxtMJ8N0MvsGDU//w6bYJlJtsWEYTTD2gwTFWBszhm+NlNhS'
    'SjyW8k+lpMKaVPbcrX9zXZaRGjSRlQN+K3tcUCanZpOnpvOm1uGpdfbdNv91u66zXblJYL/AHUHfSzLFggdcGfjMXmzh6m0L71Cv'
    '+XW928YYeZ3UhVkY/G4OVh5gZRaRmutg5QFVBj6zr6Ot99pdd3ufdmXUqTy891LxsZfKhoVMnp3Omj0gx8An+toQKsUnIeJoUhk4'
    'UM5YT5FmSTpsGZ0uTW0SnSCK4Mn+VK17gaWQFqOGgSV8eylGe2iqmJoHLBn4zJ7tQcC71bd+s/nYud2afAFUym0CIOFA2qietElc'
    'IsrCBi0gJQeC7TZSQfhQ3uU/fL3+gOXBP+580eV1ICmezupo8zgGUwcJADLjj95/qHddkinmZ+ueTDyvCFtz+6TRpTpx0iBcqAyt'
    '7FfflLDlTfHlYVQp5eUidc4aMb2SPzP2igXhsoQSXqFZlIJwoarJYPubw9SWS5BgIM125TUI/uub1TfYzVaPwdt3OPYxAsDH9i9I'
    'm2W0MNrA59AK4ofqgJH4ervZ/oZwtTvf6Nt/8Rfr4tJnhsiUDcky2mL2H88cBBI1rTrZtMMFI/yjb8FeEZKDpKAEbX06tof2Z8dY'
    'U5BIy2jhTZhFK4gkarMXyECwpuPwsllpUypBxNHuMjTGc+t5wI6ZQ+Ege38KhQAZA5/Zj/49KGNdkEPJBK8NiVOyh1zY5mImZ2aT'
    'ZpZtStbUmYP4YTR7/eaLF253jP5dbxH6Ikbjn7rnw6GjPhwekGIWUqLzKAUhCDfu2Ter6KePYURhyRhUG+1BtY03vOYBJ2bKzHQQ'
    'Hb5t5iDigF9+7esGuKEEqHGmuRtfrl5ceKxDYXcED+EBJmYhrZnoITxAxcBnhqVav0ev52VZgIYABPcbt6kD7FQdWyp6IHV1zw2q'
    'kqTYMlJq4Oi9jVSQxXAP/36xRpTiYYkTo0anMClkjGjJrl1JHNMkETaXCK09hNOJBAkMXOLr9fuLd5iQ3c8tq4vAabJ9YZP0R3u9'
    'F9nYNQjCdzaZfr70JDJB7sLpvXWxTVo/Mpfg77TtxKBv78TAAzzN1NkP+jxMmj1IdmaAx++u19c3HnsqhzaSfTLG3vkoglSfT2f2'
    'WQSJDrr32zer1/5qvVnXEYEOiGrYFvPYg1DrxGqcBptPo26OMZFGQJ+Bzxjm+H29+ml95Y4aAjDJUhj0JGIfE9VFSD9Fhs0nE7uw'
    'zSET5D2n2Tu3+S20tECPqiWpQEON1WJux2rhAehmyrR83rRBknOWhUKxo9ZA0adjUrX9NPYLq2F1gwY0ll/DA6LNPBrBpJpFI0h3'
    'zrNYQvUOLIHr/c5fH7TuxTQInsq4OGiTIXrJkWlKbBElPuiScRulINm5ACtx94u/ge2qtjtgJb2ebvHOylTZcsiEl21uZsiAkl2m'
    'Lx1zRwbgm4VU2V2o8kVUQ+LUYqpBy+AyKDQFtq3qOwR7/nw29EI1F7S9ldH+oiM9T5AWW0qreeEm0wo6B1eg2ACxK4TFal+4pr88'
    'xtlEIufSROe16cHQslEYWh6QfxZSYvMoBc2D6+z5/325D22Tr37ps5E62oaV3KlO4zEk1isp1A0ScpIUW0YqOEFnkQo6CTfZd/7G'
    'vQHV7aYp65YpIOIAmk+Spcs84ADBZ/anj/vN9aH614o3kmomR2NRVgdHYcfNjQABNJdKkHZzqAT0H/hExwii5B9gGyhlU6jvLHbi'
    'ZR240Gh0MmD/TJ4+AvLMmT6oAIJmz3f+5rDaoN4iyjlPeYejX4GoLleYjKQw8QD6M58Im0ckqAcCk5fKBrBK3jEpjgesn1smnZv/'
    'xgPOD3y2rrTn6+11qIA/iNihR4cnQrO87hnASde6rB4nabGltOCmzqMVlABE+6lp/fn5Dz+vfthconf+IIWQ2VRKcZPoy3qJvmJs'
    'D4PoX0YLb6+YQyuIXgG29zrff9wfmEYsFaendXJ0l4Erx9XBAC40efa+eJ00exCvohZ6q2exb93l5fp90FPWv4asJo9gBHcsTOUB'
    'W2gRqbmFqTwADcFn9sfVC4Rs7PX/GkJFDCDk43HTMasoQAtNnRzPQ82ZPAhRAQb3/v/4HaaRFS5sVshngJsb+3kiikLSz3JQyFkb'
    '3WKcFltKC/mkmEMrSHRhs2/95jogYx1kaDStW3VCXJG6R3lXACGacZIOW0IHWz7OoRMgheAze7tHd9u324/V/vIg++u6BrRliXqR'
    'cEyml+HNRjO8eUAUWkiJzaMUZL2k2auXz1KtdUEZkqnaS521kFKDnIiALASf2YtLt4O/SlUKb94jnqpKlb6YkFDW4fmZ0faSPOAL'
    'LaEU29zNoRTks+TZX/cXCA7uYgXJy2j7NGRfYHJ3yndMTKZHikh4gAOCT9Ch3zvE9OpgnZRKoxjFwuCuPx8bbb6As7PJs7PWWTh1'
    '9iAGEW8IAwR/3Q+hG2v0Em5TqVg9ML9BpkKAHOIIBNQgNTzfgvwY5MWpBKI8rCJqKaxxTwRONeqqCKBDC2k1uJ2TaQU5JXX2Yvdx'
    'W22L/TVen5A23Kt+wkp6dudlBam1iNTsVQUZJk32ky+LbV0j8LGR8I1lxTVJt15tm60O5g0yRFosDC8/HqFjtznExIz1qaW9PrUN'
    'AkkiyhEQf2aToX2gkylkAtgPfGYvtwgnlKbDeIKri2iWi176Khk3y1WQH7PpxPaic+gE6aH+X3vXthw3jmR/heGX3Y6gJeLCW/vJ'
    'dttqd/s2lty2OzpCwQtK4rhE1pJVkjUbHTH/sPuF8yWbmeAFVQW2irTcszExfqBLRRYOCCSQCSDzJHNfb64SWE63oSEUyterEVRU'
    'vrCGF8cuGwnaEcQvBFf38aZZP26KZNsI9r8y9loQz9CE8idGXgviG4Krtnwv6HDRiDbsNBRG7FuzepgMyr7BoOxbofhsKDYNinRV'
    'IDXTPu2eaGMLLZLXSUnLrhcl+iCGY0SxgZEEk4+m3BXEOzQLasghfCgUabHAR2+cdY1cczsp+7qNoVDY8kPpWAXPiFUYS4wiiH5o'
    'DlK3f3Y4EqnPAMaVKtOiFevAdsK55RgqDcfQVsLHpnwiHboLQey5nk5DEIcgBDqb5RwEUsZB6L7cNJcUw7NjY3d6K7C/1gCzWzDp'
    'wyByMToI+RCM2SXgttxTNLv4RsgLGw15EcQwdGjp2uVnSumkc4PYbbX4h+S2S8AYWvbOBCYtE9Zc9oK4fOAK3/fpuDIXFGnwdZmM'
    'BDH3HFYwm1YwqbSQuR+Pfjpyni1Vtq7RYP5xgykABrKEHUYg4fXL43ZFDM0yrI4DKxCfAdROpFOASJOGvF26fPr10xAJYOR8soaS'
    '6YgDbsQBjK72icJnDoyYBkN6NRTuabFKCkswagslmDyU9GpsbUykPvOgOq79g6FIr4bS/UD7+EpRtPOy0tlynlZls6432RpsUBHb'
    'InNaun3POLMZde4PSbHOw+on80OxSLOGPiZsulhvhXO3XraaRDCOrPlBpig8oheaBTVZtxLBEFxpTV2U+xlPWodRS+bEncN6bhyh'
    'MysQnwPE9s7q7wIiPRiG7g+q2ZTOO4Xkrc06WSP/5rVq1jprE+7TejYu3Ha7c8isFYxm1hLEejQTC70LJ2GRGg4j91WVq7okhlrV'
    'qJ4/ppUJFkgxxsbHDDa+bvuZWZH4LCS2t9F9FxKp5zBuiZbQOxNk4eekWoKEDF4cSNdnNVc55ku2O48TBZFAhqAXr5xfkowqPoTm'
    '2DJ90HEcTgeBsYM65sJGDEQTyscd2nBK+aS8I+a+P3186jzZFHq4nN42a3XV73kgwSSz5cvYWt8FxvousELxmVB7S8m7oEiBR9x9'
    'Xi0xkL4w1Kn10GzLr14afvXSWjg/uPA9p/27CidlHYlZJ02td2RoeEf2npJWKD4bik2DImUdSfeduqi0P0WSUWbFvukse3deO3N5'
    '7WzVpqoJxkSZtPREEPQjmARC6jny3ZdPbd7D0u72N8Fai0gpTwKYZg4SAxNcNUNNUQ72C+589IYNOg/b9j5YG1/FBltmjC9eEAPT'
    'LChpuo8fBEXqOArdnz5RdGSbfs6WFU5rjk51tCe/Y2ekRLd0QLmd/93B5ZKajWA9qsN0aWP5Eo8E9PKZM/tBz3D4slsg6bgIdNxm'
    'XUEjw7fNlpmqxUcKG++o9qbjhgcdHzt6JYKj6TjSjKI5BIfojeCKrObFjVpm1ZUyUwa1foJ4ohfbDBGTr7s1D+IR8yAmHTsHSlsi'
    'U6BI3cbM/anY/LXYUN/T3u8WeYTLYm4LdWA6ceJgx3XkR3YkPguJ9RxOhyKRso25+xcQ3/XF/pwV7LJLUByKZrluBUAHh/ThVU4f'
    'amWF49Pg6DzxK+DEZDh/Phzpf+Q4Is74vD8o1bku4tB2Jjjw7e1NC0ReBFf3VVLcbtRuijjue8wSwM9ah6XeExtV4lgULZEWTYNA'
    'rcumQJDWjX33l6Ssrs2DkrZdLIZw563befXogTnmQEscRVMQdPT7FATSu0RNVF4k5DVSlgoNFDyGH9hemQz3V1Y7h/zcOHpnVig+'
    'D4rtnfLfBUV6Nw61Ebkukh0lr6PETtA7XtgMF+0iOiSvjcfPh4gfaRaUMBIRHAZF2jmOkIoJWoySOKzqotnNesli/4D9ClyLsbH2'
    'o0XwDCRGS8kpSGQgxLF7trltnTBFbHMh83RcqGewvYylshVEhXR3oW0q3sMKlUR6BFf38duzdpu3KrfcjX32lTz2kjiQJmJMJLCX'
    'RIkEV/dtdaPqm6I0E1ZpM8maeHmLufuPQ0YkcSFNwxB77OB3YXB6D+7+WpWUI3A7jIFM4si3EnhpLegFg2mkW4tZUfhkFDLAupD1'
    'g1AEvYsgGoGWVDhJdek8krZViA6eY0ZgGxsJrpLEszShdDatdEl1l+5L1WwdYPnSszj5Q7OIEe4pSbRKcEVeq/LiRhW7nmwDI3I4'
    'iTjSUmuf2mQG1B1kkRaogN4qcJ8/OzHLr1dVrQ/29fLNklpLaINEGAwzcfuZWYH4HCBtW08BCumNwlmLdan3BWS7F+BrzxQ+Jl0h'
    'vVM4Z7GOTjB8AlREbxX1NJc/VnXxt6qkZXDpnG5Wqr5KkJeJ+xbe1625KzDmLt8KxedB+XvT5F1QMb1V7D6pkjr/gCvHYUMt6+1i'
    'mAk8Oea5HRq7B2x090ASH9RcKDYJisih4EpEQReUp2CpEjAnm63ITgyktbk/w2wV9D7du0WTfoTh8OrszOpJK3zPZshp2l8eDAS5'
    'Y0nFEYJPhhDTIEg7wjT1STXkFbUNwL3YxqFpnq9z43ydWwH4RAA2DYB0IoM156bIyKdjm8qydwoWQtp9BdhIiJokLii4us+S5tY5'
    'vcKEbrZukJLbds4jHbgUGUyN0Vg3SGql6UgDmdOhSKQ7YQz9uEmay8KGwpgl7mA3fFQaMSfSCsNnwOxFjt4FQzqTBW1080mdrC6L'
    'tTLNCy7CMSKGnnzB2P8XVhA+A4RNAyFdCb+ABlskPSO+5gqzZHHDRYPr9wFAu6WRjgLzjFabzSUlg0zKW7W9++N7h+ZCHrPqiBJq'
    'BpA0/cQOAiL9xGL3tFqsd4a4aUwG1hzGekNw2MOKRwmQJNFCzQDy+9n9QCBih4KrCy21xLXP+zpNSucN6HRthWm3NNRNtlM/Rtk6'
    '2yCjvbJJOXEGk2J5uUmG8zDTL/hHJJn1rdTqeuXTFU5m65gGIRqnWVB6IE6BImUFCoHyJ53iifswpQx9w6T0xkgYmZF9trPFmRWJ'
    'z0Jie1b/XUikvpB7aZMUKa0vlCqJfqVlTERnvGtsPhthYOsIyb2h+cYS7UmieZoFRT2lSWUOhCLNyTEKpbhMqp04nS7gN5bWPHUd'
    'sZ7hSDXG/imJ42kOkMngdxAQ6U20ER6/s1sYjU656bPgUEaekQ1jSURPM7HM7IsHYZEGhbnqpbrASQfXTVW9LjZX7bqWx2OZnPvs'
    'zVrCxZiEk/qcgMB6gtlDEUh3ctCdqk4bDHrqptDtuDDux8Li0GokQt8zAYmECa7uB1Usi13zO+Rfl3NeEvfSwcVPzTwviXIJru7H'
    '4j/QZwh5SjDuuyVq7CZpsI19+XVBCZKYl+ZCTQpKkETABFf3yeZ21N53XuA8Zp1ftpKrCyO5urBi8blYe4nc78IipS0YCMNy+RTj'
    'ePZ2mWRoI8Le0jd/7GogiZhpIojcU2p3gZCiFtx9WV1gdnP4r2iIc+Nsl+Oe2+jI9WGHZxx2jAXpSmJrmgVFy/FJUKSqhXA//GwP'
    'nRX+xI1MokyCq3vydOc0aGuhtEe+NjlmWhJN0iykiVHTkkiK4Op+qi6r3XZa0rxm18aeTk/rGekaRk4jJZETTcPQfqtTMEgxisD9'
    '+PG5fS0eiq90xpNERzQRY6JPniQiIri6zzcUQrU/3OPQIrZEZiIGMhNcyYTjEHwahHYsmwJBelhE7skmKf+mejvvRbmo6iut7MGI'
    'RS+58JAMI2gSjW2RCFLK87BQ/0/CIg0tYvfk/ePXJz+8eX3inH56ffbu/emZc/Kzc/bs9OzF6xM3CrhtudnlbegJqcQoIZUkJqSZ'
    'SGISErEhwbWLDKlAKnbkOopszBYsRL4oZmUZkMRSBFew9K8UzF/JWl20cUp7qzCt51szeOBICUZzIEqiJzq49C5Z98Glkz6UHLf3'
    'tlLp7WoQCxfYJNZISWRFc5CmEkdKYjCCq3uSLBMLUEeDGn3l/gKxF03Hmby5QMxFcCXDyHmJbNSnKtvUxbpQzZCZMrLtue+YYGKw'
    'jlhkReJzkFi0Z4fdhUS6WPrukyeULGk3NWEUMM/OfByO8M1LIhOCKzHANZt6kWSG01T/AlC0sOZsYW6X0mDPGiJGH7gS5aDzc0Wn'
    '6FDlNTmTld1mnBvxXScNvsfcawTmWxuGNNYMqLYPoglQpLlk5L5Pq7MtoY1sG3voLkWsedw+C5K+kHG7F32aLNboOrG7LsFpXIy5'
    'fIUGK2O3z25ZVRPNziwo1vuvHQhFTDtwdU/qpMwd8gYxLPluw82NbPbpzm6E+OO9AqLamQW1vy1xFxRpLB8zfTVrGuf7mW4RKo4i'
    'bqPr3PLI5oZHtmfF4vOwvH3v77uwSJv53P3pySvT07hNgMMs/n/kldXRY2mnPE+ONRupsEOLJ5+/ScWT3vJBn0CzYDY/lSws2/ww'
    'OuX+UQKtheQIh4UkZh+4uu9f/2XXAy4MD8w/NEK9I4nY5/DC2UGFP3CdB0W52qzPi6bZqOYoa64ffO88ANuhztR5Xd24dXJzDkPj'
    's6rpYw6WUeMii3BVmr+/SspiAbJ39Fe4gUX892+l4/zW3l0US/UbfPvbgxdv3zjqy6qApToWdfRl2VxBfcyHm8uE+4F+XMWhDBd5'
    'HGdRyNM49XKeRDJPWAQKNPZzpWTm5XEe+54fxZ4XIAX+IlGZx5PQU4u+6PaNGmxNXbRu2P4BrM35KqmbtqLvT52rqlxfHufJ7fGt'
    'SupHWO9lkRXr5a2TVeWiqNHFOb11rpMltJ3TbLJLJ2mIwrcPJW3Ry2TVXFbrc2i2hS6/S2XMo906Qqs3+AwOZ31Dt/95Vm1Kqjyq'
    'E30Hq1TlKtf9gvciHAxGa+qOxTte+32trgt1A79pkqvVUuknav3jkUf60nH+08DL4qJI21vn6hqUpn4g4O0TWXWt6uSCmpOEAb98'
    'X8KXxaKAhsMfOlXtXOI4rJ0mq1b0rB/r18bHn7Uo9LBZPN58vllvaioCjKcaKvAQusqp0XwvaHMDfwCr0u75x6DZoanMJ76HPs2W'
    'G/iWCvf6Z99s1g1mJciquoS2WVelGmoo8anf29dcQHGXXTv1j0DViR23bVvXaRS6aQI6rCGhyjB6k3YIwE3dBA/p105eLBaqJt4T'
    '10kxlSo9OXzroOLCCOccTZXmkVNWaydxUJ9VV46uyVEvVKu6qOpz3aHbksdCaIpV1RRr3DiD564QSD8J/22atrbqS5KtHS2Dx9Rr'
    'sAjOLlWPTE3gLDbL5cNNWUAHN8pJNnmxHmoxVLcdfMWXtm0eLtE5y0mXVfa5AVg0wVR+BCZxUsN4KhpnndQXCprOJTj15TLZNBgZ'
    't/WOmWpLPruEJgTphGpA/eEXKq2qz1Dwf8EwWDfOk2VVXaWqvnDefjx/8/bZaxc/vHwMy2AyCD6e//Lm5ftXz46c1xW8KNJ7O7p8'
    'B7VQCv2xhIKL0lkPSLfOry/e0u/LCuaBFOq33kBDrerqukDphtfA9hyq3KgG85Sd6zGvqz5U7cfTF7jFoL44Vdqo+po0U9O9BA4e'
    'WC6AgscqDK+4KYfysVrnaVfgufoCK4xuTCz05P+7nr+NId6KSacF2gkkdw11UEJPuq0+0Pfxh4ZqyDc4SWJhdNtdJg0+c3uOHQyd'
    'iLdwp9btp1C3F47ztH2oBdzUS7efuFEpdH+tYF5xV5t02XIetZWg2m99xsZx2zcEkcXCz1H+mccjqEqZF/TeqJ/aOdT41qjFUEX9'
    'ZdMrZlxi+ZSjUV2hAyXIwuOrFQjfFyNWQYT7K3tGcV6DstA62bz3dJh83LN6o9onLtfrVfP98fHNzQ07uvysvpTqpoEPx0vM3pFn'
    'FX2A5oWvjxsFN7DMY89vgfCDx7gfHq3yhevlfgS1S/IkzZQMotTzFn6QJlxk2SLFK8NlYh4zj/leGga5SvOFCOMolSKRsFgJBhL/'
    '2B3UmvuDMVWQpFIrOygQDs7Uuq8f9TMP3EaZbrvJqUqHR86bbF2BBDtY7pFuhaGFnP80mui7e26YvodxxQvr3SJXyfa5svBsnsvQ'
    'Hl3s2Vaf6lv32adoiLZ7G/jB82KfUdVd5v/j7/8DC+KezDwwO+btdoPvmzVgwwTOK5p/seqPnM4egCkOpgfnLc6Hzk0BCucGpzac'
    'l9sJMcmvijXOUE1yDdcsge58ZEy6OIUVMDO3fXnPL973GW4lhO6vyAvttB4OpoWvEz9x7lu9hWJXbKeHNG/d86CMW1om/OCJMIro'
    'PWKVqoUfgd0LAy3lcZT4UrKQxYsgYXmQxWrBcrB8vVyEmeA5y+MwiZPU98M8BxvY1SPRb0OH/pxB2bYQDspyXVfLJTzxj7//LwoV'
    '/vkQdDuy9EIRMD0WVf7dI2f7V8NQ7u7Ek8q7507ohYl8z6BF+zTapwm5OpmMA3jE6PneLBI889HhXe9DugLtrePjB1j4iYBejLNM'
    'JbCgyiOZxrBcWggmOby2nyw8xnwWCZjyeZRx7qVpmPpxwNOEpbEMEpHni8SFVUDPpO+Z0gVGOYkTWYm1WqJCddZV12XwVs2x+oJ7'
    'LdCK2upt3NaKbBSsvvB3ps2tl4po1hhTP7L9mVKhxcW4sS9Iwe5P7rlJTVnhs2TlQBZD89F7noq+lbDIEWH5tvbB/2ch0QlvOofi'
    'd6q1KZ0nSi+Zu6gEGXgW9i+dda2LUNEJGexS4m3PKH84e96DNQJGtd61xg8e83xBry3SJMxFEiuwL5Moh7fK/HDhx16UxTl8GeeL'
    'OI6DOI+ihS9yJbLQYxmPQLWLbOHlbmjG6fzZsw0GIvGJOs341Y6UtZ3yrXXaaF+YIshniiC924EiqJ+9X6P3W4lZPCJm33Se+hcV'
    'L+0DgL7tOmDtXZVWeMCC+Yqsc5poBUoYAiVG5jRxn1YSvYvo3gWWdLHU7xJ7cST8MMrC3Pczf+GFXCwS0HWZ9JM0VDzPPRCoXHlg'
    'dMOqOA8SHqjQy6SUiyjibjTIk/inTFtiVELEPai50YYzhYAfLARU4QOFQD9777PKN5EC5o2IwTefVv7Z3X/v+c9d9335uaxuyvZP'
    'vYMA/4xWPd2kDe6SgpltDhFszIebFQwVmAAb2vkHobwocDMgq65WSXnbbkTTqMySEsdbrVa1arCsZLnshuKR86w9M4BJtMyHQfqw'
    'WamMktDUChorg/7SzmL9fr3KH1LxK2P/pNnZ1eLBfe1qta2HHzwmWEy9M7M0xtrS8IPHuJRY2iPHKO7g0jCkVjvEBfgBC43OMyzv'
    '2yW4Nx/dmzfocr/NrFjGwPiQKl5EsApIfQ6TSBLnHotgnCxkhGujKILFQZBlacxkCmspmaQ8ZokSuIHjDwRnwb8F/F9OwL99EjAz'
    '9sgwB4+hc9RDPIbvTIuv1Z3imPE2Egw/YFdIPQhyP43B1OZp7kUsTSOZcMllJr0szKI0iD1PyVD5ME7CLIQhE0eSp2HOWArNlvmu'
    'GBx4mTkGPlT1Zzry2jSgO0uQeWddJySpoDkfdaoSpLBca926o1VRxPWm93DCCqOpVnQqjOeaDsWSK7TL+nF008IOAt3GW1kUbdvy'
    '+zsQ/Y23bTccO7vG3nf33A9/Zjo4MwrtPg20bydkwb+F7H6E7MHv/wejqmbx'
)
_EMBEDDED_INPUTS_SHA256 = 'f3bc6619f59b97655dec647b3fc6bc51136a0bb995db01b9e2e1e437b82d36f1'

def _read_embedded_inputs():
    import base64, zlib
    raw=zlib.decompress(base64.b64decode(_EMBEDDED_INPUTS_B64))
    if hashlib.sha256(raw).hexdigest()!=_EMBEDDED_INPUTS_SHA256:
        raise ValueError("Embedded input snapshot failed its integrity check. Copy the complete original file again.")
    return json.loads(raw)


def run_lockup_backtest():
    import io
    import matplotlib.pyplot as plt
    plt.close("all")
    try:
        from IPython.display import display, FileLink, HTML
        from IPython import get_ipython
        if get_ipython() is not None:
            get_ipython().run_line_magic("matplotlib", "inline")
    except ImportError:
        def display(value):
            print(value.to_string(index=False) if isinstance(value,pd.DataFrame) else value)
        FileLink = None
        HTML = None

    embedded = _read_embedded_inputs()
    def data_file(name):
        return pd.read_csv(io.StringIO(embedded[name]))
    manifest = json.loads(embedded["input_manifest.json"])
    CFG = Settings(asof=AS_OF, day_basis=DAY_BASIS, anchor=T0_CONVENTION,
                   borrow_rate=0.0, fee_bps_per_leg=0.0, require_complete_30=True)
    strategies = make_strategies(EXTENDED_COMBINATIONS)
    BASE = Path.cwd()
    run_id = pd.Timestamp.now(tz="UTC").strftime("%Y%m%d_%H%M%S_%f")
    OUTPUT = BASE / OUTPUT_ROOT / run_id
    OUTPUT.mkdir(parents=True, exist_ok=True)
    pd.set_option("display.max_columns", 18)
    pd.set_option("display.max_colwidth", 65)
    pd.set_option("display.max_rows", 60)
    market_ready = False

    def show_returns(frame):
        view=frame.copy()
        for col in {"mean_gross","median_issuer_gross","win_rate","worst_event","gross_return",
                    "ci_low","ci_high","benchmark_return","short_plus_hsi_return"}.intersection(view.columns):
            view[col]=view[col].map(lambda x:f"{x:+.2%}" if pd.notna(x) else "Unavailable")
        display(view)

    def heading(title):
        if HTML is not None:
            display(HTML("<h3 style='color:#24282b;border-bottom:2px solid #cb2125;padding-bottom:6px'>"+title+"</h3>"))
        else:
            print("\n"+title)

    heading("HK IPO lock-up backtest")
    print("PRICE SOURCE: Bloomberg Desktop API only. Gross short price returns; no assumed costs.")
    print("INPUT: embedded snapshot of your workbook, 451 rows / 448 tickers / 804 date entries.")
    print("OFFSETS:", DAY_BASIS, "| One month = 21 sessions in session mode. Cutoff:", AS_OF)
    print("T0:", "Last HK session on/before the filed last day locked." if CFG.anchor=="last_locked" else "First HK session strictly after the filed last day locked.")
    print("Date confirmation: 13 dates across 8 issuers freshly checked; other accepted dates reuse the 17 October 2026 primary review.")
    print("Unverified/amended dates are excluded. This is not certification of every supplied date or later amendment.")
    print("Gross return = (entry - exit) / entry. Positive means a profitable short; losses can exceed 100%.")
    print("OUTPUT FOLDER:", OUTPUT.resolve())

    heading('Strategies')
    display(pd.DataFrame([{"Strategy":s.key,"Entry and exit":s.label_for(CFG.day_basis)} for s in strategies]))

    heading('Date validation')
    if WORKBOOK_PATH is None:
        inputs = data_file("input_events.csv")
        input_issues = data_file("input_issues.csv")
        input_hash = manifest["input_sha256"]
    else:
        inputs, input_issues = read_workbook(Path(WORKBOOK_PATH))
        input_hash = hashlib.sha256(Path(WORKBOOK_PATH).read_bytes()).hexdigest()
    review = data_file("event_review.csv")
    audit, eligible = select_events(inputs, review, CFG)
    display(audit.status.value_counts().rename_axis("Date review").reset_index(name="Count"))
    if not input_issues.empty:
        display(input_issues)
    sample = data_file("sample_date_review.csv")
    display(sample[["source_row","ticker","input_date","last_day_locked","category","review_note","source_url","source_page"]])
    audit.to_csv(OUTPUT / "input_audit.csv", index=False)
    input_issues.to_csv(OUTPUT / "input_issues.csv", index=False)
    print(f"{len(inputs):,} supplied date entries; {len(eligible):,} accepted historical cornerstone dates.")
    print("Prices and complete ±30-session endpoints still need to pass the next checks.")

    print("Only confirmed historical cornerstone restrictions enter. Other holder groups remain separate in the audit.")

    heading('Bloomberg prices')
    market_ready = False
    try:
        if eligible.empty:
            raise RuntimeError("No confirmed historical cornerstone dates are eligible. Review the input audit first.")
        if importlib.util.find_spec("blpapi") is None:
            print("Installing Bloomberg's official SDK in this kernel ...")
            subprocess.check_call([sys.executable,"-m","pip","install",
                "--index-url=https://blpapi.bloomberg.com/repository/releases/python/simple/","blpapi"])
            importlib.invalidate_caches()
        prices, sessions, benchmark, provider_errors = download_bloomberg(
            eligible, CFG, BASE / OUTPUT_ROOT / "bloomberg_cache", refresh=REFRESH_BLOOMBERG)
        if prices.empty:
            raise RuntimeError("Bloomberg returned no usable stock prices; inspect the provider errors.")
        market_ready = True
    except (RuntimeError, ImportError, OSError, TimeoutError, subprocess.CalledProcessError) as exc:
        print("BLOOMBERG DATA UNAVAILABLE:", str(exc))
        print("Log in to Bloomberg on an entitled Desktop API workstation and run this notebook there.")
        print("If package installation is blocked, use the IT-approved blpapi installation in this kernel.")
        print("No price backtest was produced. No substitute provider or old report was used.")
        (OUTPUT / "connection_status.txt").write_text("Bloomberg unavailable: " + str(exc), encoding="utf-8")
    except ValueError as exc:
        print("DATA VALIDATION FAILED:", str(exc))
        print("No backtest was produced. Review the returned data and calendar before rerunning.")
        (OUTPUT / "connection_status.txt").write_text("Data validation failed: " + str(exc), encoding="utf-8")
    if market_ready:
        source_description = "Bloomberg Desktop API / exact-request Bloomberg cache"
        prices.to_csv(OUTPUT / "prices_used.csv", index=False)
        pd.DataFrame({"date":sessions}).to_csv(OUTPUT / "sessions_used.csv", index=False)
        benchmark.to_csv(OUTPUT / "benchmark_used.csv", index=False)
        provider_errors.to_csv(OUTPUT / "provider_errors.csv", index=False)
        print(f"Bloomberg: {len(prices):,} stock observations, {prices.ticker.nunique()} tickers, {len(sessions):,} market sessions.")
        print("Actual price date range:", prices.date.min(), "to", prices.date.max())
        if not provider_errors.empty:
            display(provider_errors)

    heading('Backtest and coverage')
    if market_ready:
        trades = backtest(eligible, prices, sessions, strategies, CFG, benchmark=benchmark)
        summary = summarize(trades, CFG)
        common_summary = summarize(trades, CFG, common_sample=True)
        completed = trades[trades.status.eq("Completed")] if not trades.empty else trades
        completed_ids = set(completed.event_id) if not completed.empty else set()
        admitted = eligible[eligible.event_id.isin(completed_ids)]
        paths = event_paths(admitted, prices, sessions, CFG)
        print(f"Completed: {len(completed):,} event–strategy observations, {len(completed_ids)} events, "
              f"{completed.ticker.nunique() if not completed.empty else 0} issuers.")
        if not summary.empty:
            show_returns(summary[["strategy_label","events","issuers","mean_gross",
                                  "median_issuer_gross","win_rate","worst_event","ci_low","ci_high"]])
        else:
            print("No complete trades. The audit and skip reasons explain the missing coverage.")
        if not trades.empty:
            display(trades.loc[~trades.status.eq("Completed"),"reason"].value_counts()
                    .rename_axis("Skip reason").reset_index(name="Event–strategy rows"))
    else:
        print("Backtest waiting for Bloomberg data.")

    print("Admission requires complete ±30-session endpoints. Missing prices are never filled or shifted.")

    heading('Charts and common sample')
    if market_ready:
        figures = plot_results(summary, trades, paths, OUTPUT / "charts", settings=CFG)
        plt.show()
        if not common_summary.empty:
            print("Common completed cohort across every meaningful strategy:")
            show_returns(common_summary[["strategy_label","events","issuers","mean_gross","win_rate"]])
            best = common_summary.loc[common_summary.mean_gross.idxmax()]
            print(f"Highest mean within this common sample: {best.strategy_label}, {best.mean_gross:+.2%}.")
            print("This is an in-sample comparison, not a tested trading signal.")
        else:
            print("No complete common cohort; available-per-strategy results remain separately labelled.")
    else:
        print("Charts waiting for Bloomberg data.")

    print("Means give each issuer one weight. The bootstrap does not correct calendar clustering or multiple testing.")
    print("Retrospective selection and unverified historical borrow/fills prevent interpreting this as executable net P&L.")

    heading('Individual stock')
    if market_ready:
        if not trades.empty:
            detail = trades[trades.ticker.eq(SELECTED_TICKER)]
            show_returns(detail[["ticker","last_day_locked","t0","strategy_label","entry_date","exit_date",
                                 "entry_price","exit_price","gross_return","status","reason"]])
        if not paths.empty and SELECTED_TICKER in set(paths.ticker):
            stock_figure = plot_stocks(paths, [SELECTED_TICKER], page=1)
            stock_figure.savefig(OUTPUT / "charts" / "selected_stock.png", dpi=160, bbox_inches="tight")
            plt.show()
        else:
            print("No admitted price path for this ticker; see its input audit and skip reasons.")
    else:
        print("Individual trades waiting for Bloomberg data.")

    heading('Excel and data exports')
    if market_ready:
        run_metadata = {
            "data_source":source_description, "input_sha256":input_hash,
            "review_sha256":hashlib.sha256(embedded["event_review.csv"].encode()).hexdigest(),
            "actual_price_last_date":str(pd.to_datetime(prices.date).max().date()),
            "original_input_dates":len(inputs), "completed_events":len(completed_ids),
            "completed_issuers":int(completed.ticker.nunique()) if not completed.empty else 0,
            "bloomberg_session_established":True,
            "return_basis":"Gross short price P&L / initial notional; no assumed cost deduction",
        }
        export_results(OUTPUT,audit,trades,summary,common_summary,paths,CFG,run_metadata)
        excel_path = export_excel(OUTPUT / "HK_IPO_Lockup_Backtest.xlsx",audit,trades,summary,
                                 common_summary,prices,paths,CFG,run_metadata)
        write_html_report(OUTPUT / "backtest_report.html",audit,trades,summary,common_summary,
                          figures,run_metadata,settings=CFG)
        saved_manifest = json.loads((OUTPUT / "run_manifest.json").read_text())
        saved_manifest["excel_sha256"] = hashlib.sha256(excel_path.read_bytes()).hexdigest()
        (OUTPUT / "run_manifest.json").write_text(json.dumps(saved_manifest,indent=2,default=str))
        print("FINISHED: Excel, HTML, full trade ledger, audit, input prices, charts and hashes saved.")
        for filename in ["HK_IPO_Lockup_Backtest.xlsx","backtest_report.html","trades.csv","input_audit.csv"]:
            file = OUTPUT / filename
            if FileLink is not None:
                display(FileLink(str(file.relative_to(BASE))))
            else:
                print(file.resolve())
    else:
        print("DATE REVIEW COMPLETE; PRICE BACKTEST NOT RUN. Resolve the connection or data issue shown above.")
        print("Input audit:", (OUTPUT / "input_audit.csv").resolve())
        print("Primary date checks and links are displayed above and included in the exported input audit.")

    return {"output_folder":OUTPUT, "audit":audit, "market_ready":market_ready,
            "trades":trades if market_ready else None, "summary":summary if market_ready else None}


if __name__ == "__main__":
    RESULTS = run_lockup_backtest()
