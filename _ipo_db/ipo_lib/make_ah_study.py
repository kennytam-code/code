#!/usr/bin/env python3
"""Regenerate the Word A/H study from dated, auditable analysis observations."""
from pathlib import Path
from datetime import date
import json,math
from collections import defaultdict
from statistics import median
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.section import WD_SECTION_START,WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
ROOT=Path(__file__).resolve().parent.parent

def table(doc,heads,rows,widths,size=10):
    t=doc.add_table(rows=1,cols=len(heads));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for cell,width in zip(t.columns,widths):cell.width=Inches(width)
    for j,h in enumerate(heads):t.rows[0].cells[j].text=h
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        for cell,v in zip(t.add_row().cells,row):cell.text=str(v)
    borders=OxmlElement('w:tblBorders')
    for name in ['top','left','bottom','right','insideH','insideV']:
        el=OxmlElement('w:'+name);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
    t._tbl.tblPr.append(borders)
    for i,row in enumerate(t.rows):
        pr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');pr.append(cant)
        for j,cell in enumerate(row.cells):
            cell.width=Inches(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc=cell._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'1F3864' if i==0 else 'F2F5F9' if i%2==0 else 'FFFFFF');tc.append(sh)
            marg=OxmlElement('w:tcMar')
            for key,val in [('top',90),('bottom',90),('left',90),('right',90)]:
                x=OxmlElement('w:'+key);x.set(qn('w:w'),str(val));x.set(qn('w:type'),'dxa');marg.append(x)
            tc.append(marg)
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=1.05
                p.alignment=WD_ALIGN_PARAGRAPH.LEFT if j<2 else WD_ALIGN_PARAGRAPH.RIGHT
                for r in p.runs:r.font.name='Arial';r.font.size=Pt(size);r.font.bold=i==0;r.font.color.rgb=RGBColor.from_string('FFFFFF' if i==0 else '000000')
    gap=doc.add_paragraph();gap.paragraph_format.space_after=Pt(0);gap.paragraph_format.space_before=Pt(0);gap.paragraph_format.line_spacing=Pt(1);gap.add_run().font.size=Pt(1)
    return t

def pct(v):return 'N/A' if v is None else f'{v:+.1f}%'

def link(doc,label,url):
    p=doc.add_paragraph();h=OxmlElement('w:hyperlink');h.set(qn('r:id'),p.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True))
    r=OxmlElement('w:r');pr=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'1F3864');pr.append(color);r.append(pr);tx=OxmlElement('w:t');tx.text=label; r.append(tx);h.append(r);p._p.append(h)

def main():
    batch=json.loads((ROOT/'data/batches/ah_study.json').read_text());recs=batch['pairs'];usable=[r for r in recs if r['eligible_1m']];bycode={r['code']:r for r in recs}
    out=ROOT/'out';out.mkdir(exist_ok=True);figdir=out/'study_figures';figdir.mkdir(exist_ok=True)
    doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
    sec.top_margin=sec.bottom_margin=Inches(.65);sec.left_margin=sec.right_margin=Inches(.7)
    for style,size in [('Normal',10.5),('Title',21),('Heading 1',14),('Heading 2',12)]:
        st=doc.styles[style];st.font.name='Arial';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
        st.paragraph_format.space_after=Pt(7);st.paragraph_format.line_spacing=1.06
    doc.styles['Heading 1'].paragraph_format.space_before=Pt(11)
    doc.core_properties.title='A and H Share IPO Premiums';doc.core_properties.author='HK IPO Research';doc.core_properties.subject='Dated first-month A/H premium analysis'
    foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT;run=foot.add_run();field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');run._r.addnext(field)
    narrowed=sum(abs(r['premium_1m_pct'])<abs(r['premium_day1_pct']) for r in usable)
    doc.add_paragraph('A and H Share IPO Premiums',style='Title')
    doc.add_paragraph(f"Hong Kong IPO database. Closing observations through {batch['asof']}. Prepared 6 October 2026.")
    doc.add_paragraph(f"{narrowed} of {len(usable)} usable A/H IPOs had a smaller absolute price gap after one month. Convergence is common, but deal size and company-capitalisation proxies do not explain it reliably. GigaDevice narrowed mainly through H-share gains; Montage narrowed mainly through A-share declines. Several other pairs retained large A-share premiums.")
    doc.add_paragraph('Scope and price convention',style='Heading 1')
    doc.add_paragraph('The study covers all 62 A/H pairs in the IPO database. Sixty have usable first-month observations. The first-month endpoint is the 21st H-traded session after the listing-day close. A premium is the A-share CNY price converted to HKD and divided by the H-share HKD price, less one. A positive premium means A is dearer; a negative premium means H is dearer.')
    doc.add_paragraph('GigaDevice and Montage',style='Heading 1')
    cases=[]
    for code in ['3986','6809']:
        r=bycode[code];cases.append([f"{r['name']} ({code})",pct(r['offer_reference_premium_pct']),pct(r['premium_day1_pct']),pct(r['premium_1m_pct']),pct(r['h_return_1m_pct']),pct(r['a_return_1m_pct'])])
    table(doc,['Pair','Offer ref','Day 1','1 month','H move','A move'],cases,[1.85,1,.9,.9,.9,1.05])
    doc.add_paragraph('GigaDevice reached an 8.7% A premium on 11 February 2026: H rose 32.9% while A rose 8.0% from Day 1. Montage crossed to a 4.9% H premium on 13 March: H rose only 0.3%, while A fell 15.5%. Similar-looking convergence therefore came from different price legs. Offer-reference gaps of 80.5% and 72.3% had already narrowed sharply by their first H closing prices.')
    fig,ax=plt.subplots(figsize=(8,2.75))
    for code,color in [('3986','#1F568A'),('6809','#64748B')]:
        r=bycode[code];hist=r['history'][:22];ax.plot(range(len(hist)),[x[4]*100 for x in hist],label=r['name'],color=color,lw=1.9)
    ax.axhline(0,color='#999999',lw=.8);ax.set_xlabel('H trading sessions after Day 1');ax.set_ylabel('A premium (%)');ax.spines[['top','right']].set_visible(False);ax.grid(axis='y',color='#e6e6e6');ax.legend(frameon=False,loc='upper right');fig.tight_layout();fig.savefig(figdir/'first_month.png',dpi=220);plt.close(fig)
    doc.add_picture(str(figdir/'first_month.png'),width=Inches(6.8))
    doc.add_page_break()
    doc.add_paragraph('Deal size and capitalisation',style='Heading 1')
    doc.add_paragraph('Larger deals and larger listing-company capitalisations have weak negative rank correlations with the change in A premium. Their confidence intervals include zero. None of the six exploratory tests has an adjusted p-value below 0.05. The data do not justify a rule that a larger issue will converge faster.')
    labels={'deal_size_hkdm':'Deal size HK$m','company_cap_at_a_reference_hkdm':'Listing cap at A reference','offer_size_to_a_reference_cap_pct':'Offer value divided by A reference cap','premium_day1_pct':'Day 1 premium','a_runup_pre_ipo_pct':'A run up before H IPO','current_company_cap_at_a_hkdm':'Current company cap at A price'}
    rows=[]
    for st in batch['correlations']:
        lo,hi=st['spearman_ci95'];rows.append([labels[st['field']],st['n'],f"{st['spearman']:.3f}",f'{lo:.3f}',f'{hi:.3f}',f"{st['bh_adjusted_p']:.3f}"])
    table(doc,['Factor','n','Rank rho','95% low','95% high','Adjusted p'],rows,[2.7,.4,.85,.85,.85,.95],10)
    doc.add_paragraph('Capitalisation must be defined carefully. The historical proxy revalues the database listing-company share count at the A and FX closes before H debut. It is not the market cap of the A-share float. Current company cap uses a later snapshot and has look-ahead bias, so it is descriptive only. Offer value divided by company cap measures funding size, not tradable free float. These proxies cannot settle whether A-line float size explains convergence.')
    fig,axes=plt.subplots(1,2,figsize=(8,2.7))
    for ax,field,title in zip(axes,['deal_size_hkdm','company_cap_at_a_reference_hkdm'],['Deal size','Listing cap at A reference']):
        rs=[r for r in usable if r.get(field) and r[field]>0];ax.scatter([r[field] for r in rs],[r['premium_change_1m_pp'] for r in rs],s=18,c='#1F568A',alpha=.7);ax.set_xscale('log');ax.set_xlabel(title+' (HK$m, log scale)');ax.set_ylabel('Premium change (pp)');ax.axhline(0,color='#999999',lw=.8);ax.spines[['top','right']].set_visible(False);ax.grid(axis='y',color='#e6e6e6')
    fig.tight_layout();fig.savefig(figdir/'correlations.png',dpi=220);plt.close(fig);doc.add_picture(str(figdir/'correlations.png'),width=Inches(6.8))
    doc.add_paragraph('The outcome is the first-month premium minus Day 1 premium, in percentage points. A negative change can reduce a positive A premium or deepen an existing H premium; absolute-gap narrowing is counted separately. Initial premium is also part of the change metric mechanically, which complicates interpretation. Spearman uses ranks. Confidence intervals resample pairs 1,000 times; permutation p-values use 1,999 draws and are adjusted across six tests.')
    doc.add_page_break()
    doc.add_paragraph('Why some premiums remain large',style='Heading 1')
    doc.add_paragraph('Large remaining A premiums occur across chemicals, electronics, instruments and futures brokers. This breadth argues against a single semiconductor-specific explanation. A company can also converge partly and still retain a large absolute gap. The examples below are first-month closing premiums, not today\'s live premiums.')
    largest=sorted(usable,key=lambda r:r['premium_1m_pct'],reverse=True)[:7]
    table(doc,['Pair','First month date','A premium','H move','A move'],[[f"{r['code']} {r['name']}",r['date_1m'],pct(r['premium_1m_pct']),pct(r['h_return_1m_pct']),pct(r['a_return_1m_pct'])] for r in largest],[2.65,1.25,.9,.9,.9])
    doc.add_paragraph('The arithmetic identifies the price leg responsible for a change, but investor motives require additional evidence. Segmented investor access can sustain different valuations. H-line liquidity, stock-borrow costs and investable free float can restrict trades that would otherwise narrow the gap. Stock Connect eligibility, IPO price discovery, stabilisation and company news can change demand over time. Their timing and effects differ by pair; none is established as the cause of GigaDevice or Montage\'s first-month move by this price dataset.')
    doc.add_paragraph('What to test next',style='Heading 2')
    for text in ['Use admission dates and daily Stock Connect turnover or flows to test changes around access events.','Measure A and H tradable float, turnover and borrow availability at the H listing date rather than using current total-company cap.','Match stabilisation notices, lock-up releases and company announcements to daily price-leg contributions. Control for the sector, listing period and starting premium.']:
        doc.add_paragraph(text,style='List Bullet')
    doc.add_paragraph('Sector comparisons',style='Heading 2')
    grouped=defaultdict(list)
    for r in usable:grouped[r['sector']].append(r)
    table(doc,['Sector','n','Median Day 1','Median first month'],[[sector,len(rs),pct(median(r['premium_day1_pct'] for r in rs)),pct(median(r['premium_1m_pct'] for r in rs))] for sector,rs in sorted(grouped.items())],[2.7,.6,1.7,1.9])
    doc.add_paragraph('These are medians of each period, not the median of within-pair changes. Small groups and different listing dates prevent causal sector comparisons. Semiconductor peer groups now separate compute and AI, memory and interface, analog and power, image and display, equipment and materials, foundries, and diversified businesses. Product exposure, profitability and legal listing regimes should remain separate comparison attributes.')
    doc.add_page_break()
    doc.add_paragraph('Methods and sources',style='Heading 1')
    doc.add_paragraph('H and A use unadjusted daily traded closes. The latest prior known A and FX observations are aligned to each H session, with maximum carry of 10 and 5 calendar days. Quote dates are retained. Different closing times mean these are dated relative prices, not simultaneous executable arbitrage quotes. No dividends are included. Raw price-leg movements around share-scale actions require care.')
    doc.add_paragraph('Offer reference uses A and FX closes before H debut divided by the H offer price. It is not a premium measured at the book pricing timestamp. Day 1 and first month use their own dated closes. Latest paired closing premiums are separate from live Bloomberg ratios in Excel.')
    doc.add_paragraph('BIOCYTOGEN has no aligned A close at its H debut and is excluded from first-month tests. HUAQIN has a first-month H share-scale action; its class conversion needs validation and it is also excluded. Unverified HUAQIN premiums after the action are withheld in the appendix and the live monitor. All 62 pairs remain in the appendix. Historical company-cap and deal-size inputs inherit their source share-count and proceeds basis, including differing gross, net and shoe conventions.')
    doc.add_paragraph('The exact log-ratio change equals the A-price log change plus the FX log change minus the H-price log change. The underlying record for each horizon retains those contributions and a reconciliation residual. These are log points, not percentage points. All calculated bridges reconcile within numerical tolerance. Repeated exploratory tests, sample heterogeneity and unobserved liquidity mean association cannot establish causation.')
    link(doc,'Tencent raw daily price history','https://web.ifzq.gtimg.cn/appstock/app/fqkline/get')
    link(doc,'Yahoo dated CNYHKD closing history','https://finance.yahoo.com/quote/CNYHKD=X/history/')
    link(doc,'AAStocks current A and H market snapshot','https://www.aastocks.com/en/stocks/market/ah.aspx')
    link(doc,'HKEX Stock Connect investor access and eligibility framework','https://www.hkex.com.hk/Mutual-Market/Connect-Hub/Stock-Connect?sc_lang=en')
    doc.add_paragraph('The companion database retains original IPO and market-data provenance. The study observations and correlation inputs are stored in data/batches/ah_study.json, with raw dated history in data/ah_study_cache. Current company-cap snapshot retrieval date is recorded per pair; it is not treated as a historical explanatory variable.')
    sec=doc.add_section(WD_SECTION_START.NEW_PAGE);sec.orientation=WD_ORIENT.LANDSCAPE;sec.page_width=Inches(11.69);sec.page_height=Inches(8.27);sec.left_margin=sec.right_margin=Inches(.55);sec.top_margin=sec.bottom_margin=Inches(.55)
    doc.add_paragraph('All A and H IPO pairs',style='Heading 1')
    doc.add_paragraph(f"Closing observations through {batch['asof']}. Premiums are A over H. Change is first month minus Day 1 in percentage points. N/A means the aligned period is unavailable; review flags override the raw calculation.")
    appendix=[]
    for r in recs:
        eligible=r['eligible_1m'];review='share-scale' in r.get('exclusion','')
        appendix.append([r['code']+' '+r['name'],r['sector'],r['ipo_date'],pct(r.get('premium_day1_pct')),pct(r.get('premium_1m_pct')) if eligible else 'N/A',f"{r['premium_change_1m_pp']:+.1f}" if eligible else 'N/A',pct(r.get('latest_premium_pct')) if not review else 'Withheld',r.get('latest_date','N/A'),'Included' if eligible else 'Class review' if review else 'A absent at H debut'])
    table(doc,['Pair','Sector','H IPO date','Day 1','1 month','Change pp','Latest','H close date','First month status'],appendix,[2.15,1.1,.95,.75,.75,.75,.75,.95,1.4],9.5)
    dest=out/'AH_Premium_Study.docx';doc.save(dest);print(f'Wrote {dest}: {len(recs)} pairs, {len(usable)} first-month cases')
if __name__=='__main__':main()
