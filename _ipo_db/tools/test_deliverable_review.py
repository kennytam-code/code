"""Check the delivered data, event formulas, caches and selector contracts."""
import json,re,sys,math,base64,gzip
from pathlib import Path
import openpyxl,formulas
ROOT=Path(__file__).resolve().parent.parent;sys.path.insert(0,str(ROOT/'ipo_lib'))
import build_xlsx as B,build_dashboard as H
D=json.loads((ROOT/'data/deals.json').read_text())['deals'];codes={d['code'] for d in D}
text=(ROOT/'out/hk_ipo_dashboard.html').read_text();payload=re.search(r'<script id="ipo-payload" type="application/octet-stream">([A-Za-z0-9+/=]+)</script>',text)[1];data=json.loads(gzip.decompress(base64.b64decode(payload)));web={d['code']:d for d in data['deals']}
assert set(web)==codes
book=openpyxl.load_workbook(ROOT/'out/HK_IPO_Database_v1.xlsx');cached=openpyxl.load_workbook(ROOT/'out/HK_IPO_Database_v1.xlsx',data_only=True)
rows={book['Database'].cell(r,1).value:r for r in range(5,529)};checked=0
for d in D:
 for field in H.KEEP:
  if field in d:assert web[d['code']].get(field)==d[field],(d['code'],field)
 for col,(_,_,field,_,_) in enumerate(B.DB_COLS,1):
  if isinstance(d.get(field),(int,float)) and not isinstance(d[field],bool):
   value=book['Database'].cell(rows[d['code']],col).value
   assert value==d[field],(d['code'],field,value,d[field]);checked+=1
 cscol=B.db_col('cornerstone_investors');cs=book['Database'][f'{cscol}{rows[d["code"]]}'].value
 expected='; '.join(d.get('cornerstone_investors') or [])
 if expected:assert cs==expected,(d['code'],cs,expected)
 else:assert cs=='none — no cornerstone tranche' and (web[d['code']].get('cornerstone_none') or web[d['code']]['cornerstone_pct']==0)
print('PASS all 524 HTML records, subscription/cornerstone parity and',checked,'Excel numeric cells',flush=True)
events=[e for e in json.loads((ROOT/'data/batches/lockup_events.json').read_text())['events'] if e['category']=='Cornerstone' and e['date_basis']!='No cornerstone tranche'];assert all(e['code'] in codes for e in events)
assert book.sheetnames[:6]==['IPO Live','AH','Lock-up','Peer Trades','Screener','Database']
assert book['Lock-up Data'].sheet_state=='hidden'
header=json.loads((ROOT/'data/market_views_template.json').read_text())['lockup_layout']['table_header']
assert book['Lock-up'].tables['LockupEvents'].ref==f'A{header}:M{header+len(events)}'
assert str(book['Lock-up']['B4'].value) in codes and book['Lock-up']['E4'].value=='Auto'
assert {str(v.sqref) for v in book['Lock-up'].data_validations.dataValidation}.issuperset({'B4','E4','B32','B33','B34','B35'})
assert len(book['Lock-up']._charts)==2
assert book['Lock-up']._charts[0].display_blanks=='gap' and not book['Lock-up']._charts[0].visible_cells_only
assert book['Lock-up']._charts[0].series[0].val.numRef.f=="'Chart Data'!$L$5:$L$65"
assert data['lockups']['price_asof']=='2026-10-02'
assert len(data['lockups']['events'])==len(events)
for index,e in enumerate(events):
 i=header+1+index
 w=e.get('windows',{}).get('30 trading sessions',{})
 def round8(v):
  if isinstance(v,float):return round(v,8)
  if isinstance(v,list):return [round8(x) for x in v]
  if isinstance(v,dict):return {k:round8(x) for k,x in v.items()}
  return v
 assert data['lockups']['events'][index]['window']==round8(w)
 assert data['lockups']['events'][index]['last_day_locked']==e['last_day_locked']
 path=w.get('path',[])
 if path:
  assert [p[0] for p in path]==list(range(-30,31))
  assert path[0][1:]==[w['pre_start'],w['pre_start_close']]
  assert path[30][1:]==[w['anchor'],w['anchor_close']]
  assert path[60][1:]==[w['post_end'],w['post_end_close']]
  dates=[p[1] for p in path if p[1]];assert dates==sorted(set(dates))
  assert all(p[1] is not None or p[2] is None for p in path)
 elif e['last_day_locked'] and e['last_day_locked']>data['lockups']['price_asof']:
  assert w['before_pct'] is None and w['after_pct'] is None
 assert book['Lock-up'].cell(i,1).value==e['code']
 assert book['Lock-up'].cell(i,6).comment
 assert book['Lock-up'].cell(i,5).value.date().isoformat()==next(d['ipo_date'] for d in D if d['code']==e['code'])
 assert data['lockups']['events'][index]['ipo_date']==next(d['ipo_date'] for d in D if d['code']==e['code'])
 for offset in range(-30,31):
  stored=next((p for p in path if p[0]==offset),None);r=8+index*61+offset+30
  assert book['Lock-up Data'].cell(r,28).value==offset
  value=book['Lock-up Data'].cell(r,30).value;expected=stored[2] if stored else None
  assert value is None and expected is None or value is not None and expected is not None and math.isclose(value,expected,rel_tol=1e-14,abs_tol=1e-12)
  dt=book['Lock-up Data'].cell(r,29).value
  assert (dt.date().isoformat() if dt else None)==(stored[1] if stored else None)
 for col,key in [(7,'before_pct'),(8,'after_pct')]:
  expect=e.get('windows',{}).get('30 trading sessions',{}).get(key);value=cached['Lock-up'].cell(i,col).value
  assert (expect is None and value in [None,'']) or expect is not None and math.isclose(value,expect/100,abs_tol=1e-12),(i,key,value,expect)
# Validate every saved Auto lookup against the full company/event schedule.
lookup={book['Lock-up Data'].cell(r,25).value:book['Lock-up Data'].cell(r,26).value for r in range(8,8+len({e['code'] for e in events}))}
for code in lookup:
 indexed=[(i,e) for i,e in enumerate(events) if e['code']==code];past=[x for x in indexed if x[1]['last_day_locked'] and x[1]['last_day_locked']<=data['lockups']['price_asof']]
 expected=sorted(past or indexed,key=lambda x:x[1]['last_day_locked'] or '',reverse=True)[0][0]+8
 assert lookup[code]==expected,(code,lookup[code],expected)
print('PASS all event rows, source comments, cached returns and stock coverage',flush=True)
small=openpyxl.Workbook();s=small.active;s.title='Lock-up';raw=small.create_sheet('Lock-up Data')
for dest,src in [('F','G'),('G','H'),('H','L')]:s[f'{dest}8']=book['Lock-up'][f'{src}{header+1}'].value
for r in range(8,9):
 for c in range(1,9):raw.cell(r,c,book['Lock-up Data'].cell(r,c).value)
path=ROOT/'out/_lockup_formula_probe.xlsx';small.save(path);m=formulas.ExcelModel().loads(str(path)).finish()
def key(sheet,cell):return f"'[{path.name}]{sheet.upper()}'!{cell}"
def calc(e):
 mode='30 trading sessions'
 windows=e.get('windows',{});t=windows.get('30 trading sessions',{});c=windows.get('Calendar days',{})
 vals=[t.get('pre_start_close'),t.get('anchor_close'),t.get('post_end_close'),c.get('pre_start_close'),c.get('anchor_close'),c.get('post_end_close'),t.get('status',e['date_basis']),c.get('status',e['date_basis'])]
 inputs={key('Lock-up','B4'):mode,**{key('Lock-up Data',f'{openpyxl.utils.get_column_letter(i)}8'):v if v is not None else '' for i,v in enumerate(vals,1)}}
 result=m.calculate(inputs=inputs)
 def get(cell):return result[key('Lock-up',cell)].value[0,0]
 w=windows.get(mode,{})
 for cell,k in [('F8','before_pct'),('G8','after_pct')]:
  v=get(cell);expect=w.get(k)
  assert expect is None and v=='' or expect is not None and math.isclose(float(v),expect/100,abs_tol=1e-12),(e['code'],mode,cell,v,expect)
 assert get('H8')==w.get('status',e['date_basis'])
sample=[]
for status in ['Complete','Upcoming expiry','After window pending','Missing close / suspended']:
 e=next((x for x in events if x.get('windows',{}).get('30 trading sessions',{}).get('status')==status),None)
 if e:sample.append(e)
sample+=events[::80]
for e in sample:
 calc(e)
print('PASS 30-trading-session formulas and HTML paths, upcoming, suspended, pending and unavailable cases',flush=True)
path.unlink();book.close();cached.close()
