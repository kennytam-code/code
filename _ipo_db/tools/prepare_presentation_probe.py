"""Truthful small projections of delivered cells/charts; never loads full workbook."""
import io,sys,zipfile,posixpath,re,xml.etree.ElementTree as E
from pathlib import Path
import openpyxl
from openpyxl.chart.data_source import AxDataSource,StrRef,NumRef,NumDataSource
from openpyxl.chart.series import SeriesLabel
from openpyxl.chart.axis import TextAxis
ROOT=Path(__file__).resolve().parent.parent;sys.path.insert(0,str(ROOT/'ipo_lib'))
from patch_peer_workbook import maps,q,N,R,P,T
with zipfile.ZipFile(ROOT/'out/HK_IPO_Database_v1.xlsx') as z:
 _,_,sheets=maps({p:z.read(p) for p in ['xl/workbook.xml','xl/_rels/workbook.xml.rels']})
 for name,limit in [('Lock-up',61),('AH',92),('Screener',80),('Peer Trades',45)]:
  sp=sheets[name];doc=E.fromstring(z.read(sp));mini={};data=doc.find(q('sheetData'))
  for row in list(data):
   if int(row.get('r'))>limit:data.remove(row);continue
   for cell in list(row):
    col=re.match('[A-Z]+',cell.get('r'))[0]
    if openpyxl.utils.column_index_from_string(col)>15:row.remove(cell);continue
    f=cell.find(q('f'))
    if f is not None:cell.remove(f)
    if cell.get('t')=='str':
     v=cell.find(q('v'));value=v.text or '' if v is not None else ''
     if v is not None:cell.remove(v)
     cell.set('t','inlineStr');E.SubElement(E.SubElement(cell,q('is')),q('t')).text=value
  for tag in ['autoFilter','dataValidations','conditionalFormatting','hyperlinks','tableParts','legacyDrawing']:
   for node in list(doc.findall(q(tag))):doc.remove(node)
  dim=doc.find(q('dimension'))
  if dim is not None:dim.set('ref','A1:O'+str(limit))
  rp=sp.replace('/worksheets/','/worksheets/_rels/')+'.rels'
  mini={'xl/styles.xml':z.read('xl/styles.xml'),'xl/worksheets/sheet1.xml':E.tostring(doc)}
  def copy_related(path):
   if path in mini:return
   mini[path]=z.read(path);rel=posixpath.dirname(path)+'/_rels/'+posixpath.basename(path)+'.rels'
   if rel in z.namelist():
    mini[rel]=z.read(rel)
    for x in E.fromstring(mini[rel]):
     if x.get('TargetMode')=='External':continue
     target=x.get('Target');target=target.lstrip('/') if target.startswith('/') else posixpath.normpath(posixpath.join(posixpath.dirname(path),target))
     copy_related(target)
  if rp in z.namelist():
   rel=E.fromstring(z.read(rp))
   for x in list(rel):
    if x.get('Type')!=R+'/drawing':rel.remove(x);continue
    target=x.get('Target');target=target.lstrip('/') if target.startswith('/') else posixpath.normpath(posixpath.join(posixpath.dirname(sp),target));copy_related(target)
   mini['xl/worksheets/_rels/sheet1.xml.rels']=E.tostring(rel)
  if 'xl/sharedStrings.xml' in z.namelist():mini['xl/sharedStrings.xml']=z.read('xl/sharedStrings.xml')
  w=E.Element(q('workbook'));E.SubElement(E.SubElement(w,q('sheets')),q('sheet'),{'name':name,'sheetId':'1','{'+R+'}id':'sheet'});mini['xl/workbook.xml']=E.tostring(w)
  rr=E.Element('{'+P+'}Relationships')
  for k,t,target in [('sheet','worksheet','/xl/worksheets/sheet1.xml'),('styles','styles','/xl/styles.xml')]:E.SubElement(rr,'{'+P+'}Relationship',{'Id':k,'Type':R+'/'+t,'Target':target})
  if 'xl/sharedStrings.xml' in mini:E.SubElement(rr,'{'+P+'}Relationship',{'Id':'strings','Type':R+'/sharedStrings','Target':'/xl/sharedStrings.xml'})
  mini['xl/_rels/workbook.xml.rels']=E.tostring(rr)
  rr=E.Element('{'+P+'}Relationships');E.SubElement(rr,'{'+P+'}Relationship',{'Id':'book','Type':R+'/officeDocument','Target':'xl/workbook.xml'});mini['_rels/.rels']=E.tostring(rr)
  content=E.fromstring(z.read('[Content_Types].xml'))
  for x in list(content):
   if x.tag.endswith('Override') and x.get('PartName','').lstrip('/') not in mini:content.remove(x)
  E.SubElement(content,'{'+T+'}Override',{'PartName':'/xl/worksheets/sheet1.xml','ContentType':'application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml'});mini['[Content_Types].xml']=E.tostring(content)
  buffer=io.BytesIO()
  with zipfile.ZipFile(buffer,'w',zipfile.ZIP_DEFLATED) as small:
   for p,raw in mini.items():small.writestr(p,raw)
  book=openpyxl.load_workbook(buffer)
  for ci,ch in enumerate(book[name]._charts):
   helper=book.create_sheet('Chart '+str(ci))
   for j,ser in enumerate([s for group in ch._charts for s in group.series]):
    def values(ref):
     if ref is None:return {}
     if ref.numRef is not None and ref.numRef.numCache is not None:return {p.idx:float(p.v) for p in ref.numRef.numCache.pt}
     if hasattr(ref,'strRef') and ref.strRef is not None and ref.strRef.strCache is not None:return {p.idx:p.v for p in ref.strRef.strCache.pt}
     return {}
    xs=values(ser.xVal if ser.xVal is not None else ser.cat);ys=values(ser.yVal if ser.yVal is not None else ser.val);count=max(max(xs,default=0),max(ys,default=0))+1
    for i in range(count):
     x=xs.get(i);label=(x if isinstance(x,str) else openpyxl.utils.datetime.from_excel(x).strftime('%d-%b-%y') if x and x>30000 else str(int(x)) if x is not None else None)
     helper.cell(i+2,j*2+1,x if ser.xVal is not None else label);helper.cell(i+2,j*2+2,ys.get(i))
    a=openpyxl.utils.get_column_letter(j*2+1);b=openpyxl.utils.get_column_letter(j*2+2);ser.cat=AxDataSource(strRef=StrRef(f"'Chart {ci}'!${a}$2:${a}${count+1}"));ser.val=NumDataSource(numRef=NumRef(f"'Chart {ci}'!${b}$2:${b}${count+1}"))
    if ser.xVal is not None:
     ser.xVal=AxDataSource(numRef=NumRef(f=f"'Chart {ci}'!${a}$2:${a}${count+1}"));ser.yVal=NumDataSource(numRef=NumRef(f=f"'Chart {ci}'!${b}$2:${b}${count+1}"))
    if ser.tx and ser.tx.strRef and ser.tx.strRef.strCache:ser.tx=SeriesLabel(v=ser.tx.strRef.strCache.pt[0].v)
   if ch.__class__.__name__=='ScatterChart':ch.x_axis.numFmt='mmm-yy';continue
   old=ch.x_axis;ch.x_axis=TextAxis(axId=old.axId,crossAx=old.crossAx,axPos='b',tickLblPos='low',delete=False,tickLblSkip=max(1,count//5),tickMarkSkip=max(1,count//5));helper.sheet_state='hidden'
  book.save(ROOT/'out'/('_presentation_'+name.replace('-','')+'.xlsx'));book.close()
  print('Prepared actual delivered',name,flush=True)
