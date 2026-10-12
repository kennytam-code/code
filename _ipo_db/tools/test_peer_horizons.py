"""Bounded evaluation of the shipped peer formulas with deliberately gapped data.

Bloomberg responses here are fixtures, never delivered quote stand-ins.
"""
import io,json,re,math,base64,zipfile,sys,xml.etree.ElementTree as E
from pathlib import Path
from datetime import date
import openpyxl,formulas
from openpyxl.formula.translate import Translator
from openpyxl.utils.datetime import to_excel
ROOT=Path(__file__).resolve().parent.parent;sys.path.insert(0,str(ROOT/'ipo_lib'))
from patch_peer_workbook import maps,q
with zipfile.ZipFile(io.BytesIO(base64.b64decode(json.loads((ROOT/'data/peer_views_template.json').read_text())['xlsx_base64']))) as z:
 parts={p:z.read(p) for p in ['xl/workbook.xml','xl/_rels/workbook.xml.rels']};_,_,sheets=maps(parts)
 source={}
 for name in ['Peer Chart','Peer Trades']:
  doc=E.fromstring(z.read(sheets[name]));source[name]={c.get('r'):c.find(q('f')).text for c in doc.iter(q('c')) if c.find(q('f')) is not None}
wb=openpyxl.Workbook();wb.remove(wb.active)
for name in ['Peer Chart','Peer Trades','Peer Data','Peer Live']:wb.create_sheet(name)
def bounded(f):
 f=f.replace('3004','34');return re.sub(r'\$([AB])\$(\d{4,})',lambda m:'$'+m[1]+'$100',f)
for a in ['H2','H5','H6','H3','H4','I2','J2','J3','K2']:wb['Peer Chart'][a]='='+bounded(source['Peer Chart'][a])
for col in 'ABCDEFG':
 f='='+bounded(source['Peer Chart'][col+'5'])
 for r in range(5,35):wb['Peer Chart'][f'{col}{r}']=Translator(f,origin=col+'5').translate_formula(col+str(r))
ps=wb['Peer Trades'];ps['B4']='9856';ps['E4']='3308';ps['H4']='Stored';ps['B5']=to_excel(date(2026,8,25));ps['E5']=to_excel(date(2026,9,25));ps['A9']='Stock A';ps['A10']='Stock B'
dates=[date(2026,9,d) for d in [22,23,24,25]];apr=[10,11,12,13];bdates=[dates[1],dates[3]];bpr=[18,20]
pd=wb['Peer Data'];pd['H3']='9856';pd['K3']=2;pd['L3']=4;pd['H4']='3308';pd['K4']=6;pd['L4']=2
for i,(d,p) in enumerate(zip(dates,apr),2):pd.cell(i,1,to_excel(d));pd.cell(i,2,p)
for i,(d,p) in enumerate(zip(bdates,bpr),6):pd.cell(i,1,to_excel(d));pd.cell(i,2,p)
path=ROOT/'out/_peer_formula_probe.xlsx';wb.save(path);m=formulas.ExcelModel().loads(str(path)).finish()
def key(s,a):return f"'[{path.name}]{s.upper()}'!{a}"
def calc(inputs,outs):
 result=m.calculate(inputs={key(s,a):v for (s,a),v in inputs.items()},outputs=[key('Peer Chart',a) for a in outs]);return {a:result[key('Peer Chart',a)].value[0,0] for a in outs}
outs=['H2','H5','H6','J2','J3','A5','A6','B6','C6','D7','D8'];got=calc({},outs)
assert got['H2']==1 and got['H5']==2 and got['H6']==1,got
assert got['J2']==4 and got['J3']==6 and got['A5']==to_excel(dates[0]),got
assert abs(float(got['B6']))<1e-12 and abs(float(got['C6']))<1e-12,got
assert str(got['D7'])=='#N/A',got
expected=100*(13/11-20/18);assert math.isclose(float(got['D8']),expected,abs_tol=1e-10),got
live={('Peer Trades','H4'):'Bloomberg',('Peer Trades','B4'):'0005',('Peer Trades','E4'):'2888'}
for col,values in [('A',['Date',*[to_excel(d) for d in dates]]),('B',['PX_LAST',*apr]),('D',['Date',*[to_excel(d) for d in bdates]]),('E',['PX_LAST',*bpr])]:
 for r,v in enumerate(values,5):live[('Peer Live',col+str(r))]=v
got=calc(live,outs);assert got['H2']==0 and got['H5']==0 and got['J3']==6,got
assert math.isclose(float(got['D8']),expected,abs_tol=1e-10),got
bad=calc({('Peer Trades','B4'):'0005'},['J3','A5']);assert bad['J3']==0 and bad['A5']=='',bad
later=calc({('Peer Trades','B5'):to_excel(dates[3])},['A5','B5','C5','D5','J2']);assert later['A5']==to_excel(dates[3]) and later['J2']==1 and later['D5']==0,later
path.unlink()
controls=openpyxl.Workbook();cs=controls.active;cs.title='Peer Trades'
for a in ['B5','E5']:cs[a]='='+source['Peer Trades'][a]
cs['H4']='Stored';cs['K4']='1Y';cs['H5']=to_excel(date(2026,7,1));cs['K5']=to_excel(date(2026,9,30))
controls.save(path);cm=formulas.ExcelModel().loads(str(path)).finish()
for horizon,expected_date in [('1M',date(2026,9,2)),('3M',date(2026,7,2)),('6M',date(2026,4,2)),('1Y',date(2025,10,2)),('3Y',date(2023,10,2)),('5Y',date(2021,10,2)),('Custom',date(2026,7,1))]:
 result=cm.calculate(inputs={key('Peer Trades','K4'):horizon});got=result[key('Peer Trades','B5')].value[0,0]
 assert got==to_excel(expected_date),(horizon,got)
path.unlink();print('PASS stored/live parity, unrestricted codes, headered histories, common-date rebasing, gaps and all calendar/custom horizons')
