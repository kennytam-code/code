"""Extract one cached sheet for visual QA without loading the full workbook."""
import io,sys,zipfile,xml.etree.ElementTree as E
from pathlib import Path
import openpyxl
from openpyxl.chart.data_source import AxDataSource,StrRef,NumRef,NumDataSource
from openpyxl.chart.series import SeriesLabel
from openpyxl.chart.axis import TextAxis
ROOT=Path(__file__).resolve().parent.parent;sys.path.insert(0,str(ROOT/'ipo_lib'))
from patch_peer_workbook import maps,q,N,R,P,T
with zipfile.ZipFile(ROOT/'out/HK_IPO_Database_v1.xlsx') as z:
 parts={p:z.read(p) for p in ['xl/workbook.xml','xl/_rels/workbook.xml.rels']};_,_,sheets=maps(parts);sp=sheets['Peer Trades'];doc=E.fromstring(z.read(sp))
 for row in list(doc.find(q('sheetData'))):
  if int(row.get('r'))>89:doc.find(q('sheetData')).remove(row);continue
  for c in row:
   f=c.find(q('f'))
   if f is not None:c.remove(f)
   if c.get('t')=='str':
    v=c.find(q('v'));text=v.text or '' if v is not None else ''
    if v is not None:c.remove(v)
    c.set('t','inlineStr');E.SubElement(E.SubElement(c,q('is')),q('t')).text=text
 mini={'xl/styles.xml':z.read('xl/styles.xml'),'xl/worksheets/sheet1.xml':E.tostring(doc),'xl/worksheets/_rels/sheet1.xml.rels':z.read(sp.replace('/worksheets/','/worksheets/_rels/')+'.rels')}
 for p in z.namelist():
  if 'DeskPeer' in p:mini[p]=z.read(p)
 w=E.Element(q('workbook'));E.SubElement(E.SubElement(w,q('sheets')),q('sheet'),{'name':'Peer Trades','sheetId':'1','{'+R+'}id':'sheet'});mini['xl/workbook.xml']=E.tostring(w)
 rr=E.Element('{'+P+'}Relationships');E.SubElement(rr,'{'+P+'}Relationship',{'Id':'sheet','Type':R+'/worksheet','Target':'/xl/worksheets/sheet1.xml'});E.SubElement(rr,'{'+P+'}Relationship',{'Id':'styles','Type':R+'/styles','Target':'/xl/styles.xml'});mini['xl/_rels/workbook.xml.rels']=E.tostring(rr)
 rr=E.Element('{'+P+'}Relationships');E.SubElement(rr,'{'+P+'}Relationship',{'Id':'book','Type':R+'/officeDocument','Target':'xl/workbook.xml'});mini['_rels/.rels']=E.tostring(rr)
 content=E.fromstring(z.read('[Content_Types].xml'))
 for n in list(content):
  if n.tag.endswith('Override') and n.get('PartName','').lstrip('/') not in mini:content.remove(n)
 E.SubElement(content,'{'+T+'}Override',{'PartName':'/xl/worksheets/sheet1.xml','ContentType':'application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml'});mini['[Content_Types].xml']=E.tostring(content)
 buffer=io.BytesIO()
 with zipfile.ZipFile(buffer,'w',zipfile.ZIP_DEFLATED) as small:
  for p,raw in mini.items():small.writestr(p,raw)
 book=openpyxl.load_workbook(buffer)
for ci,ch in enumerate(book['Peer Trades']._charts):
 helper=book.create_sheet('Chart '+str(ci))
 for j,ser in enumerate(ch.series):
  def values(ref):
   ca=ref.numRef.numCache if ref and ref.numRef else None
   return {p.idx:float(p.v) for p in ca.pt} if ca else {}
  xs=values(ser.cat);ys=values(ser.val);count=max(max(xs,default=0),max(ys,default=0))+1
  for i in range(count):helper.cell(i+2,j*2+1,openpyxl.utils.datetime.from_excel(xs[i]).strftime('%d-%b-%y') if i in xs else None);helper.cell(i+2,j*2+2,ys.get(i))
  a=openpyxl.utils.get_column_letter(j*2+1);b=openpyxl.utils.get_column_letter(j*2+2);ser.cat=AxDataSource(strRef=StrRef(f="'"+helper.title+"'!$"+a+'$2:$'+a+'$'+str(count+1)));ser.val=NumDataSource(numRef=NumRef(f="'"+helper.title+"'!$"+b+'$2:$'+b+'$'+str(count+1)))
  if ser.tx and ser.tx.strRef and ser.tx.strRef.strCache:ser.tx=SeriesLabel(v=ser.tx.strRef.strCache.pt[0].v)
 old=ch.x_axis;ch.x_axis=TextAxis(axId=old.axId,crossAx=old.crossAx,axPos='b',tickLblPos='low',delete=False,tickLblSkip=max(1,count//5),tickMarkSkip=max(1,count//5));helper.sheet_state='hidden'
book.save(ROOT/'out/_peer_visual_probe.xlsx');print('Prepared bounded peer preview from actual delivered cells and chart caches')
