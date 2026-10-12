"""Bounded copies of the final cached workbook for visual QA only."""
from pathlib import Path
from copy import copy,deepcopy
from datetime import datetime
import openpyxl
from openpyxl.chart.axis import TextAxis
from openpyxl.chart.data_source import AxDataSource,StrRef,NumRef,NumDataSource
ROOT=Path(__file__).resolve().parent.parent
src=openpyxl.load_workbook(ROOT/'out/HK_IPO_Database_v1.xlsx');cache=openpyxl.load_workbook(ROOT/'out/HK_IPO_Database_v1.xlsx',data_only=True,read_only=True)
w=openpyxl.Workbook();w.remove(w.active)
for name,rows,cols in [('IPO Live',64,15),('AH',64,15),('Lock-up',60,13),('Peer Trades',89,12),('Screener',38,15),('Database',38,15)]:
 s=src[name];t=w.create_sheet(name)
 for row in s.iter_rows(max_row=rows,max_col=cols):
  for c in row:
   if c.__class__.__name__=='MergedCell':continue
   n=t.cell(c.row,c.column,cache[name][c.coordinate].value)
   for attr in ['font','fill','border','alignment']:setattr(n,attr,copy(getattr(c,attr)))
   n.number_format=c.number_format
 for k,d in s.column_dimensions.items():t.column_dimensions[k]=copy(d);t.column_dimensions[k].parent=t
 for k,d in s.row_dimensions.items():
  if k<=rows:t.row_dimensions[k]=copy(d);t.row_dimensions[k].parent=t
 for m in s.merged_cells.ranges:
  if m.max_row<=rows and m.max_col<=cols:t.merge_cells(str(m))
 t.sheet_view.showGridLines=False
 t._charts=deepcopy(s._charts) if name in ['IPO Live','AH','Lock-up','Peer Trades'] else []
 for ci,ch in enumerate(t._charts):
  helper=w.create_sheet(name.replace('-','')+' Chart '+str(ci));row=2
  for j,ser in enumerate(ch.series):
   def values(ref):
    r=ref.numRef or ref.strRef
    ca=r.numCache if ref.numRef else r.strCache
    return {x.idx:float(x.v) if ref.numRef else x.v for x in ca.pt} if ca else {}
   cats=values(ser.cat or ser.xVal);vals=values(ser.val or ser.yVal);count=max(max(cats,default=0),max(vals,default=0))+1
   for i in range(count):
    x=cats.get(i);y=vals.get(i)
    if ch.x_axis.__class__.__name__=='DateAxis' and isinstance(x,(int,float)):x=openpyxl.utils.datetime.from_excel(x).strftime('%d-%b-%y')
    helper.cell(i+2,j*2+1,x);helper.cell(i+2,j*2+2,y)
   a=openpyxl.utils.get_column_letter(j*2+1);b=openpyxl.utils.get_column_letter(j*2+2)
   if ser.cat is not None:
    ser.cat=AxDataSource(strRef=StrRef(f="'"+helper.title+"'!$"+a+'$2:$'+a+'$'+str(count+1)));ser.val=NumDataSource(numRef=NumRef(f="'"+helper.title+"'!$"+b+'$2:$'+b+'$'+str(count+1)))
   if ser.cat is None:
    ser.xVal=AxDataSource(numRef=NumRef(f="'"+helper.title+"'!$"+a+'$2:$'+a+'$'+str(count+1)));ser.yVal=NumDataSource(numRef=NumRef(f="'"+helper.title+"'!$"+b+'$2:$'+b+'$'+str(count+1)))
  if ch.x_axis.__class__.__name__=='DateAxis':
   old=ch.x_axis;ch.x_axis=TextAxis(axId=old.axId,crossAx=old.crossAx,tickLblSkip=max(1,count//5),tickMarkSkip=max(1,count//5))
  helper.sheet_state='hidden'
w.save(ROOT/'out/_final_visual_probe.xlsx')
print('Extracted truthful cached charts and cell styles for six bounded previews')
