"""Compute delivered scalar formulas using explicitly simulated Bloomberg data."""
import json,math,re
from pathlib import Path
from datetime import datetime,date
import openpyxl,formulas
from openpyxl.utils.datetime import to_excel
ROOT=Path(__file__).resolve().parent.parent
src=openpyxl.load_workbook(ROOT/'out/HK_IPO_Database_v1.xlsx')
assert src.sheetnames[:2]==['IPO Live','AH']
assert src['Screener'].freeze_panes=='C1' and 'AH Study' not in src.sheetnames
assert src['IPO Live'].tables['IPOLiveQuotes'].ref=='A7:E721'
assert src['AH'].tables['AHLivePairs'].ref=='A7:H75'
assert src['AH']['E4'].value=='Stored'
assert any(str(d.sqref)=='E4' and d.formula1=='"Stored,Bloomberg"' for d in src['AH'].data_validations.dataValidation)
for s in ['Market Data','Live History']:
 for row in src[s]:
  for c in row:
   if c.data_type=='f' and ('BDP(' in c.value or 'BDH(' in c.value) and (s=='Live History' or c.column>2):assert 'AH!$E$4' in c.value,c.coordinate
for r in range(8,722):
 assert src['IPO Live'].cell(r,4).value==f'=IF($A{r}="","",BDP($A{r},"LAST_PRICE"))'
 assert 'PX_YEST_CLOSE' in src['Market Data'].cell(r,2).value
for name in ['BBG Verify']:
 fs=[c.value for row in src[name] for c in row if c.data_type=='f']
 assert not any('INDEX(BDH(' in f or 'run on terminal' in f for f in fs)
 assert any('BDH(' in f for f in fs)
wb=openpyxl.Workbook();wb.remove(wb.active)
starts=[src['Market Data'].cell(r,24).value for r in range(3,71)]
for name in ['IPO Live','AH','Market Data','AH History','Live History','Chart Data','AH Delta']:
 w=wb.create_sheet(name)
 for row in src[name]:
  for c in row:
   if c.value is None:continue
   if name=='AH Delta' and c.column not in [12,13]:continue
   if name=='IPO Live' and not (c.column<=5 or (c.column<=9 and (c.row<=14 or 33<=c.row<=44))):continue
   if name=='AH' and c.coordinate not in ['B4','E4','J24','J60'] and c.column!=16:continue
   if name=='Chart Data' and (c.row not in [5,6,7,8] or c.column>9):continue
   if name=='Live History':continue
   if name=='AH History' and not any(st<=c.row<st+3 for st in starts):continue
   v=c.value
   if isinstance(v,str) and 'BDP(' in v:
    v=re.sub(r'BDP\([^()]*\)',f'$AN${c.row}',v) if name=='Market Data' else ''
   if isinstance(v,str) and 'BDH(' in v:v=''
   w.cell(c.row,c.column,v)
md=wb['Market Data'];md['AN2']=1.1
for r in range(8,722):wb['IPO Live'].cell(r,4,20);md.cell(r,2,20)
for r in range(3,71):md.cell(r,40,30)
path=ROOT/'out/_monitor_probe.xlsx';wb.save(path)
model=formulas.ExcelModel().loads(str(path)).finish()
def key(s,c):return f"'[{path.name}]{s.upper()}'!{c}"
def calc(inputs={},outputs=None):
 args={'inputs':{key(s,c):to_excel(v) if isinstance(v,datetime) else v for (s,c),v in inputs.items()}}
 if outputs:args['outputs']=[key(s,c) for s,c in outputs]
 return model.calculate(**args)
def get(s,c):
 for k,v in sol.items():
  if k.upper().endswith(key(s,c).upper()):return v.value[0,0]
 raise AssertionError((s,c))
import sys
sys.path.insert(0,str(ROOT/'ipo_lib'))
from market_universe import histories
records=histories(ROOT)['pairs']
for code,expect in [('3986','3986'),('603986.SS','3986'),('6809','6809'),('688008','6809'),(501,'0501'),('0501 HK Equity','0501'),('unknown',None)]:
 sol=calc({('AH','B4'):code});idx=get('AH','P3')
 if expect:
  assert idx>0 and md.cell(int(idx)+2,16).value==expect,(code,idx)
  rec=next(x for x in records if x['code']==expect)
  assert get('Chart Data','A5')==to_excel(datetime.fromisoformat(rec['h_history'][0][0]))
  if rec.get('history') and rec['history'][0][0]==rec['h_history'][0][0]:
   assert math.isclose(float(get('Chart Data','B5')),rec['history'][0][4],abs_tol=1e-10)
 else:assert idx==0 and get('Chart Data','A5')==''
 print('PASS selected code and own listing date',code,flush=True)
for rec in records:
 sol=calc({('AH','B4'):rec['code']},outputs=[('AH','P3'),('AH','P11'),('Chart Data','A5')]);idx=int(get('AH','P3'))
 assert idx>0 and md.cell(idx+2,16).value==rec['code']
 assert get('Chart Data','A5')==to_excel(datetime.fromisoformat(rec['h_history'][0][0]))
 assert int(get('AH','P11'))==len(rec['h_history'])
 start=src['Market Data'].cell(idx+2,24).value
 aligned={x[0]:x for x in rec.get('history',[])}
 last=None
 for offset,(dt,h) in enumerate(rec['h_history']):
  row=start+offset;actual=src['AH History'].cell(row,1).value.date().isoformat()
  assert actual==dt and (last is None or dt>last) and dt>=rec['ipo_date']
  assert src['AH History'].cell(row,2).value==h
  x=aligned.get(dt)
  if x:
   assert x[5]<=dt and x[6]<=dt
   assert (datetime.fromisoformat(dt)-datetime.fromisoformat(x[5])).days<=10
   assert (datetime.fromisoformat(dt)-datetime.fromisoformat(x[6])).days<=5
   assert math.isclose(x[4],x[2]*x[3]/h-1,abs_tol=1e-10)
  last=dt
print('PASS all 68 Excel selectors, complete cached histories, dates and premium maths',flush=True)
sol=calc(outputs=[('Market Data',f'AE{r}') for r in range(3,71)])
for r in range(3,71):
 rec=next(x for x in records if x['code']==md.cell(r,16).value)
 x=next((x for x in rec['history'] if x[0]==rec['h_history'][-1][0]),None)
 v=get('Market Data',f'AE{r}')
 if md.cell(r,26).value==1:assert v=='Class ratio review'
 elif x:assert math.isclose(float(v),x[2]*x[3]/x[1]-1,abs_tol=1e-10)
 else:assert v==''
print('PASS Stored latest premiums for all 68 pairs use their own dated A/H/FX snapshot',flush=True)
# The filed Huaqin ex dates differ. Test both price modes at each boundary,
# including a carried pre-ex A quote after its own ex date.
for mode in ['Stored','Bloomberg']:
 for hd,ad,valid in [('2026-05-21','2026-05-21',True),('2026-05-22','2026-05-22',False),('2026-06-16','2026-06-16',False),('2026-06-17','2026-06-17',True),('2026-06-17','2026-06-16',False)]:
  sol=calc({('AH','B4'):'3296',('AH','E4'):mode,('Chart Data','A5'):datetime.fromisoformat(hd),('Chart Data','C5'):80,('Chart Data','D5'):90,('Chart Data','E5'):1.17,('Chart Data','F5'):datetime.fromisoformat(ad),('Chart Data','G5'):datetime.fromisoformat(hd)},outputs=[('Chart Data','B5')])
  v=get('Chart Data','B5')
  if valid:assert math.isclose(float(v),90*1.17/80-1,abs_tol=1e-10),(mode,hd,ad,v)
  else:assert str(v)=='#N/A',(mode,hd,ad,v)
print('PASS Huaqin stored/live entitlement-date boundaries and stale pre-ex A carry',flush=True)
sol=calc({('AH','E4'):'Bloomberg',('IPO Live','D8'):22,('Market Data','B9'):0,('IPO Live','D10'):'#N/A'})
assert math.isclose(float(get('IPO Live','E8')),.1,abs_tol=1e-12)
assert get('IPO Live','E9')=='' and get('IPO Live','E10')==''
sec=md['C8'].value
n=sum(md.cell(r,3).value==sec for r in range(8,722))-sum(md.cell(r,3).value==sec for r in [9,10])
sr=next(r for r in range(6,15) if wb['IPO Live'].cell(r,7).value==sec)
assert math.isclose(float(get('IPO Live',f'H{sr}')),.1/n,abs_tol=1e-12)
assert get('IPO Live',f'I{sr}')==n
sol=calc({('AH','E4'):'Bloomberg',('IPO Live','D8'):18});assert math.isclose(float(get('IPO Live','E8')),-.1,abs_tol=1e-12)
print('PASS daily returns, missing quotes, zero prior close and sector averages',flush=True)
assert len(src['IPO Live']._charts)==2
assert any(str(d.sqref)=='B5' for d in src['IPO Live'].data_validations.dataValidation)
assert any(str(d.sqref)=='H34' for d in src['IPO Live'].data_validations.dataValidation)
sol=calc({('IPO Live','B5'):'Subsector'})
assert get('IPO Live','C8')==md['D8'].value
sol=calc({('IPO Live','B5'):'Sector'})
assert get('IPO Live','C8')==md['C8'].value
tax=json.loads((ROOT/'data/taxonomy.json').read_text())['sectors']
for sector,groups in tax.items():
 sol=calc({('IPO Live','H34'):sector},outputs=[('IPO Live',f'G{r}') for r in range(37,45)])
 assert [get('IPO Live',f'G{r}') for r in range(37,37+len(groups))]==[x['label'] for x in groups]
 assert all(get('IPO Live',f'G{r}')=='' for r in range(37+len(groups),45))
print('PASS classification dropdown and all nine sector/subsector chart selectors',flush=True)
sol=calc({('AH','E4'):'Bloomberg',('IPO Live','D8'):18})

for r in range(3,71):
 if md.cell(r,26).value==1:assert get('Market Data',f'AE{r}')=='Class ratio review'
 else:assert math.isclose(float(get('Market Data',f'AE{r}')),30*1.1/(18 if src['IPO Live'].cell(8,1).value==md.cell(r,17).value else 20)-1,abs_tol=1e-12)
print('PASS 68 current premiums and share-class exclusion',flush=True)
# Headered native outputs: holiday carry, explicit raw-debut check and TODAY append.
ah=wb['AH'];ah['B4']='3986';ah['E4']='Bloomberg';ah['P16']=20;ah['P17']=datetime(2026,1,15)
lh=wb['Live History'];ipo=datetime(2026,1,13)
for col in [1,4,7]:lh.cell(5,col,'Date');lh.cell(5,col+1,'PX_LAST')
for row,dt,h in [(6,ipo,20),(7,datetime(2026,1,14),22),(8,datetime(2026,1,15),24)]:lh.cell(row,1,dt);lh.cell(row,2,h)
for row,dt,a in [(6,datetime(2026,1,12),30),(7,datetime(2026,1,15),36)]:lh.cell(row,4,dt);lh.cell(row,5,a)
for row,dt,fx in [(6,datetime(2026,1,12),1.1),(7,datetime(2026,1,15),1.2)]:lh.cell(row,7,dt);lh.cell(row,8,fx)
wb.save(path);model=formulas.ExcelModel().loads(str(path)).finish();sol=calc()
assert get('AH','P10')==1 and get('AH','P13')==2 and get('AH','P21')==1 and get('AH','P11')==4
for row,h,a,fx in [(5,20,30,1.1),(6,22,30,1.1),(7,24,36,1.2)]:
 assert math.isclose(float(get('Chart Data',f'B{row}')),a*fx/h-1,abs_tol=1e-10)
assert get('Chart Data','A8')==to_excel(datetime.combine(date.today(),datetime.min.time()))
assert math.isclose(float(get('Chart Data','B8')),30*1.1/20-1,abs_tol=1e-10)
print('PASS headers, holiday alignment and live TODAY endpoint',flush=True)
sol=calc({('Live History','D6'):datetime(2026,1,14)});assert str(get('Chart Data','B5'))=='#N/A'
sol=calc({('AH','P17'):datetime.combine(date.today(),datetime.min.time())});assert get('AH','P10')==0
sol=calc({('Live History','B6'):22});assert get('AH','P10')==0
sol=calc({('AH','B4'):'unknown'});assert get('AH','P3')==0 and get('Chart Data','A5')==''
sol=calc({('AH','P9'):1});assert str(get('Chart Data','B5'))=='#N/A'
print('PASS future A price, stale/wrong-debut series, unknown code and class-review gaps',flush=True)
sol=calc({('AH','E4'):'Stored'})
rec=next(x for x in records if x['code']=='3986')
assert get('AH','P10')==0 and get('AH','P21')==0
assert math.isclose(float(get('Chart Data','B5')),rec['history'][0][4],abs_tol=1e-10)
sol=calc({('AH','P16'):22})
assert get('Chart Data','A5')=='' and 'unavailable' in get('AH','J24')
print('PASS switching ignores lingering Bloomberg results in Stored mode; unavailable Bloomberg never silently plots Stored',flush=True)
# Identical dated raw inputs must produce identical premiums and match checks.
first=rec['history'][:3];ah['P16']=first[0][1];ah['P17']=datetime.fromisoformat(first[-1][0])
for row,x in enumerate(first,6):
 for col,v in [(1,datetime.fromisoformat(x[0])),(2,x[1]),(4,datetime.fromisoformat(x[5])),(5,x[2]),(7,datetime.fromisoformat(x[6])),(8,x[3])]:lh.cell(row,col,v)
wb.save(path);model=formulas.ExcelModel().loads(str(path)).finish();sol=calc()
for row,x in enumerate(first,5):
 assert math.isclose(float(get('Chart Data',f'B{row}')),x[4],abs_tol=1e-10)
 assert get('Chart Data',f'I{row}')==1
assert 'closes agree' in get('AH','J60')
sol=calc({('Live History','E7'):first[1][2]+1})
assert get('Chart Data','I6')==-1 and 'dates differ' in get('AH','J60')
print('PASS same-date Stored / Bloomberg parity and deliberate mismatch detection',flush=True)
c=next(c for c in src['AH']._charts if len(c.series)==1)
assert c.series[0].cat.numRef.f=="'Chart Data'!AHChartDates"
assert c.series[0].val.numRef.f=="'Chart Data'!AHChartPremium"
assert c.x_axis.__class__.__name__=='DateAxis'
cache=c.series[0].val.numRef.numCache
assert cache is not None and cache.ptCount==len(next(x for x in records if x['code']=='3986')['h_history'])
assert len(json.loads((ROOT/'data/taxonomy.json').read_text())['sectors']['Technology'])==6
print('PASS dynamic native chart, six Technology groups and column-only Screener freeze',flush=True)
path.unlink()
