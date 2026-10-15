"""Evaluate shipped comparison formulae with headers, stale data and missing quotes.
Bloomberg inputs in this test are simulated; no synthetic quotes are delivered.
"""
import json,math,sys
from pathlib import Path
from datetime import date,datetime,timedelta
import openpyxl,formulas
from openpyxl.utils.datetime import to_excel
ROOT=Path(__file__).resolve().parent.parent
sys.path.insert(0,str(ROOT/'ipo_lib'))
from market_universe import histories
from corp_actions import load_actions,bbg_factor
src=openpyxl.load_workbook(ROOT/'out/HK_IPO_Database_v1.xlsx')
cached=openpyxl.load_workbook(ROOT/'out/HK_IPO_Database_v1.xlsx',data_only=True)
meta=json.loads((ROOT/'data/desk_chart_data.json').read_text());recs={r['code']:r for r in histories(ROOT)['pairs']}
wb=openpyxl.Workbook();wb.remove(wb.active)
needed=['AH','AH Delta','AH Recent','AH Compare','AH Compare Live','AH History','Market Data','Peer Trades','Peer Data','Peer Live','Peer Chart']
for name in needed:wb.create_sheet(name)
def copy(name,addr):
 c=src[name][addr];v=c.value
 if isinstance(v,str) and ('BDP(' in v or 'BDH(' in v):v=''
 wb[name][addr]=v
for r in range(3,71):
 for c in range(1,15):copy('AH Delta',f'{openpyxl.utils.get_column_letter(c)}{r}')
 for c in [16,17,18]:copy('Market Data',f'{openpyxl.utils.get_column_letter(c)}{r}')
for a in ['B4','E4']:copy('AH',a)
wb['AH']['P3']=1
for a in ['K28','K29','K30']:copy('AH',a)
for k in range(4):
 b=k*9
 for r in [2,3,4,5,6,7,8,20,500,1504,2204]:
  for c in range(b+1,b+10):copy('AH Compare',f'{openpyxl.utils.get_column_letter(c)}{r}')
for p in meta['pairs']:
 hs=recs[p['code']]['h_history'];start=p['cache_start']
 for off in sorted(set([0,1,2,len(hs)-1])):
  for c in range(1,7):copy('AH History',f'{openpyxl.utils.get_column_letter(c)}{start+off}')
path=ROOT/'out/_desk_compare_probe.xlsx';wb.save(path)
m=formulas.ExcelModel().loads(str(path)).finish()
def key(s,c):return f"'[{path.name}]{s.upper()}'!{c}"
def calc(inputs,outs):
 result=m.calculate(inputs={key(s,c):to_excel(v) if isinstance(v,(date,datetime)) else v for (s,c),v in inputs.items()},outputs=[key(s,c) for s,c in outs])
 return {c:result[key(s,c)].value[0,0] for s,c in outs}
for index,p in enumerate(meta['pairs'],1):
 start=p['cache_start'];hs=recs[p['code']]['h_history'];inp={('AH','P3'):index};outs=[('AH Compare','A5'),('AH Compare','A6'),('AH Compare','B5')]
 got=calc(inp,outs)
 assert got['A5']==to_excel(date.fromisoformat(hs[0][0])) and got['B5']==hs[0][1]
 if len(hs)>1:assert got['A6']==to_excel(date.fromisoformat(hs[1][0])) and got['A6']>got['A5']
 # Check an interior long-history row after filling down, not merely template exemplars.
 if len(hs)>495:
  got=calc({**inp,('AH History',f'A{start+495}'):date.fromisoformat(hs[495][0])},[('AH Compare','A500')])
  assert got['A500']==to_excel(date.fromisoformat(hs[495][0])),p['code']
print('PASS stored dates advance for all 68 A/H selectors',flush=True)
# Headered BDH replies and a today live endpoint. Deliberately not a production snapshot.
p=meta['pairs'][0];ipo=date.fromisoformat(p['ipo_date']);now=date.today();last=now-timedelta(days=3)
inp={('AH','E4'):'Bloomberg',('AH','P3'):1}
for col,values in [('A',['Date',ipo,last]),('B',['PX_LAST',p['raw_debut'],80]),('D',['Date',ipo,last]),('E',['PX_LAST',90,90]),('G',['Date',ipo,last]),('H',['PX_LAST',1.17,1.17])]:
 for r,v in enumerate(values,5):inp[('AH Compare Live',f'{col}{r}')]=v
for c,v in [('AC',95),('AD',82),('AK',1.17),('AE',95*1.17/82-1)]:inp[('Market Data',c+'3')]=v
for c in ['J','K','L','M']:inp[('AH Delta',c+'3')]=0
outs=[('AH Compare','I2'),('AH Compare','I3'),('AH Compare','A5'),('AH Compare','A6'),('AH Compare','A7'),('AH Compare','G6'),('AH Compare','G7')]
got=calc(inp,outs)
assert got['I2']==1 and got['I3']==3 and got['A5']==to_excel(ipo) and got['A6']==to_excel(last) and got['A7']==to_excel(now),got
assert math.isclose(float(got['G6']),90*1.17/80-1,abs_tol=1e-12) and math.isclose(float(got['G7']),95*1.17/82-1,abs_tol=1e-12)
assert calc({**inp,('AH Compare Live','B6'):p['raw_debut']*2},[('AH Compare','I2'),('AH Compare','A5')])=={'I2':0,'A5':''}
stale=calc({**inp,('AH Compare Live','A7'):date(2025,1,1)},[('AH Compare','I2')]);assert stale['I2']==0
# A/H references: calendar day / calendar month, aligned dated raw quotes.
start=5;target=now-timedelta(days=7)
for col,vs in [('A',['Date',target]),('B',['PX_LAST',80]),('D',['Date',target]),('E',['PX_LAST',90]),('G',['Date',target]),('H',['PX_LAST',1.17])]:
 for r,v in enumerate(vs,start):inp[('AH Recent',f'{col}{r}')]=v
assert math.isclose(float(calc(inp,[('AH Delta','D3')])['D3']),90*1.17/80-1,abs_tol=1e-12)
print('PASS headered A/H history, exact starting date, today endpoint, stale/wrong security rejection and calendar reference',flush=True)
f=load_actions(ROOT)
assert math.isclose(.56/bbg_factor('2147',f,as_of='2026-10-12'),8.1641025641,abs_tol=1e-9)
assert math.isclose(.56/bbg_factor('2147',f,as_of='2026-08-26'),11.2,abs_tol=1e-10)
for name in ['IPO Live','AH','Peer Trades','Screener','Database']:assert src[name].sheet_properties.tabColor is not None
assert 'Analogs' not in src.sheetnames and 'Verify (BBG)' not in src.sheetnames
assert src['Scoring Stats'].sheet_state=='hidden'
print('PASS rights-adjusted 2147 price basis with dated terminal fixture; important tab colours and combined verification',flush=True)
path.unlink()
