#!/usr/bin/env python3
"""Off-terminal checks for oi_charts.py: a fake Bloomberg terminal drives the real code path.

    python3 test_oi_charts.py            PASS/FAIL per check, exit 1 on any FAIL
    python3 test_oi_charts.py --demo     write demo_OI_charts.xlsx from the fake (open it in Excel)

The fake speaks enough blpapi to exercise everything oi_charts.py does: batched
ReferenceDataRequests split across PARTIAL_RESPONSE / RESPONSE events, HistoricalDataRequests
with the single-element securityData shape, securityError for unknown tickers, rejected
mnemonics, a silent terminal, a rejected request - and a universe with one-digit tickers
for live contracts, two-digit for expired ones, a wrong-decade one-digit ticker, a
quarterly-only product and a contract that resolves but has no prints yet.
"""
import datetime as dt
import importlib.util
import io
import os
import sys
import tempfile
import zipfile
from contextlib import redirect_stdout

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import oi_charts as M                                    # noqa: E402

from openpyxl import load_workbook                       # noqa: E402
from openpyxl.chart import LineChart                     # noqa: E402
from openpyxl.chart.axis import DateAxis                 # noqa: E402
from openpyxl.utils import get_column_letter             # noqa: E402

TEST_TODAY = dt.date(2026, 9, 17)
FAKE_PRODUCTS = [('HI', 'HSI'), ('XP', 'AS51'), ('QZ', 'SIMSCI')]
QUARTERLY = (3, 6, 9, 12)
HOLIDAYS = {(1, 1), (4, 4), (12, 25)}
DAY = dt.timedelta(days=1)


# ========================================================== FAKE BLPAPI ===
class FakeEvent:
    TIMEOUT, RESPONSE, PARTIAL_RESPONSE, REQUEST_STATUS = 0, 1, 2, 3

    def __init__(self, messages, event_type=1):
        self._messages, self._type = messages, event_type

    def __iter__(self):
        return iter(self._messages)

    def eventType(self):
        return self._type


class FakeName(str):
    pass


class FakeElement:
    """Scalar, array-of-scalars, or complex (named children), like the real thing."""

    def __init__(self, name, value=None, children=None, array=None):
        self._name, self._value = name, value
        self._children = children                 # list[(name, FakeElement)] when complex
        self._array = array                       # list[FakeElement] when an array

    def name(self):
        return FakeName(self._name)

    def isArray(self):
        return self._array is not None

    def isComplexType(self):
        if self._array is not None:
            return bool(self._array) and self._array[0]._children is not None
        return self._children is not None

    def isNull(self):
        return self._value is None and self._children is None and self._array is None

    def numValues(self):
        return len(self._array or [])

    def numElements(self):
        return len(self._children or [])

    def getValueAsElement(self, i):
        return self._array[i]

    def getElement(self, key):
        if isinstance(key, int):
            return self._children[key][1]
        for n, el in (self._children or []):
            if n == key:
                return el
        raise KeyError(key)

    def hasElement(self, key):
        return any(n == key for n, _ in (self._children or []))

    def getValue(self, i=None):
        return self._array[i]._value if i is not None else self._value

    def getValueAsString(self):
        return '' if self._value is None else str(self._value)

    def getElementAsString(self, key):
        v = self.getElement(key)._value
        return '' if v is None else str(v)

    def getElementAsFloat(self, key):
        return float(self.getElement(key)._value)

    def getElementAsDatetime(self, key):
        return self.getElement(key)._value

    def appendValue(self, v):
        self._array.append(FakeElement('item', v))

    def set(self, key, value):
        self._children = [(n, e) for n, e in (self._children or []) if n != key]
        self._children.append((key, FakeElement(key, value)))


def fake_complex(name, pairs):
    return FakeElement(name, children=[(n, e if isinstance(e, FakeElement) else FakeElement(n, e))
                                       for n, e in pairs])


def fake_array(name, items):
    return FakeElement(name, array=list(items))


class FakeMessage:
    def __init__(self, root):
        self._root = root

    def hasElement(self, key):
        return self._root.hasElement(key)

    def getElement(self, key):
        return self._root.getElement(key)

    def toString(self):
        return '<fake message %s>' % self._root.name()


class FakeRequest:
    def __init__(self, operation):
        self.operation = operation
        self.settings = {}
        self._root = FakeElement('request', children=[
            ('securities', fake_array('securities', [])),
            ('fields', fake_array('fields', [])),
        ])

    def getElement(self, key):
        return self._root.getElement(key)

    def set(self, key, value):
        self.settings[key] = value
        self._root.set(key, value)

    @property
    def securities(self):
        return [e._value for e in self._root.getElement('securities')._array]

    @property
    def fields(self):
        return [e._value for e in self._root.getElement('fields')._array]


class FakeService:
    def __init__(self, name):
        self._name = name

    def createRequest(self, operation):
        return FakeRequest(operation)


class FakeSessionOptions:
    def setServerHost(self, *a):
        pass

    def setServerPort(self, *a):
        pass


# ------------------------------------------------------------ fake market ---
def weekday_on_or_before(d):
    while d.weekday() >= 5:
        d -= DAY
    return d


def last_trade(y, m):
    """Second-last weekday on or before the 28th: always inside the month."""
    return weekday_on_or_before(weekday_on_or_before(dt.date(y, m, 28)) - DAY)


def listing(y, m):
    days = 1100 if m == 12 else (400 if m in QUARTERLY else 90)   # Dec listed 3y out, serials 3m
    return last_trade(y, m) - dt.timedelta(days=days)


def spec(y, m, **over):
    s = dict(year=y, month=m, last_trade=last_trade(y, m), listing=listing(y, m))
    s.update(over)
    return s


def build_universe():
    """ticker -> contract spec.  Which forms exist follows the real Bloomberg behaviour."""
    u = {}
    for root, _ in FAKE_PRODUCTS:
        for y, m in M.contract_months(2024, 2026):
            if root == 'XP' and m not in QUARTERLY:
                continue                                            # (v)  quarterly-only product
            t1, t2 = M.candidate_tickers(root, y, m)
            s = spec(y, m)
            expired = s['last_trade'] < TEST_TODAY
            if root == 'HI' and (y, m) == (2026, 1):                # (iv) one-digit form is the 2036 contract
                u[t1], u[t2] = spec(2036, 1), s
            elif root == 'QZ' and (y, m) == (2026, 2):              # (iv) only a wrong-decade form exists
                u[t1] = spec(2036, 2)
            elif root == 'QZ' and (y, m) == (2026, 10):             # (vi) resolves, not listed yet
                u[t1] = spec(y, m, listing=dt.date(2026, 10, 1))
            elif root == 'HI' and (y, m) in ((2026, 8), (2026, 12)):   # (iii) both forms valid
                u[t1], u[t2] = s, s
            else:                                                   # (i) expired 2-digit / (ii) live 1-digit
                u[t2 if expired else t1] = s
    return u


UNIVERSE = build_universe()
TICKER_ID = {t: i for i, t in enumerate(sorted(UNIVERSE))}


def fake_oi(ticker, d):
    """Deterministic, distinct per ticker, rising with time - a misaligned cell cannot match."""
    s = UNIVERSE[ticker]
    return float(TICKER_ID[ticker] * 100000 + (d - s['listing']).days * 3 + 100)


def is_session(d):
    return d.weekday() < 5 and (d.month, d.day) not in HOLIDAYS


def first_session_on_or_after(d):
    while not is_session(d):
        d += DAY
    return d


def last_session_on_or_before(d):
    while not is_session(d):
        d -= DAY
    return d


def fake_rows(ticker, start, end):
    s = UNIVERSE.get(ticker)
    if s is None:
        return []
    d, hi, out = max(s['listing'], start), min(s['last_trade'], end), []
    while d <= hi:
        if is_session(d):
            out.append((d, fake_oi(ticker, d)))
        d += DAY
    return out


def security_error(sec):
    return fake_complex('securityData', [
        ('security', sec),
        ('securityError', fake_complex('securityError', [
            ('category', 'BAD_SEC'), ('message', 'Unknown/Invalid Security [nid:1234]')])),
    ])


class FakeSession:
    instances = []

    def __init__(self, options=None):
        self._queue, self.log, self.stopped = [], [], False
        FakeSession.instances.append(self)

    def start(self):
        return True

    def stop(self):
        self.stopped = True

    def openService(self, name):
        return name == '//blp/refdata'

    def getService(self, name):
        return FakeService(name)

    def sendRequest(self, request):
        self.log.append(dict(op=request.operation, securities=list(request.securities),
                             fields=list(request.fields), settings=dict(request.settings)))
        if request.operation == 'ReferenceDataRequest':
            self._queue = self._ref_events(request)
        elif request.operation == 'HistoricalDataRequest':
            self._queue = self._hist_events(request)
        else:
            raise ValueError(request.operation)

    def nextEvent(self, timeout=0):
        return self._queue.pop(0) if self._queue else FakeEvent([], FakeEvent.TIMEOUT)

    @staticmethod
    def _ref_item(sec, wanted):
        s = UNIVERSE.get(sec)
        if s is None:
            return security_error(sec)
        mon = M.MONTH_ABBR[s['month'] - 1]
        known = {'LAST_TRADEABLE_DT': s['last_trade'],
                 'FUT_MONTH_YR': '%s %02d' % (mon.upper(), s['year'] % 100),
                 'NAME': 'FAKE %s FUT %s%02d' % (sec.split()[0], mon, s['year'] % 100)}
        kids = [('security', sec),
                ('fieldData', fake_complex('fieldData', [(k, known[k]) for k in wanted if k in known]))]
        bad = [f for f in wanted if f not in known]
        if bad:
            kids.append(('fieldExceptions', fake_array('fieldExceptions', [
                fake_complex('fieldExceptions', [
                    ('fieldId', f),
                    ('errorInfo', fake_complex('errorInfo', [
                        ('subcategory', 'BAD_FLD'), ('message', 'Invalid Field')])),
                ]) for f in bad])))
        return fake_complex('securityData', kids)

    def _ref_events(self, request):
        items = [self._ref_item(s, request.fields) for s in request.securities]
        half = len(items) // 2
        events = [FakeEvent([], FakeEvent.TIMEOUT)]
        for part, kind in ((items[:half], FakeEvent.PARTIAL_RESPONSE), (items[half:], FakeEvent.RESPONSE)):
            root = FakeElement('ReferenceDataResponse',
                               children=[('securityData', fake_array('securityData', part))])
            events.append(FakeEvent([FakeMessage(root)], kind))
        return events

    @staticmethod
    def _hist_event(sd, kind):
        root = FakeElement('HistoricalDataResponse', children=[('securityData', sd)])   # scalar
        return FakeEvent([FakeMessage(root)], kind)

    def _hist_events(self, request):
        start = dt.datetime.strptime(request.settings['startDate'], '%Y%m%d').date()
        end = dt.datetime.strptime(request.settings['endDate'], '%Y%m%d').date()
        field = request.fields[0]
        events = [FakeEvent([], FakeEvent.TIMEOUT)]
        secs = request.securities
        for k, sec in enumerate(secs):
            last_sec = (k == len(secs) - 1)
            if sec not in UNIVERSE:
                events.append(self._hist_event(
                    security_error(sec), FakeEvent.RESPONSE if last_sec else FakeEvent.PARTIAL_RESPONSE))
                continue
            rows = fake_rows(sec, start, end)
            half = len(rows) // 2
            for pi, part in enumerate((rows[:half], rows[half:])):     # one history, two messages
                fd = fake_array('fieldData', [fake_complex('fieldData', [('date', d), (field, v)])
                                              for d, v in part])
                sd = fake_complex('securityData', [('security', sec), ('fieldData', fd)])
                kind = FakeEvent.RESPONSE if (last_sec and pi == 1) else FakeEvent.PARTIAL_RESPONSE
                events.append(self._hist_event(sd, kind))
        return events


class FakeAPI:
    """Stands in for the blpapi module itself."""
    SessionOptions = FakeSessionOptions
    Session = FakeSession
    Event = FakeEvent


class SilentSession(FakeSession):
    def sendRequest(self, request):
        self._queue = []


class RejectingSession(FakeSession):
    def sendRequest(self, request):
        root = FakeElement('RequestFailure', children=[
            ('reason', fake_complex('reason', [('description', 'Daily capacity reached')]))])
        self._queue = [FakeEvent([FakeMessage(root)], FakeEvent.REQUEST_STATUS)]


class ResponseErrorSession(FakeSession):
    def sendRequest(self, request):
        root = FakeElement('ReferenceDataResponse', children=[
            ('responseError', fake_complex('responseError', [('message', 'Not logged in')]))])
        self._queue = [FakeEvent([FakeMessage(root)], FakeEvent.RESPONSE)]


def api_with(session_cls):
    return type('API', (), dict(SessionOptions=FakeSessionOptions, Session=session_cls, Event=FakeEvent))


# ================================================================ CHECKS ===
FAILS, COUNT = [], [0]


def check(name, cond, detail=''):
    COUNT[0] += 1
    if cond:
        print('PASS  ' + name)
    else:
        print('FAIL  ' + name + ('  -- %s' % (detail,) if detail != '' else ''))
        FAILS.append(name)


def quiet(fn, *a, **k):
    buf = io.StringIO()
    with redirect_stdout(buf):
        out = fn(*a, **k)
    return out, buf.getvalue()


def test_helpers():
    months = M.contract_months(2024, 2026)
    check('36 contract months, Jan 24 first, Dec 26 last',
          len(months) == 36 and months[0] == (2024, 1) and months[-1] == (2026, 12))
    check('candidate forms HI Sep 24', M.candidate_tickers('HI', 2024, 9) == ('HIU4 Index', 'HIU24 Index'))
    check('candidate forms HCT Jan 26', M.candidate_tickers('HCT', 2026, 1) == ('HCTF6 Index', 'HCTF26 Index'))
    check('candidate forms TWT Dec 25', M.candidate_tickers('TWT', 2025, 12) == ('TWTZ5 Index', 'TWTZ25 Index'))
    check('labels', M.month_label(2024, 1) == 'Jan 24' and M.month_label(2026, 12) == 'Dec 26')
    check('years_back clips 29 Feb', M.years_back(dt.date(2024, 2, 29), 2) == dt.date(2022, 2, 28))
    check('years_back plain', M.years_back(dt.date(2026, 9, 28), 2) == dt.date(2024, 9, 28))
    check('as_date datetime', M.as_date(dt.datetime(2026, 1, 28, 0, 0)) == dt.date(2026, 1, 28))
    check('as_date date', M.as_date(dt.date(2026, 1, 28)) == dt.date(2026, 1, 28))
    check('as_date text', M.as_date('2026-01-28T00:00:00') == dt.date(2026, 1, 28))
    check('as_date None / garbage', M.as_date(None) is None and M.as_date('n.a.') is None)
    check('10 real products, distinct legal tab names',
          len({n for _, n in M.PRODUCTS}) == len(M.PRODUCTS) == 10
          and all(M.safe_sheet_name(n) == n for _, n in M.PRODUCTS))
    check('real roots match the request note',
          [r for r, _ in M.PRODUCTS] == ['HI', 'HC', 'HCT', 'KM', 'XP', 'FT', 'TWT', 'FPO', 'HJA', 'QZ'])


def test_resolution():
    bbg = M.Bloomberg(blpapi_module=FakeAPI).connect()
    months = M.contract_months(2024, 2026)
    results = M.resolve_contracts(bbg, FAKE_PRODUCTS, months, TEST_TODAY)
    by = {(n, c.label): c for n, cs in results for c in cs}
    refs = [e for e in bbg.session.log if e['op'] == 'ReferenceDataRequest']
    cands = [t for _, cs in results for c in cs for t in (c.ticker_1, c.ticker_2)]
    requested = [s for e in refs for s in e['securities']]
    check('one reference pass, chunked at REF_CHUNK',
          len(refs) == -(-len(cands) // M.REF_CHUNK) and all(len(e['securities']) <= M.REF_CHUNK for e in refs),
          len(refs))
    check('every candidate requested exactly once', sorted(requested) == sorted(cands))
    check('reference fields as configured', all(e['fields'] == M.REF_FIELDS for e in refs))
    c = by[('HSI', 'Sep 25')]
    check('(i)   expired -> two-digit form', c.ticker == 'HIU25 Index' and c.status == M.OK, c)
    check('(ii)  live -> one-digit form',
          by[('HSI', 'Oct 26')].ticker == 'HIV6 Index' and by[('HSI', 'Sep 26')].ticker == 'HIU6 Index')
    c = by[('HSI', 'Aug 26')]
    check('(iii) both forms valid, expired -> two-digit', c.ticker == 'HIQ26 Index' and c.note == 'both forms valid', c)
    check('(iii) both forms valid, live -> one-digit', by[('HSI', 'Dec 26')].ticker == 'HIZ6 Index')
    c = by[('HSI', 'Jan 26')]
    check('(iv)  wrong-decade one-digit rejected, two-digit used',
          c.ticker == 'HIF26 Index' and c.status == M.OK and c.last_trade.year == 2026, c)
    c = by[('SIMSCI', 'Feb 26')]
    check('(iv)  only a wrong-decade form -> NOT FOUND, note shows the date seen',
          c.status == M.NOT_FOUND and c.ticker == '' and c.last_trade is None and '2036' in c.note, c.note)
    as51 = results[1][1]
    nf = [c for c in as51 if c.status == M.NOT_FOUND]
    check('(v)   quarterly-only product: 24 NOT FOUND, 12 OK',
          len(nf) == 24 and sum(c.status == M.OK for c in as51) == 12, [c.label for c in nf])
    check('(v)   the NOT FOUND months are exactly the serial months', all(c.month not in QUARTERLY for c in nf))
    check('(v)   NOT FOUND note carries Bloomberg\'s message', 'Unknown/Invalid' in by[('AS51', 'Jan 24')].note,
          by[('AS51', 'Jan 24')].note)
    check('(vii) unknown tickers collected, not raised',
          'HIU5 Index' in bbg.bad_securities and 'HIU25 Index' not in bbg.bad_securities)
    check('no field errors on the reference pass', not bbg.field_errors, bbg.field_errors)
    ok = [c for _, cs in results for c in cs if c.status == M.OK]
    check('LAST_TRADEABLE_DT kept as a date, inside the contract month',
          all(isinstance(c.last_trade, dt.date) and (c.last_trade.year, c.last_trade.month) == (c.year, c.month)
              for c in ok))
    c = by[('HSI', 'Sep 25')]
    check('FUT_MONTH_YR / NAME captured', c.fut_month_yr == 'SEP 25' and c.name.startswith('FAKE HIU25'), (c.fut_month_yr, c.name))
    counts = [(n, sum(c.status == M.OK for c in cs)) for n, cs in results]
    check('resolution counts HSI 36 / AS51 12 / SIMSCI 35', counts == [('HSI', 36), ('AS51', 12), ('SIMSCI', 35)], counts)
    return bbg, results


def test_history(bbg, results):
    for _, cs in results:
        for c in cs:
            M.fetch_open_interest(bbg, c, 2, TEST_TODAY)
    by = {(n, c.label): c for n, cs in results for c in cs}
    hist = [e for e in bbg.session.log if e['op'] == 'HistoricalDataRequest']
    resolved = [c for _, cs in results for c in cs if c.status in (M.OK, M.NO_DATA)]
    check('one historical request per resolved contract (83)', len(hist) == len(resolved) == 83, (len(hist), len(resolved)))
    check('each request: one security, OPEN_INT only',
          all(len(e['securities']) == 1 and e['fields'] == ['OPEN_INT'] for e in hist))
    check('settings: DAILY / ACTUAL / ACTIVE_DAYS_ONLY',
          all(e['settings'].get('periodicitySelection') == 'DAILY'
              and e['settings'].get('periodicityAdjustment') == 'ACTUAL'
              and e['settings'].get('nonTradingDayFillOption') == 'ACTIVE_DAYS_ONLY' for e in hist))
    by_ticker = {c.ticker: c for c in resolved}
    bad = []
    for e in hist:
        c = by_ticker[e['securities'][0]]
        exp = (M.years_back(c.last_trade, 2).strftime('%Y%m%d'), min(c.last_trade, TEST_TODAY).strftime('%Y%m%d'))
        got = (e['settings']['startDate'], e['settings']['endDate'])
        if got != exp:
            bad.append((c.ticker, got, exp))
    check('request window = [last trade - 2y, min(last trade, today)] as YYYYMMDD', not bad, bad[:3])
    c = by[('SIMSCI', 'Oct 26')]
    check('(vi)  resolves but no prints -> NO DATA, no rows, window kept for the audit',
          c.status == M.NO_DATA and c.rows == [] and c.req_start is not None and 'no OPEN_INT rows' in c.note, c)
    ok = [c for c in resolved if c.status == M.OK]
    check('every OK contract has rows', all(c.rows for c in ok))
    check('rows sorted, strictly increasing dates',
          all(all(a[0] < b[0] for a, b in zip(c.rows, c.rows[1:])) for c in ok))
    check('rows inside the request window', all(c.req_start <= c.rows[0][0] and c.rows[-1][0] <= c.req_end for c in ok))
    check('rows are sessions only (no filled holidays / weekends)', all(is_session(d) for c in ok for d, _ in c.rows))
    check('values are floats equal to the print of that day',
          all(isinstance(v, float) and v == fake_oi(c.ticker, d) for c in ok for d, v in c.rows))
    check('two messages of one history are accumulated, not overwritten',
          all(len(c.rows) == len(fake_rows(c.ticker, c.req_start, c.req_end)) for c in ok))
    c = by[('HSI', 'Feb 25')]
    check('serial contract: starts on its listing day, ~3 months long',
          c.rows[0][0] == first_session_on_or_after(UNIVERSE[c.ticker]['listing'])
          and (c.rows[-1][0] - c.rows[0][0]).days < 100, (c.rows[0], c.rows[-1]))
    c = by[('HSI', 'Dec 25')]
    check('Dec contract listed 3y out: clipped to the 2-year window',
          c.rows[0][0] == first_session_on_or_after(c.req_start) and c.req_start == M.years_back(c.last_trade, 2),
          (c.rows[0], c.req_start))
    c = by[('HSI', 'Jun 25')]
    check('expired contract: ends on its last trade date', c.rows[-1][0] == last_session_on_or_before(c.last_trade))
    c = by[('HSI', 'Dec 26')]
    check('live contract: ends today', c.rows[-1][0] == last_session_on_or_before(TEST_TODAY))
    c3 = M.Contract(**{k: getattr(by[('HSI', 'Dec 25')], k) for k in
                        ('product', 'root', 'year', 'month', 'label', 'ticker_1', 'ticker_2', 'ticker', 'status', 'last_trade')})
    M.fetch_open_interest(bbg, c3, 3, TEST_TODAY)
    check('--years-back 3 widens the window', c3.req_start == M.years_back(c3.last_trade, 3)
          and bbg.session.log[-1]['settings']['startDate'] == c3.req_start.strftime('%Y%m%d'))
    return results


def test_workbook(results):
    by = {(n, c.label): c for n, cs in results for c in cs}
    tmp = tempfile.mkdtemp()
    path = os.path.join(tmp, 'test_OI.xlsx')
    M.write_workbook(path, results)
    wb = load_workbook(path)
    check('tabs: products in order, then Contracts', wb.sheetnames == ['HSI', 'AS51', 'SIMSCI', M.AUDIT_SHEET], wb.sheetnames)
    with zipfile.ZipFile(path) as z:
        chart_xml = {n: z.read(n) for n in z.namelist() if n.startswith('xl/charts/chart')}
        sheet_xml = {n: z.read(n) for n in z.namelist() if n.startswith('xl/worksheets/sheet')}
    check('no empty <v></v> cells anywhere (the Excel "repair" trigger)',
          all(b'<v></v>' not in x and b'<v/>' not in x for x in sheet_xml.values()))
    check('one chart part per product tab', len(chart_xml) == 3, len(chart_xml))
    check('chart XML: gaps, date axis, no markers, straight lines',
          all(b'dispBlanksAs val="gap"' in x and b'<dateAx>' in x and b'symbol val="none"' in x
              and b'smooth val="0"' in x for x in chart_xml.values()))
    for name, cs in results:
        ws = wb[name]
        dates, found = M.align_product(cs)
        n, N = len(found), len(dates)
        check('%s: row 1 tickers, row 2 labels, expiry order' % name,
              [ws.cell(1, j).value for j in range(2, n + 2)] == [c.ticker for c in found]
              and [ws.cell(2, j).value for j in range(2, n + 2)] == [c.label for c in found]
              and ws['A1'].value == 'Ticker' and ws['A2'].value == 'Date')
        check('%s: no column beyond the %d found contracts' % (name, n), ws.max_column == n + 1, ws.max_column)
        col_a = [ws.cell(i, 1).value for i in range(3, N + 3)]
        col_a = [v.date() if isinstance(v, dt.datetime) else v for v in col_a]
        check('%s: column A = sorted union of dates, nothing below' % name,
              col_a == dates and all(a < b for a, b in zip(dates, dates[1:])) and ws.max_row == N + 2)
        good, why = True, ''
        for j, c in enumerate(found, start=2):
            lk = dict(c.rows)
            for i, d in enumerate(dates, start=3):
                v, exp = ws.cell(i, j).value, lk.get(d)
                if (exp is None and v is not None) or (exp is not None and v != exp):
                    good, why = False, (c.ticker, d, v, exp)
                    break
            if not good:
                break
        check('%s: every cell is the print of that (contract, date), blank elsewhere' % name, good, why)
        check('%s: no NaN cells' % name,
              all(not (isinstance(v, float) and v != v) for row in ws.iter_rows(values_only=True) for v in row))
        check('%s: freeze panes B3' % name, ws.freeze_panes == 'B3')
        check('%s: exactly one chart' % name, len(ws._charts) == 1)
        ch = ws._charts[0]
        check('%s: LineChart with %d series' % (name, n), isinstance(ch, LineChart) and len(ch.series) == n, len(ch.series))
        check('%s: gaps, date axis mmm-yy by month, both axes shown, crossAx wired' % name,
              ch.display_blanks == 'gap' and isinstance(ch.x_axis, DateAxis)
              and ch.x_axis.number_format.formatCode == 'mmm-yy' and ch.x_axis.majorTimeUnit == 'months'
              and ch.x_axis.delete is False and ch.y_axis.delete is False
              and ch.x_axis.axId == 500 and ch.x_axis.crossAx == 100
              and ch.y_axis.axId == 100 and ch.y_axis.crossAx == 500)
        title = ch.title.tx.rich.p[0].r[0].t
        ytitle = ch.y_axis.title.tx.rich.p[0].r[0].t
        check('%s: chart title and OI axis title' % name,
              title == '%s Futures Open Interest' % name and ytitle == 'OI', (title, ytitle))
        check('%s: legend at the bottom' % name, ch.legend is not None and ch.legend.position == 'b')
        ok_series, why = True, ''
        for k, s in enumerate(ch.series):
            col = get_column_letter(k + 2)
            want = ("'%s'!%s2" % (name, col), "'%s'!$%s$3:$%s$%d" % (name, col, col, N + 2), "'%s'!$A$3:$A$%d" % (name, N + 2))
            got = (s.tx.strRef.f, s.val.numRef.f, s.cat.numRef.f)
            if got != want or s.smooth is not False or s.marker.symbol is not None \
                    or s.graphicalProperties.line.width != 12700:
                ok_series, why = False, (got, want)
                break
        check('%s: series titles from row 2, values rows 3..%d, dates as categories' % (name, N + 2), ok_series, why)
        check('%s: chart anchored one column right of the data' % name,
              ch.anchor._from.col == n + 2 and ch.anchor._from.row == 1, (ch.anchor._from.col, ch.anchor._from.row))
        check('%s: chart 32 x 16 cm' % name, ch.anchor.ext.cx == 11520000 and ch.anchor.ext.cy == 5760000)
    ws = wb[M.AUDIT_SHEET]
    total = sum(len(cs) for _, cs in results)
    check('Contracts: header + one row per contract', ws.max_row == total + 1 and [c.value for c in ws[1]] == M.AUDIT_COLUMNS,
          (ws.max_row, total))
    rows = {(r[0], r[1]): r for r in ws.iter_rows(min_row=2, values_only=True)}
    r = rows[('SIMSCI', 'Feb 26')]
    check('Contracts: NOT FOUND row - no ticker, no window, no rows, a note',
          r[5] == M.NOT_FOUND and r[4] is None and r[9] is None and r[13] is None and r[16], r)
    r = rows[('SIMSCI', 'Oct 26')]
    check('Contracts: NO DATA row - ticker, window, Rows == 0',
          r[5] == M.NO_DATA and r[4] == 'QZV6 Index' and r[9] is not None and r[13] == 0, r)
    c, r = by[('HSI', 'Sep 25')], rows[('HSI', 'Sep 25')]
    check('Contracts: OK row - both forms, ticker used, dates, rows, last and max OI',
          r[2] == 'HIU5 Index' and r[3] == 'HIU25 Index' and r[4] == c.ticker and r[5] == M.OK
          and r[6].date() == c.last_trade and r[7] == 'SEP 25' and r[9].date() == c.req_start
          and r[10].date() == c.req_end and r[11].date() == c.first_dt and r[12].date() == c.last_dt
          and r[13] == len(c.rows) and r[14] == c.last_oi and r[15] == c.max_oi, r)
    check('Contracts: frozen header', ws.freeze_panes == 'A2')
    empty = [M.Contract(product='EMPTY', root='ZZ', year=2024, month=m, label=M.month_label(2024, m),
                        ticker_1='a', ticker_2='b') for m in range(1, 13)]
    p2 = os.path.join(tmp, 'empty.xlsx')
    M.write_workbook(p2, [('EMPTY', empty)])
    wb2 = load_workbook(p2)
    check('product with nothing found: note in the tab, no chart, no crash',
          wb2.sheetnames == ['EMPTY', M.AUDIT_SHEET] and len(wb2['EMPTY']._charts) == 0
          and 'No contract' in str(wb2['EMPTY']['B1'].value))
    for p in (path, p2):
        os.remove(p)
    os.rmdir(tmp)


def test_guards():
    bbg = M.Bloomberg(blpapi_module=api_with(SilentSession)).connect()
    try:
        bbg.ref(['HIU6 Index'], M.REF_FIELDS)
        raised = ''
    except RuntimeError as e:
        raised = str(e)
    check('silent terminal -> one sentence after the spin budget, no hang', 'did not answer' in raised, raised)
    bbg = M.Bloomberg(blpapi_module=api_with(RejectingSession)).connect()
    try:
        bbg.history('HIU6 Index', 'OPEN_INT', dt.date(2024, 1, 1), dt.date(2024, 2, 1))
        raised = ''
    except RuntimeError as e:
        raised = str(e)
    check('REQUEST_STATUS (RequestFailure) -> one sentence', 'rejected the request' in raised, raised)
    bbg = M.Bloomberg(blpapi_module=api_with(ResponseErrorSession)).connect()
    try:
        bbg.ref(['HIU6 Index'], M.REF_FIELDS)
        raised = ''
    except RuntimeError as e:
        raised = str(e)
    check('responseError -> one sentence with Bloomberg\'s message', 'Not logged in' in raised, raised)
    bbg = M.Bloomberg(blpapi_module=FakeAPI).connect()
    out = bbg.ref(['HIU25 Index', 'XXXX FAKE Index'], ['LAST_TRADEABLE_DT', 'NOT_A_FIELD'])
    check('rejected mnemonic + dead ticker collected, good row still returned',
          bbg.field_errors.get('NOT_A_FIELD') == 'BAD_FLD / Invalid Field'
          and 'XXXX FAKE Index' in bbg.bad_securities and 'HIU25 Index' in out
          and isinstance(out['HIU25 Index']['LAST_TRADEABLE_DT'], dt.date))
    hist = bbg.history('XXXX FAKE Index', 'OPEN_INT', dt.date(2024, 1, 1), dt.date(2024, 2, 1))
    check('history of a dead ticker: empty, error collected', hist == [] and 'XXXX FAKE Index' in bbg.bad_securities)
    if importlib.util.find_spec('blpapi') is None:
        rc, text = quiet(M.main, ['--tickers-only'])
        check('no blpapi on this machine -> ERROR sentence, exit code 1', rc == 1 and 'blpapi is not installed' in text, text)


def test_run():
    tmp = tempfile.mkdtemp()
    path = os.path.join(tmp, 'run.xlsx')
    FakeSession.instances.clear()
    out, text = quiet(M.run, products=FAKE_PRODUCTS, today=TEST_TODAY, blpapi_module=FakeAPI, out=path)
    check('run(): returns the path and writes the file', out == path and os.path.exists(path))
    check('run(): per-product counts, NOT FOUND months, NO DATA months, output path printed',
          ('%-8s %2d/%d' % ('AS51', 12, 36)) in text and 'NOT FOUND: Jan 24' in text
          and 'NO DATA:  Oct 26' in text and ('Written: %s' % path) in text, text[-800:])
    check('run(): session closed', FakeSession.instances and FakeSession.instances[-1].stopped)
    FakeSession.instances.clear()
    out, text = quiet(M.run, products=FAKE_PRODUCTS, today=TEST_TODAY, blpapi_module=FakeAPI, tickers_only=True)
    sess = FakeSession.instances[-1]
    check('run(tickers_only): no historical request, table printed, returns ""',
          out == '' and not any(e['op'] == 'HistoricalDataRequest' for e in sess.log)
          and 'HIU25 Index' in text and 'NOT FOUND' in text)
    try:
        quiet(M.main, ['--help'])
        code = None
    except SystemExit as e:
        code = e.code
    check('--help exits 0', code == 0, code)
    os.remove(path)
    os.rmdir(tmp)


def demo(out=None):
    out = out or os.path.join(os.getcwd(), 'demo_OI_charts.xlsx')
    M.run(products=FAKE_PRODUCTS, today=TEST_TODAY, blpapi_module=FakeAPI, out=out)
    print('open it in Excel: 3 product tabs with a chart each, plus Contracts')


if __name__ == '__main__':
    if '--demo' in sys.argv:
        i = sys.argv.index('--demo')
        demo(sys.argv[i + 1] if len(sys.argv) > i + 1 else None)
    else:
        test_helpers()
        _bbg, _results = test_resolution()
        test_history(_bbg, _results)
        test_workbook(_results)
        test_guards()
        test_run()
        print('\n%d checks, %d failed' % (COUNT[0], len(FAILS)))
        for f in FAILS:
            print('  FAIL  ' + f)
        sys.exit(1 if FAILS else 0)
